# Wallbox-Artikelimport (v34, PLAN_V21 Phase 178, Stand Nachtrag 15.10.2026): liest den
# TAIFUN-Angebotsexport „WB Angebot Excel 1.xlsx“ im Ordner Artikel-Preislisten/Wallbox/
# (Blatt „Tabelle 1“, Spalten GUID · Position · Menge · Einheit (unbenannt) · Beschreibung ·
# E-Preis · G-Preis · Einkaufspreis Material; Werte als Text) und legt die Positionen als
# eigenen Wallbox-Artikelbereich WB001 … WB008 an (Quelle „wb“):
#
#   WB001/WB002  Mennekes AMTRON 4You 310 / 560          Pos. 001 / 002   (Stück)
#   WB003        Anschlussvorbereitung Ladestation       Pos. 003         (psl.)
#   WB004        Netzanmeldung beim EVU nach §14a        Pos. 005         (psl., 0,00 €)
#   WB005        lfd. Meter NYY-J / Erdung / CAT7        Pos. 006         (m)
#   WB006        Erdarbeiten bis zur Hauseinführung      Pos. 007         (m)
#   WB007        Fundament und Standsäule                Pos. 008         (Export „m“ → Blatt psl.)
#   WB008        Zusätzliche Bohrung/Durchbrüche         Pos. 004         (Stück)
#
# Anker ist das Blatt „WB-Artikel“ der Logik-Excel: Datei + Position; der
# Beschreibungsanfang (erste 40 Zeichen, Groß-/Kleinschreibung und doppelte Leerzeichen
# ignoriert) muss passen, sonst Fehler im Protokoll und die Zeile wird nicht importiert.
# Die Spalte „Einheit“ des Blatts überstimmt die Export-Einheit (Protokollzeile).
# Die GUIDs der Exporte sind Angebots-GUIDs (je Export verschieden) – sie werden nur als
# „zuletzt gesehen“ nachgeführt (artikel.guid und optional Blatt „WB-Artikel“).
# Weitere xlsx-Dateien im Ordner sind optional und dienen nur dem Preisvergleich: dieselbe
# Position mit anderem Preis → Preiskonflikt im Protokoll, die Ankerdatei gewinnt (Rückfrage 1).
# Die Textzeile ohne Position („Komplettpaket Wallbox inkl. Montage“) wird nicht importiert,
# sondern bereinigt als Absatz des WB-Vortexts angeboten (komplettpaket_text, seed_wb).
#
# Beschreibung: _x000D_\n → \n, „?“-Platzhalter der TAIFUN-Symbole samt umgebendem
# Doppel-Leerzeichen entfernen, Leerzeichen vor Satzzeichen („E-Auto ,“) entfernen, führende
# Leerzeichen je Zeile entfernen; Bezeichnung = erste Zeile, Beschreibung = Rest. VK = E-Preis;
# EK („Einkaufspreis Material“) nur in artikel.ek_cent – nie im Protokoll (Innendienst/Admin
# sehen das Protokoll, der Außendienst nie). Re-Import idempotent (Anker WB-Nr.); WP-/PV-/
# KL-Importe fassen WB nie an (eigene Quelle) und umgekehrt. v27: Blöcke je 200 Zeilen mit
# Commit, betrieb.import_markieren("Wallbox-Positionslisten").

import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

import openpyxl
from sqlalchemy.orm import Session

from app import config, models
from app.import_preisliste import (Diff, FELD_NAMEN, _header_indizes, euro_zu_cent)
from app.models import Artikel

# QUELLE_WB gehört nach app/models.py (Orchestrator) – bis dahin derselbe Wert als Fallback
QUELLE_WB = getattr(models, "QUELLE_WB", "wb")

WB_ORDNER_STANDARD = config.PROJEKT_ORDNER / "Artikel-Preislisten" / "Wallbox"
# Nachtrag 15.10.2026: EINE Pflichtdatei (TAIFUN-Export „WB Angebot Excel 1.xlsx“, Positionen
# 001–008 + Textzeile „Komplettpaket Wallbox inkl. Montage“); Extraleitung.xlsx entfällt. Weitere
# xlsx-Dateien im Ordner sind optional und werden nur auf Preiskonflikte geprüft (dieselbe WB-Nr.
# mit anderem Preis → Ankerdatei laut Blatt „WB-Artikel“ gewinnt, Konflikt im Protokoll).
PFLICHT_DATEI = ("WB Angebot Excel 1", "WB Angebot Excel 1.xlsx")
DATEIEN = {PFLICHT_DATEI[0]: PFLICHT_DATEI[1]}
BLATT_ARTIKEL = "WB-Artikel"
BLATT_NAME = "Tabelle 1"
ANFANG_ZEICHEN = 40
KATEGORIE = "WB · Wallbox"
FELDER_WB = ["kategorie", "bezeichnung", "beschreibung", "menge_standard", "einheit",
             "e_preis_cent", "ep_flag", "ek_cent"]
# v27 (PLAN_V17 Phase 128): Blockgröße – Commit je BLOCK geschriebene Zeilen
BLOCK = 200
IMPORT_NAME = "Wallbox-Positionslisten"
MARKER_HINWEIS = "migration_v34_wallbox_hinweis"
HINWEIS_DATEIEN_FEHLEN = ("Wallbox-Preislisten fehlen – Import übersprungen "
                          "(Artikel-Preislisten/Wallbox/)")
EINHEITEN = {"psl": "psl.", "psl.": "psl.", "pausch": "psl.", "pauschal": "psl.", "pau": "psl.",
             "stück": "Stück", "stk": "Stück", "stk.": "Stück", "st": "Stück", "st.": "Stück",
             "stueck": "Stück", "m": "m", "lfm": "m", "lfd. m": "m", "lfd.m": "m", "meter": "m"}


def wb_ordner() -> Path:
    """Ordner der Wallbox-Positionslisten (config.WB_PREISLISTEN_ORDNER, sonst Standard)."""
    return Path(getattr(config, "WB_PREISLISTEN_ORDNER", WB_ORDNER_STANDARD))


def wb_nummer(zahl: int) -> str:
    return f"WB{zahl:03d}"


# --- Textbereinigung ----------------------------------------------------------

def text_bereinigen(wert) -> str:
    """TAIFUN-Artefakte: `_x000D_` (CR der Excel-Zelle) → Zeilenumbruch, „?“-Platzhalter
    der TAIFUN-Symbole (alleinstehende Fragezeichen – echte Fragezeichen am Wortende
    bleiben) samt umgebendem Doppel-Leerzeichen entfernen, führende/trailing Leerzeichen
    je Zeile entfernen, Leerzeilen am Rand entfernen."""
    if wert is None:
        return ""
    text = str(wert).replace("_x000D_\r\n", "\n").replace("_x000D_\n", "\n")
    text = text.replace("_x000D_", "\n").replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\xad", "")
    zeilen = []
    for zeile in text.split("\n"):
        zeile = re.sub(r"(?:(?<=\s)|^)\?+(?=\s|$)", "", zeile)   # „? Produktvorteile“ → „Produktvorteile“
        zeile = re.sub(r"\s+([,.;:!])", r"\1", zeile)            # „E-Autos , sondern“ → „E-Autos, sondern“
        zeile = re.sub(r"[ \t]{2,}", " ", zeile).strip()
        zeilen.append(zeile)
    while zeilen and not zeilen[0]:
        zeilen.pop(0)
    while zeilen and not zeilen[-1]:
        zeilen.pop()
    return "\n".join(zeilen)


def anfang_norm(text: str, zeichen: int = ANFANG_ZEICHEN) -> str:
    """Vergleichsform des Beschreibungsanfangs: Kleinschreibung, einfache Leerzeichen,
    die ersten `zeichen` Zeichen."""
    return re.sub(r"\s+", " ", str(text or "").strip().lower())[:zeichen]


def anfang_passt(soll: str, ist: str) -> bool:
    """Blatt-Anfang (bis 40 Zeichen) muss Präfix des bereinigten Dateitexts sein."""
    soll_n = anfang_norm(soll)
    if not soll_n:
        return False
    return anfang_norm(ist, len(soll_n)) == soll_n


def einheit_norm(wert: str, standard: str = "") -> str:
    roh = str(wert or "").strip()
    if not roh:
        return standard
    return EINHEITEN.get(roh.lower(), roh)


def erste_zeile(text: str) -> str:
    for zeile in (text or "").splitlines():
        if zeile.strip():
            return zeile.strip()
    return ""


# --- Anker aus der Logik-Excel ---------------------------------------------------

def _logik_workbook():
    pfad = config.LOGIK_EXCEL_PFAD
    if not pfad.exists():
        return None
    return openpyxl.load_workbook(pfad, read_only=True, data_only=True)


def anker_lesen(wb=None) -> list[dict]:
    """Blatt „WB-Artikel“ der Logik-Excel: Zeilen {nr, datei, position, anfang, guid,
    einheit, bemerkung} in Blattreihenfolge (leer = Blatt fehlt)."""
    schliessen = False
    if wb is None:
        wb = _logik_workbook()
        schliessen = wb is not None
    if wb is None or BLATT_ARTIKEL not in wb.sheetnames:
        return []
    try:
        zeilen = wb[BLATT_ARTIKEL].iter_rows(values_only=True)
        kopf = [str(z or "").strip() for z in next(zeilen, [])]
        spalten = {name: kopf.index(name) for name in kopf if name}

        def wert(row, name) -> str:
            i = spalten.get(name)
            return str(row[i] or "").strip() if i is not None and i < len(row) else ""

        anker = []
        for row in zeilen:
            row = tuple(row)
            nr = wert(row, "WB-Nr.")
            if not re.fullmatch(r"WB\d{3}", nr):
                continue
            anker.append({"nr": nr, "datei": wert(row, "Datei"),
                          "position": wert(row, "Position").zfill(3) if wert(row, "Position").isdigit()
                          else wert(row, "Position"),
                          "anfang": wert(row, "Beschreibungsanfang"),
                          "guid": wert(row, "GUID (zuletzt gesehen)"),
                          "einheit": wert(row, "Einheit"), "bemerkung": wert(row, "Bemerkung")})
        return anker
    finally:
        if schliessen:
            wb.close()


def _datei_kurz(name: str) -> str:
    """„Wallbox Artikel Excel.xlsx“ / „wallbox artikel excel“ → Vergleichsschlüssel."""
    return re.sub(r"\s+", " ", re.sub(r"\.xlsx?$", "", str(name or "").strip(), flags=re.I)).lower()


def datei_finden(ordner: Path, kurz: str) -> Optional[Path]:
    """Positionsliste im Ordner (tolerant: mit/ohne .xlsx, Groß-/Kleinschreibung,
    doppelte Leerzeichen)."""
    ziel = _datei_kurz(kurz)
    try:
        for kandidat in sorted(Path(ordner).iterdir()):
            if (kandidat.is_file() and kandidat.suffix.lower() in (".xlsx", ".xlsm")
                    and _datei_kurz(kandidat.name) == ziel):
                return kandidat
    except OSError:
        return None
    return None


def anker_dateien(anker: Optional[list[dict]] = None) -> list[str]:
    """Kurznamen der Pflichtdateien = die im Blatt „WB-Artikel“ genannten Dateien (Reihenfolge
    des Blatts); ohne Blatt die Pflichtdatei des Nachtrags."""
    anker = anker_lesen() if anker is None else anker
    namen: list[str] = []
    for a in anker:
        kurz = str(a.get("datei") or "").strip()
        if kurz and _datei_kurz(kurz) not in {_datei_kurz(n) for n in namen}:
            namen.append(kurz)
    return namen or [PFLICHT_DATEI[0]]


def dateien_vorhanden(ordner: Optional[Path] = None,
                      anker: Optional[list[dict]] = None) -> tuple[bool, list[str]]:
    """(alle Pflichtdateien da?, fehlende Dateinamen) – für Parametrierung und migrate."""
    ordner = Path(ordner) if ordner is not None else wb_ordner()
    fehlend = [f"{kurz}.xlsx" for kurz in anker_dateien(anker)
               if not ordner.exists() or datei_finden(ordner, kurz) is None]
    return (not fehlend, fehlend)


def weitere_listen(ordner: Path, ausgenommen: list[str]) -> list[Path]:
    """Optionale weitere xlsx-Dateien im Ordner (nicht die Pflichtdateien, keine Excel-
    Sperrdateien „~$…“) – nur für den Preisvergleich."""
    bekannt = {_datei_kurz(k) for k in ausgenommen}
    try:
        return sorted(p for p in Path(ordner).iterdir()
                      if p.is_file() and p.suffix.lower() in (".xlsx", ".xlsm")
                      and not p.name.startswith("~$") and _datei_kurz(p.name) not in bekannt)
    except OSError:
        return []


def komplettpaket_text(ordner: Optional[Path] = None) -> str:
    """Text der Zeile ohne Position „Komplettpaket Wallbox inkl. Montage“ aus der
    Pflichtdatei (bereinigt: Umbrüche, „?“-Platzhalter, Leerzeichen) – Absatz am Ende des
    Vortexts (angebotsprofile.seed_wb); leer, wenn Datei oder Zeile fehlt."""
    ordner = Path(ordner) if ordner is not None else wb_ordner()
    if not ordner.exists():
        return ""
    for kurz in anker_dateien():
        pfad = datei_finden(ordner, kurz)
        if pfad is None:
            continue
        try:
            for pos, _guid, roh, _einheit, _e, _ek in _positionszeilen(pfad):
                if pos is not None:
                    continue
                text = text_bereinigen(roh)
                if erste_zeile(text).lower().startswith("komplettpaket"):
                    return text
        except (OSError, ValueError, KeyError):
            continue
    return ""


# --- Dateien lesen -----------------------------------------------------------------

@dataclass
class WbLeseErgebnis:
    artikel: list[dict] = field(default_factory=list)
    warnungen: list[str] = field(default_factory=list)   # Hinweise/Fehler (ohne EK!)
    protokoll: list[str] = field(default_factory=list)   # Datei, Position, WB-Nr., VK, GUID
    konflikte: list[str] = field(default_factory=list)   # Preiskonflikte zwischen den Dateien
    textzeilen: list[str] = field(default_factory=list)  # Zeilen ohne Position (Vortext)
    guid_wechsel: dict[str, str] = field(default_factory=dict)   # WB-Nr. → neue GUID


def _positionszeilen(pfad: Path):
    """(pos, guid, beschreibung_roh, einheit, e_preis, ek) je Zeile; pos None = Textzeile."""
    wb = openpyxl.load_workbook(pfad, data_only=True, read_only=True)
    try:
        ws = wb[BLATT_NAME] if BLATT_NAME in wb.sheetnames else wb[wb.sheetnames[0]]
        zeilen = ws.iter_rows(values_only=True)
        indizes = _header_indizes(next(zeilen), [])

        def zelle(row, name):
            i = indizes.get(name)
            return row[i] if i is not None and i < len(row) else None

        for row in zeilen:
            guid, beschreibung, pos = zelle(row, "GUID"), zelle(row, "Beschreibung"), zelle(row, "Position")
            if guid is None and beschreibung is None:
                continue
            pos_text = None
            if pos not in (None, ""):
                pos_text = str(pos).strip()
                if re.fullmatch(r"\d+(?:\.0+)?", pos_text):
                    pos_text = str(int(float(pos_text))).zfill(3)
            yield (pos_text, (str(guid).strip() if guid not in (None, "") else None),
                   beschreibung, str(zelle(row, "Einheit") or "").strip(),
                   zelle(row, "E-Preis"), zelle(row, "Einkaufspreis Material"))
    finally:
        wb.close()


def _euro(cent: int) -> str:
    text = f"{cent / 100:,.2f}"
    return text.replace(",", "X").replace(".", ",").replace("X", ".") + " €"


def _artikel(anker: dict, guid, text: str, einheit: str, e_preis, ek, kurz: str, pos: str) -> dict:
    """Artikeldaten einer Positionszeile. Einheit: die Spalte „Einheit“ des Blatts „WB-Artikel“
    überstimmt die Export-Einheit (Nachtrag 15.10.2026: Fundament/Standsäule im Export „m“, im
    Tool 1 × psl.); leer im Blatt → Export-Einheit."""
    bezeichnung = erste_zeile(text)
    rest = [z.strip() for z in text.splitlines()[1:]] if "\n" in text else []
    return {
        "guid": guid,
        "pos_nr": anker["nr"],
        "kategorie": KATEGORIE,
        "bezeichnung": bezeichnung,
        "beschreibung": "\n".join(z for z in rest if z),
        "menge_standard": 1.0,
        "einheit": einheit_norm(anker.get("einheit", "")) or einheit_norm(einheit),
        "e_preis_cent": euro_zu_cent(e_preis or 0),
        "ep_flag": False,
        "ek_cent": euro_zu_cent(ek) if ek not in (None, "") else None,
        "quelle": QUELLE_WB,
        "_datei": kurz,
        "_position": pos,
    }


def lese_wb_dateien(ordner: Optional[Path] = None, anker: Optional[list[dict]] = None) -> WbLeseErgebnis:
    """Liest die Pflichtdatei(en) laut Blatt „WB-Artikel“ (Datei + Position, Beschreibungs-
    anfang geprüft) und optionale weitere xlsx-Dateien im Ordner. Dieselbe Position in einer
    anderen Datei (gleicher Beschreibungsanfang) wird nur verglichen: gleicher Preis → Hinweis
    „identisch“, anderer Preis → Konflikt (die Ankerdatei gewinnt)."""
    ergebnis = WbLeseErgebnis()
    ordner = Path(ordner) if ordner is not None else wb_ordner()
    anker = anker_lesen() if anker is None else anker
    if not anker:
        ergebnis.warnungen.append("Blatt „WB-Artikel“ der Logik-Excel fehlt oder ist leer – "
                                  "kein Anker für den Wallbox-Import.")
        return ergebnis
    nach_datei_pos = {(_datei_kurz(a["datei"]), a["position"]): a for a in anker}
    importiert: dict[str, dict] = {}            # WB-Nr. → Artikel-Daten
    kandidaten: list[tuple] = []                # Zeilen anderer Dateien zum Preisvergleich
    gelesen = 0
    pflicht = anker_dateien(anker)
    quellen: list[tuple[str, Optional[Path], bool]] = []
    for kurz in pflicht:
        quellen.append((kurz, datei_finden(ordner, kurz) if ordner.exists() else None, True))
    if ordner.exists():
        for pfad in weitere_listen(ordner, pflicht):
            quellen.append((pfad.stem, pfad, False))
    for kurz, pfad, ist_pflicht in quellen:
        if pfad is None:
            ergebnis.warnungen.append(f"Wallbox-Positionsliste fehlt: {ordner / (kurz + '.xlsx')}")
            continue
        try:
            zeilen = list(_positionszeilen(pfad))
        except (OSError, ValueError, KeyError) as exc:
            ergebnis.warnungen.append(f"{kurz}: nicht lesbar ({exc}) – übersprungen.")
            continue
        if ist_pflicht:
            gelesen += 1
        else:
            ergebnis.protokoll.append(f"{kurz}: weitere Liste (optional) – nur Preisvergleich")
        for pos, guid, roh, einheit, e_preis, ek in zeilen:
            text = text_bereinigen(roh)
            if pos is None:
                if text:
                    ergebnis.textzeilen.append(text)
                    ergebnis.protokoll.append(
                        f"{kurz}: Textzeile ohne Position übersprungen – „{erste_zeile(text)[:60]}“ "
                        "(Vortext „Komplettpaket“, Parametrierung → Textblöcke)")
                continue
            if not text:
                ergebnis.warnungen.append(f"{kurz} Pos. {pos}: leere Beschreibung – übersprungen.")
                continue
            zeile_anker = nach_datei_pos.get((_datei_kurz(kurz), pos))
            if zeile_anker is not None:
                if not anfang_passt(zeile_anker["anfang"], text):
                    ergebnis.warnungen.append(
                        f"FEHLER {kurz} Pos. {pos}: Beschreibungsanfang „{erste_zeile(text)[:40]}“ "
                        f"passt nicht zum Blatt „WB-Artikel“ ({zeile_anker['nr']}: "
                        f"„{zeile_anker['anfang'][:40]}“) – Zeile nicht importiert.")
                    continue
                if zeile_anker["nr"] in importiert:
                    ergebnis.warnungen.append(
                        f"{kurz} Pos. {pos}: Nummer {zeile_anker['nr']} ist bereits belegt – übersprungen.")
                    continue
                daten = _artikel(zeile_anker, guid, text, einheit, e_preis, ek, kurz, pos)
                importiert[zeile_anker["nr"]] = daten
                ergebnis.artikel.append(daten)
                guid_text = f", GUID {guid}" if guid else ""
                ergebnis.protokoll.append(
                    f"{kurz} Pos. {pos} → {zeile_anker['nr']} „{daten['bezeichnung'][:50]}“: "
                    f"VK {_euro(daten['e_preis_cent'])} / {daten['einheit']}{guid_text}")
                export_einheit = einheit_norm(einheit)
                if export_einheit and export_einheit != daten["einheit"]:
                    ergebnis.protokoll.append(
                        f"{zeile_anker['nr']}: Einheit laut Export „{export_einheit}“, laut Blatt "
                        f"„WB-Artikel“ „{daten['einheit']}“ – das Blatt gilt")
                if guid and zeile_anker.get("guid") and zeile_anker["guid"] != guid:
                    ergebnis.protokoll.append(
                        f"{zeile_anker['nr']}: GUID-Wechsel {zeile_anker['guid']} → {guid} "
                        "(Angebots-GUID des neuen Exports – Blatt „WB-Artikel“ „zuletzt gesehen“)")
                if guid and zeile_anker.get("guid") != guid:
                    ergebnis.guid_wechsel[zeile_anker["nr"]] = guid
                continue
            # keine Ankerzeile für Datei + Position: Doppelgänger der anderen Datei?
            treffer = next((a for a in anker if anfang_passt(a["anfang"], text)), None)
            if treffer is None:
                ergebnis.protokoll.append(
                    f"{kurz} Pos. {pos}: „{erste_zeile(text)[:50]}“ nicht im Blatt „WB-Artikel“ – übersprungen")
                continue
            kandidaten.append((treffer["nr"], kurz, pos, euro_zu_cent(e_preis or 0)))
    # Preisvergleich der Doppelgänger (Rückfrage 1): die Ankerdatei gewinnt
    for nr, kurz, pos, preis_cent in kandidaten:
        gewinner = importiert.get(nr)
        if gewinner is None:
            ergebnis.warnungen.append(
                f"{kurz} Pos. {pos}: gleiche Position wie {nr}, aber die Ankerdatei "
                f"(Blatt „WB-Artikel“) liefert {nr} nicht – nicht importiert.")
            continue
        if gewinner["e_preis_cent"] != preis_cent:
            ergebnis.konflikte.append(
                f"Preiskonflikt {nr} „{gewinner['bezeichnung'][:40]}“: {kurz} Pos. {pos} nennt "
                f"{_euro(preis_cent)}, {gewinner['_datei']} Pos. {gewinner['_position']} nennt "
                f"{_euro(gewinner['e_preis_cent'])} – es gilt {_euro(gewinner['e_preis_cent'])} "
                f"aus „{gewinner['_datei']}“ (Blatt „WB-Artikel“, Rückfrage 1).")
        else:
            ergebnis.protokoll.append(
                f"{kurz} Pos. {pos}: identisch mit {nr} ({_euro(preis_cent)}) – nur einmal importiert")
    for a in anker:
        if a["nr"] not in importiert and gelesen:
            ergebnis.warnungen.append(
                f"{a['nr']} ({a['datei']} Pos. {a['position']}) in den Positionslisten nicht gefunden.")
    ergebnis.warnungen.extend(ergebnis.konflikte)
    return ergebnis


# --- Abgleich mit der Datenbank ----------------------------------------------------

@dataclass
class WbDiff(Diff):
    protokoll: list[str] = field(default_factory=list)
    konflikte: list[str] = field(default_factory=list)
    textzeilen: list[str] = field(default_factory=list)


def _zuordnen(session: Session, ergebnis: WbLeseErgebnis):
    """Jede gelesene Zeile → Bestandsartikel über die WB-Nummer (nie über die GUID)."""
    bestand = session.query(Artikel).filter(Artikel.quelle == QUELLE_WB).all()
    nach_nr = {a.pos_nr: a for a in bestand}
    paare = [(daten, nach_nr.get(daten["pos_nr"])) for daten in ergebnis.artikel]
    return paare, bestand


def berechne_diff(session: Session, ergebnis: Optional[WbLeseErgebnis] = None,
                  ordner: Optional[Path] = None) -> WbDiff:
    ergebnis = ergebnis or lese_wb_dateien(ordner)
    diff = WbDiff(warnungen=list(ergebnis.warnungen), protokoll=list(ergebnis.protokoll),
                  konflikte=list(ergebnis.konflikte), textzeilen=list(ergebnis.textzeilen))
    paare, bestand = _zuordnen(session, ergebnis)
    gefunden = set()
    for daten, vorhanden in paare:
        if vorhanden is None:
            diff.neu.append(daten)
            continue
        gefunden.add(vorhanden.id)
        felder = [f for f in FELDER_WB if getattr(vorhanden, f) != daten[f]]
        if daten["guid"] and vorhanden.guid != daten["guid"]:
            felder.append("guid")
            if vorhanden.guid:
                diff.protokoll.append(
                    f"{daten['pos_nr']}: GUID-Wechsel {vorhanden.guid} → {daten['guid']} "
                    "(Angebots-GUID des neuen Exports – nur „zuletzt gesehen“, Anker bleibt die WB-Nr.)")
        if felder:
            diff.geaendert.append((vorhanden, daten,
                                   [FELD_NAMEN.get(f, "GUID") for f in felder]))
        else:
            diff.unveraendert += 1
        if not vorhanden.aktiv:
            diff.reaktiviert.append(vorhanden)
    if ergebnis.artikel:
        diff.entfallen = [a for a in bestand if a.id not in gefunden and a.aktiv]
        for a in diff.entfallen:
            diff.warnungen.append(
                f"{a.pos_nr} „{a.titel[:60]}“ ist in keiner Wallbox-Positionsliste mehr – wird deaktiviert.")
    return diff


def import_ausfuehren(session: Session, ordner: Optional[Path] = None,
                      guid_nachfuehren: bool = False) -> tuple[WbDiff, str]:
    """Wendet den Import an; Rückgabe (Diff, Meldung) wie import_pv/import_klima.
    Anker je Zeile ist die WB-Nummer; jede Zeile ist ein Upsert, Verwaiste werden erst
    nach allen Zeilen deaktiviert (nur wenn mindestens eine Datei gelesen wurde).
    v27: Commit je BLOCK Zeilen, Import im Betriebs-Status. `guid_nachfuehren=True`
    schreibt die aktuellen GUIDs nach dem Commit in das Blatt „WB-Artikel“ (nur aus der
    Parametrierung – migrate.py lässt die Live-Excel unangetastet)."""
    from app import betrieb
    with betrieb.import_markieren(IMPORT_NAME):
        ergebnis = lese_wb_dateien(ordner)
        diff = berechne_diff(session, ergebnis)
        if not ergebnis.artikel:
            session.rollback()
            return diff, "WB: keine Positionsliste gelesen – nichts importiert"
        paare, _bestand = _zuordnen(session, ergebnis)
        fremde_guids = {g for (g,) in session.query(Artikel.guid)
                        .filter(Artikel.quelle != QUELLE_WB, Artikel.guid.isnot(None))}
        zaehler = 0
        for daten, vorhanden in paare:
            werte = {k: v for k, v in daten.items() if not k.startswith("_")}
            if werte["guid"] in fremde_guids:    # Angebots-GUID gehört bereits einem WP-/PV-/KL-Artikel
                diff.warnungen.append(
                    f"{werte['pos_nr']}: GUID {werte['guid']} gehört bereits einem anderen "
                    "Artikel – ohne GUID angelegt.")
                werte["guid"] = None
            if vorhanden is None:
                session.add(Artikel(**werte, aktiv=True))
            else:
                for feld in FELDER_WB + ["pos_nr", "quelle"]:
                    setattr(vorhanden, feld, werte[feld])
                if werte["guid"] or vorhanden.guid != werte["guid"]:
                    vorhanden.guid = werte["guid"]
                vorhanden.aktiv = True
            zaehler += 1
            if zaehler % BLOCK == 0:
                session.commit()      # Block abschließen – Schreibsperre freigeben
        for artikel in diff.entfallen:
            artikel.aktiv = False
        session.commit()
    if guid_nachfuehren and ergebnis.guid_wechsel:
        diff.protokoll.append(guid_nachfuehren_blatt(ergebnis.guid_wechsel))
    meldung = (f"WB: {len(diff.neu)} neu, {len(diff.geaendert)} geändert, "
               f"{diff.unveraendert} unverändert, {len(diff.entfallen)} deaktiviert")
    if diff.konflikte:
        meldung += f", {len(diff.konflikte)} Preiskonflikt(e) – siehe Protokoll"
    return diff, meldung


def guid_nachfuehren_blatt(guid_wechsel: dict[str, str], pfad: Optional[Path] = None) -> str:
    """Spalte „GUID (zuletzt gesehen)“ im Blatt „WB-Artikel“ der Logik-Excel nachführen
    (best effort – die Datei kann in Excel geöffnet sein). Liefert die Protokollzeile.
    Hinweis Betrieb: schreibt in die Live-Excel; auf dem Server ist die Datei
    git-versioniert (update.bat → git pull) – Einstellung `wb_guid_nachfuehren = 0`
    schaltet es ab (siehe Parametrierung → Logik & Importe)."""
    pfad = Path(pfad) if pfad is not None else config.LOGIK_EXCEL_PFAD
    try:
        wb = openpyxl.load_workbook(pfad)
        if BLATT_ARTIKEL not in wb.sheetnames:
            return "GUID „zuletzt gesehen“: Blatt „WB-Artikel“ fehlt – nicht nachgeführt."
        ws = wb[BLATT_ARTIKEL]
        kopf = [str(c.value or "").strip() for c in ws[1]]
        if "WB-Nr." not in kopf or "GUID (zuletzt gesehen)" not in kopf:
            return "GUID „zuletzt gesehen“: Spalte fehlt im Blatt „WB-Artikel“ – nicht nachgeführt."
        i_nr, i_guid = kopf.index("WB-Nr.") + 1, kopf.index("GUID (zuletzt gesehen)") + 1
        geaendert = 0
        for r in range(2, ws.max_row + 1):
            nr = str(ws.cell(r, i_nr).value or "").strip()
            if nr in guid_wechsel and str(ws.cell(r, i_guid).value or "").strip() != guid_wechsel[nr]:
                ws.cell(r, i_guid).value = guid_wechsel[nr]
                geaendert += 1
        if geaendert:
            wb.save(pfad)
        return f"GUID „zuletzt gesehen“ im Blatt „WB-Artikel“ nachgeführt ({geaendert} Zeilen)."
    except (OSError, PermissionError) as exc:
        return (f"GUID „zuletzt gesehen“ konnte nicht in die Logik-Excel geschrieben werden "
                f"({exc.__class__.__name__}: Datei geöffnet?) – Artikelstamm ist aktuell.")


# --- Hilfen für migrate.py / Parametrierung -----------------------------------------

def wb_referenzen() -> list[str]:
    """WB-Nummern, die die Logik-Blätter referenzieren (über logik.wb_referenzen);
    sonst alle Nummern des Blatts „WB-Artikel“."""
    try:
        from app import logik as _logik
        l, _b = _logik.logik_einlesen()
        refs = sorted(_logik.wb_referenzen(l))
        if refs:
            return refs
    except Exception:   # noqa: BLE001 – Logik-Excel fehlt/defekt: Fallback Blatt
        pass
    try:
        return sorted({a["nr"] for a in anker_lesen()})
    except Exception:   # noqa: BLE001
        return []


def fehlende_wb_artikel(session: Session) -> list[str]:
    """Referenzierte WB-Nummern ohne aktiven Artikel (Auto-Import in migrate.py)."""
    refs = wb_referenzen()
    if not refs:
        return []
    vorhanden = {nr for (nr,) in session.query(Artikel.pos_nr)
                 .filter(Artikel.pos_nr.like("WB%"), Artikel.aktiv.is_(True))}
    return [r for r in refs if r not in vorhanden]


def migration_v34_wallbox(session: Session) -> list[str]:
    """migrate.py-Hook (Orchestrator registriert ihn): importiert die Wallbox-
    Positionslisten automatisch, wenn referenzierte WB-Artikel fehlen UND die Dateien
    in Artikel-Preislisten/Wallbox/ liegen. Fehlen die Dateien, meldet der erste Lauf
    „Wallbox-Preislisten fehlen – Import übersprungen (…)“; weitere Läufe prüfen jedes
    Mal, melden aber nur bei Änderung (Fingerabdruck fehlende Dateien + fehlende
    Referenzen in der Einstellung `migration_v34_wallbox_hinweis`). Legt zusätzlich die
    WB-Textblöcke an (angebotsprofile.seed_wb, idempotent). Die Live-Excel wird hier
    nie beschrieben (guid_nachfuehren=False)."""
    from app.models import einstellung_holen, einstellung_setzen
    meldungen: list[str] = []
    try:
        from app import angebotsprofile
        meldungen += list(angebotsprofile.seed_wb(session) or [])
    except Exception as exc:   # noqa: BLE001 – Texte dürfen den Import nie blockieren
        meldungen.append(f"WARNUNG: WB-Textblöcke (seed_wb) schlugen fehl: {exc}")
    try:
        fehlend = fehlende_wb_artikel(session)
    except Exception as exc:   # noqa: BLE001
        meldungen.append(f"WARNUNG: WB-Artikelreferenzen nicht ermittelbar: {exc}")
        return meldungen
    if not fehlend:
        if einstellung_holen(session, MARKER_HINWEIS, ""):
            einstellung_setzen(session, MARKER_HINWEIS, "")   # Zustand erledigt – Marker löschen
        return meldungen
    ok, fehlende_dateien = dateien_vorhanden()
    if ok:
        try:
            _diff, meldung = import_ausfuehren(session)
            meldungen.append(f"WB-Artikel fehlten ({len(fehlend)}, z. B. {fehlend[0]}) – "
                             f"Wallbox-Positionslisten-Import ausgeführt ({meldung})")
            einstellung_setzen(session, MARKER_HINWEIS, "")
        except Exception as exc:   # noqa: BLE001
            meldungen.append(f"WARNUNG: WB-Artikel fehlen und der Wallbox-Import schlug fehl: {exc}")
        return meldungen
    fingerabdruck = json.dumps({"dateien": sorted(fehlende_dateien), "refs": sorted(fehlend)},
                               ensure_ascii=False)
    if einstellung_holen(session, MARKER_HINWEIS, "") != fingerabdruck:
        einstellung_setzen(session, MARKER_HINWEIS, fingerabdruck)
        meldungen.append(HINWEIS_DATEIEN_FEHLEN)
    return meldungen
