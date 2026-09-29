# Bestandsimport laufender Projekte (PLAN_PROJ_V3 Phase 85): Die meisten
# laufenden Aufträge sind nie durch das Tool gelaufen. Eine Excel-Vorlage
# (eine Zeile je Gewerk) wird hochgeladen, je Zeile geprüft (Vorschau) und nur
# fehlerfreie Zeilen werden importiert. Idempotent über TAIFUN-Nummer + Sparte:
# ein zweiter Lauf aktualisiert statt zu duplizieren. Jeder Import wird
# protokolliert und lässt sich rückgängig machen (nur, was er angelegt hat und
# was seither unverändert ist). Bestandsgewerke tragen das Badge „Bestand“,
# ihre Angebote zählen nicht in der Statistik und lösen keine monday-
# Rückspielung aus.

import io
import json
from dataclasses import dataclass, field
from datetime import date, datetime

from sqlalchemy.orm import Session

from app import projektierung as kern
from app.models import (Angebot, Aufgabe, AufgabenpaketInstanz, Benutzer,
                        Bestandsimport, Gewerk, Kunde, Projekt, ProjektDokument,
                        ProjektTermin, ProjektVerlauf, SteckbriefWert, Team,
                        Vorgang, GEWERK_PHASEN_NAMEN, angebot_status_setzen)

BLATT = "Projekte"

# (Schlüssel, Spaltenüberschrift, Pflicht, Hinweis für das Blatt „Anleitung“)
SPALTEN = [
    ("anrede", "Kunde Anrede", False, "Herr | Frau | Firma"),
    ("vorname", "Vorname", False, ""),
    ("nachname", "Nachname", True, "bei Firmen: Firmenname"),
    ("strasse", "Straße", False, "Straße und Hausnummer (Rechnungs-/Kundenadresse)"),
    ("plz", "PLZ", True, "5 Ziffern"),
    ("ort", "Ort", False, ""),
    ("telefon", "Telefon", False, ""),
    ("email", "E-Mail", False, ""),
    ("ausfuehrung", "Ausführungsadresse (falls abweichend)",
     False, "„Straße Nr, PLZ Ort“ – leer = Kundenadresse"),
    ("sparte", "Sparte", True, "WP | PV | KL | WB"),
    ("taifun_nummer", "TAIFUN-Angebotsnummer", True,
     "z. B. AN261234 – zusammen mit der Sparte der Schlüssel (zweiter Lauf aktualisiert)"),
    ("auftragswert", "Auftragswert brutto", True, "Euro, z. B. 32.450,00"),
    ("auftragsdatum", "Auftragsdatum", True, "TT.MM.JJJJ"),
    ("vertriebler", "Vertriebler (Name)", False, "wie in der Benutzerverwaltung"),
    ("kanal", "Kanal", False, "Standard | Enni | SWD | Sparkasse DU …"),
    ("projektleiter", "Projektleiter (Name)", False,
     "wie in der Benutzerverwaltung; leer = Standard-Projektleiter"),
    ("phase", "Phase", True,
     "Auftragseingang | Feinplanung VOT | Planung | Montagevorbereitung | "
     "Montage | Abnahme | Freigabe"),
    ("montage_von", "Montage von", False, "TT.MM.JJJJ"),
    ("montage_bis", "Montage bis", False, "TT.MM.JJJJ (leer = von + 4 Arbeitstage)"),
    ("montageteam", "Montageteam", False, "Teamname, z. B. Montageteam 1"),
    ("elektroteam", "Elektro-Team", False, "Teamname"),
    ("subteam", "Subteam", False, "Teamname, z. B. Subteam 1"),
    ("bestaetigt", "Termin bestätigt (J/N)", False, "J | N"),
    ("hersteller", "Hersteller", False, "Steckbrief"),
    ("leistungsklasse", "Leistungsklasse", False, "Steckbrief, kW"),
    ("innengeraet", "Innengerät", False, "Steckbrief"),
    ("zaehlerschrank", "Zählerschrank", False, "Steckbrief"),
    ("oeltank", "Öltankentsorgung (J/N)", False, "Steckbrief"),
    ("bemerkung", "Bemerkung", False, "landet als Kopfnotiz am Projekt"),
]
PHASEN_AUSWAHL = ["auftragseingang", "feinplanung_vot", "planung",
                  "montagevorbereitung", "montage", "abnahme", "freigabe"]
STECKBRIEF_SPALTEN = ("hersteller", "leistungsklasse", "innengeraet",
                      "zaehlerschrank", "oeltank")


# --- Vorlage ------------------------------------------------------------------

def vorlage_bytes() -> bytes:
    """Excel-Vorlage mit Blatt „Projekte“ (Kopfzeile + Beispielzeile) und
    Blatt „Anleitung“ (Ausfüllhinweise, erlaubte Werte)."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = BLATT
    ws.append([titel for _, titel, _, _ in SPALTEN])
    for zelle in ws[1]:
        zelle.font = Font(bold=True, color="FFFFFF")
        zelle.fill = PatternFill("solid", fgColor="1F4E79")
    ws.append(["Herr", "Max", "Beispiel", "Musterweg 1", "47139", "Duisburg",
               "0203 123456", "max@beispiel.de", "", "WP", "AN269999",
               "32.450,00", "15.08.2026", "", "Standard", "", "Planung",
               "", "", "", "", "", "N", "Bosch", "10", "AWM", "bleibt", "J",
               "BEISPIELZEILE – vor dem Import löschen"])
    for i, (_, titel, _, _) in enumerate(SPALTEN, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(12, len(titel) + 2)
    ws.freeze_panes = "A2"
    anleitung = wb.create_sheet("Anleitung")
    anleitung.append(["Spalte", "Pflicht", "Erlaubte Werte / Hinweis"])
    for zelle in anleitung[1]:
        zelle.font = Font(bold=True)
    for _, titel, pflicht, hinweis in SPALTEN:
        anleitung.append([titel, "ja" if pflicht else "", hinweis])
    anleitung.append([])
    for zeile in (
            "Eine Zeile je Gewerk (zwei Sparten beim selben Kunden = zwei Zeilen).",
            "Nur OFFENE Projekte eintragen; die Phase ehrlich setzen.",
            "Alle Pflichtaufgaben der Pakete VOR der angegebenen Phase werden "
            "pauschal als erledigt markiert, die Aufgaben der Phase bleiben offen.",
            "Kunden werden über Nachname + Vorname + PLZ wiedererkannt (wie beim "
            "monday-Sync), sonst neu angelegt.",
            "Schlüssel ist TAIFUN-Nummer + Sparte: ein zweiter Import derselben "
            "Datei aktualisiert die Einträge statt sie zu verdoppeln.",
            "Vor dem echten Import immer die Vorschau prüfen; Fehlerzeilen werden "
            "nicht importiert."):
        anleitung.append([zeile])
    anleitung.column_dimensions["A"].width = 40
    anleitung.column_dimensions["C"].width = 90
    puffer = io.BytesIO()
    wb.save(puffer)
    return puffer.getvalue()


# --- Einlesen + Prüfen -----------------------------------------------------------

@dataclass
class Zeile:
    nr: int                                  # Excel-Zeilennummer
    werte: dict
    fehler: list[str] = field(default_factory=list)
    hinweise: list[str] = field(default_factory=list)
    aktion: str = ""                         # neu | aktualisieren
    kunde_id: int | None = None
    phase: str = ""
    datum: datetime | None = None
    montage_von: datetime | None = None
    montage_bis: datetime | None = None
    betrag_cent: int = 0
    ids: dict = field(default_factory=dict)  # Team-/Benutzer-IDs

    @property
    def ok(self) -> bool:
        return not self.fehler


def _text(wert) -> str:
    if wert is None:
        return ""
    if isinstance(wert, float) and wert == int(wert):
        wert = int(wert)
    if isinstance(wert, (datetime, date)):
        return wert.strftime("%d.%m.%Y")
    return str(wert).strip()


def _datum(wert) -> datetime | None:
    if isinstance(wert, datetime):
        return wert
    if isinstance(wert, date):
        return datetime.combine(wert, datetime.min.time())
    text = _text(wert)
    for muster in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, muster)
        except ValueError:
            continue
    return None


def _betrag_cent(wert) -> int | None:
    if isinstance(wert, (int, float)):
        return int(round(float(wert) * 100))
    from app.routers.artikel import preis_parsen
    return preis_parsen(_text(wert).replace("€", "").strip())


def _normal(text: str) -> str:
    from app.monday_sync import _normal as monday_normal
    return monday_normal(text or "")


def _phase(wert: str) -> str:
    text = (wert or "").strip().lower()
    for key in PHASEN_AUSWAHL:
        if text in (key, GEWERK_PHASEN_NAMEN[key].lower()):
            return key
    return ""


def einlesen(inhalt: bytes) -> tuple[list[Zeile], str]:
    """Excel → Zeilen (roh). Liefert (zeilen, fehlertext)."""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(inhalt), data_only=True, read_only=True)
    except Exception as problem:
        return [], f"Datei ist keine lesbare Excel-Datei ({problem})."
    if BLATT not in wb.sheetnames:
        return [], f"Blatt „{BLATT}“ fehlt – bitte die Vorlage verwenden."
    ws = wb[BLATT]
    zeilen_roh = list(ws.iter_rows(values_only=True))
    if not zeilen_roh:
        return [], "Das Blatt ist leer."
    kopf = [_text(z).lower() for z in zeilen_roh[0]]
    index = {}
    for schluessel, titel, _, _ in SPALTEN:
        if titel.lower() in kopf:
            index[schluessel] = kopf.index(titel.lower())
    fehlend = [t for s, t, p, _ in SPALTEN if p and s not in index]
    if fehlend:
        return [], "Pflichtspalten fehlen: " + ", ".join(fehlend)
    zeilen = []
    for nr, roh in enumerate(zeilen_roh[1:], start=2):
        if roh is None or not any(_text(z) for z in roh):
            continue
        werte = {s: (roh[i] if i < len(roh) else None) for s, i in index.items()}
        if "BEISPIELZEILE" in _text(werte.get("bemerkung")).upper():
            continue
        zeilen.append(Zeile(nr=nr, werte=werte))
    return zeilen, ""


def _benutzer_nach_name(session: Session) -> dict[str, int]:
    return {_normal(b.name): b.id for b in session.query(Benutzer)}


def _teams_nach_name(session: Session) -> dict[str, int]:
    return {_normal(t.name): t.id for t in session.query(Team).filter(Team.aktiv.is_(True))}


def _bestandsangebot(session: Session, taifun_nummer: str, sparte: str) -> Angebot | None:
    return (session.query(Angebot)
            .filter(Angebot.extern.is_(True), Angebot.taifun_nummer == taifun_nummer,
                    Angebot.konfigurator_typ == sparte)
            .order_by(Angebot.id).first())


def pruefen(session: Session, zeilen: list[Zeile]) -> list[Zeile]:
    """Prüfung je Zeile (Vorschau): Pflichtfelder, Kunde vorhanden?, Team/
    Projektleiter bekannt?, Phase gültig?, Datum plausibel?, Doppel in der Datei."""
    benutzer = _benutzer_nach_name(session)
    teams = _teams_nach_name(session)
    standard_pl = kern.parameter_holen(session, "standard_projektleiter", "")
    gesehen: dict[tuple[str, str], int] = {}
    heute = datetime.now()
    for z in zeilen:
        w = {k: _text(v) for k, v in z.werte.items()}
        for schluessel, titel, pflicht, _ in SPALTEN:
            if pflicht and not w.get(schluessel):
                z.fehler.append(f"{titel} fehlt")
        sparte = w.get("sparte", "").upper()
        if sparte and sparte not in ("WP", "PV", "KL", "WB"):
            z.fehler.append(f"Sparte „{w['sparte']}“ ungültig (WP/PV/KL/WB)")
        if w.get("plz") and not (w["plz"].isdigit() and len(w["plz"]) == 5):
            z.fehler.append(f"PLZ „{w['plz']}“ ungültig")
        if w.get("phase"):
            z.phase = _phase(w["phase"])
            if not z.phase:
                roh = w["phase"].lower()
                z.fehler.append(
                    f"Phase „{w['phase']}“ ungültig"
                    + (" – seit V4 getrennt: Abnahme oder Freigabe"
                       if "abnahme" in roh and "freigabe" in roh else ""))
        if w.get("auftragswert"):
            betrag = _betrag_cent(z.werte.get("auftragswert"))
            if betrag is None or betrag <= 0:
                z.fehler.append(f"Auftragswert „{w['auftragswert']}“ ungültig")
            else:
                z.betrag_cent = betrag
        if w.get("auftragsdatum"):
            z.datum = _datum(z.werte.get("auftragsdatum"))
            if z.datum is None:
                z.fehler.append(f"Auftragsdatum „{w['auftragsdatum']}“ ungültig")
            elif z.datum > heute or z.datum.year < 2015:
                z.fehler.append(f"Auftragsdatum {w['auftragsdatum']} unplausibel")
        for schluessel in ("montage_von", "montage_bis"):
            if w.get(schluessel):
                wert = _datum(z.werte.get(schluessel))
                if wert is None:
                    z.fehler.append(f"{schluessel.replace('_', ' ').title()} "
                                    f"„{w[schluessel]}“ ungültig")
                setattr(z, schluessel, wert)
        if z.montage_bis and not z.montage_von:
            z.fehler.append("Montage bis ohne Montage von")
        if z.montage_von and z.montage_bis and z.montage_bis < z.montage_von:
            z.fehler.append("Montage bis liegt vor Montage von")
        if z.montage_von and z.datum and z.montage_von < z.datum:
            z.hinweise.append("Montagebeginn liegt vor dem Auftragsdatum")
        if z.phase in ("montage", "abnahme", "freigabe") and not z.montage_von:
            z.hinweise.append("Phase ab Montage, aber kein Montagetermin")
        if w.get("bestaetigt") and w["bestaetigt"].upper() not in ("J", "N", "JA", "NEIN"):
            z.fehler.append("Termin bestätigt: nur J oder N")
        if w.get("oeltank") and w["oeltank"].upper() not in ("J", "N", "JA", "NEIN"):
            z.fehler.append("Öltankentsorgung: nur J oder N")
        # Personen und Teams
        if w.get("projektleiter"):
            pl = benutzer.get(_normal(w["projektleiter"]))
            if pl is None:
                z.fehler.append(f"Projektleiter „{w['projektleiter']}“ unbekannt")
            z.ids["projektleiter"] = pl
        elif standard_pl.isdigit():
            z.ids["projektleiter"] = int(standard_pl)
            z.hinweise.append("ohne Projektleiter → Standard-Projektleiter")
        else:
            z.fehler.append("Projektleiter fehlt (kein Standard-Projektleiter gesetzt)")
        if w.get("vertriebler"):
            z.ids["vertriebler"] = benutzer.get(_normal(w["vertriebler"]))
            if z.ids["vertriebler"] is None:
                z.hinweise.append(f"Vertriebler „{w['vertriebler']}“ unbekannt – bleibt leer")
        for schluessel, name in (("montageteam", "Montageteam"),
                                 ("elektroteam", "Elektro-Team"), ("subteam", "Subteam")):
            if w.get(schluessel):
                team_id = teams.get(_normal(w[schluessel]))
                if team_id is None:
                    z.fehler.append(f"{name} „{w[schluessel]}“ unbekannt")
                z.ids[schluessel] = team_id
        if (w.get("elektroteam") or w.get("subteam")) and not z.montage_von:
            z.hinweise.append("Teams ohne Montagetermin werden nur zugeordnet")
        # Doppel in der Datei
        schluessel = (w.get("taifun_nummer", ""), sparte)
        if all(schluessel):
            if schluessel in gesehen:
                z.fehler.append(f"doppelt in der Datei (wie Zeile {gesehen[schluessel]})")
            else:
                gesehen[schluessel] = z.nr
            z.aktion = ("aktualisieren" if _bestandsangebot(session, *schluessel)
                        is not None else "neu")
        # Kunde vorhanden? (Name + PLZ wie beim monday-Sync)
        kunde = _kunde_finden(session, w)
        z.kunde_id = kunde.id if kunde is not None else None
    return zeilen


def _kunde_finden(session: Session, w: dict) -> Kunde | None:
    """Duplikatabgleich Nachname + Vorname + PLZ (wie beim monday-Sync)."""
    if not (w.get("nachname") and w.get("plz")):
        return None
    for kandidat in session.query(Kunde).filter(Kunde.plz == w["plz"]):
        if (_normal(kandidat.nachname) == _normal(w["nachname"])
                and _normal(kandidat.vorname) == _normal(w.get("vorname", ""))):
            return kandidat
    return None


# --- Import -----------------------------------------------------------------------

def _ausfuehrung(text: str, kunde: Kunde) -> dict:
    """„Straße Nr, PLZ Ort“ → Adresse; leer = Kundenadresse."""
    text = (text or "").strip()
    if not text:
        return {"strasse": kunde.strasse, "plz": kunde.plz, "ort": kunde.ort}
    strasse, _, rest = text.partition(",")
    rest = rest.strip()
    plz, _, ort = rest.partition(" ")
    if not plz.isdigit():
        plz, ort = "", rest
    return {"strasse": strasse.strip()[:200], "plz": plz[:10], "ort": ort.strip()[:100]}


def _pflicht_vorher_erledigen(session: Session, gewerk: Gewerk, benutzer) -> int:
    """Alle Pflichtaufgaben der Pakete VOR der aktuellen Phase → erledigt
    (pauschal, Verlaufseintrag); die Aufgaben der aktuellen Phase bleiben offen."""
    rang = kern._PHASEN_RANG.get(gewerk.phase, 0)
    instanzen = {i.id: i for i in session.query(AufgabenpaketInstanz)
                 .filter(AufgabenpaketInstanz.gewerk_id == gewerk.id)}
    anzahl = 0
    jetzt = datetime.now()
    for aufgabe in session.query(Aufgabe).filter(Aufgabe.gewerk_id == gewerk.id):
        instanz = instanzen.get(aufgabe.paket_instanz_id)
        if instanz is None or not aufgabe.pflicht:
            continue
        if aufgabe.status in ("erledigt", "entfaellt"):
            continue
        if kern.paket_rang(instanz.paket_name) < rang:
            aufgabe.status = "erledigt"
            aufgabe.erledigt_am = jetzt
            aufgabe.erledigt_von = benutzer.id if benutzer else None
            anzahl += 1
    if anzahl:
        kern.verlauf(session, gewerk.projekt_id,
                     f"Bestand – pauschal erledigt: {anzahl} Pflichtaufgaben vor "
                     f"„{GEWERK_PHASEN_NAMEN.get(gewerk.phase, gewerk.phase)}“",
                     benutzer=benutzer, gewerk_id=gewerk.id)
    return anzahl


def _zeile_importieren(session: Session, z: Zeile, imp: Bestandsimport,
                       benutzer) -> dict:
    """Legt eine Zeile an bzw. aktualisiert sie; liefert den Protokolleintrag."""
    w = {k: _text(v) for k, v in z.werte.items()}
    sparte = w["sparte"].upper()
    eintrag = {"zeile": z.nr, "taifun_nummer": w["taifun_nummer"], "sparte": sparte,
               "neu": {}, "aktion": ""}
    angebot = _bestandsangebot(session, w["taifun_nummer"], sparte)
    # Kunde
    # erneut prüfen: eine frühere Zeile derselben Datei kann den Kunden eben
    # erst angelegt haben (zwei Sparten beim selben neuen Kunden)
    kunde = session.get(Kunde, z.kunde_id) if z.kunde_id else _kunde_finden(session, w)
    if angebot is not None:
        kunde = session.get(Kunde, angebot.kunde_id) or kunde
    if kunde is None:
        kunde = Kunde(anrede=w.get("anrede", ""), vorname=w.get("vorname", ""),
                      nachname=w["nachname"], strasse=w.get("strasse", ""),
                      plz=w["plz"], ort=w.get("ort", ""), telefon=w.get("telefon", ""),
                      email=w.get("email", ""), interesse=sparte,
                      vertriebskanal=w.get("kanal", ""))
        session.add(kunde)
        session.flush()
        eintrag["neu"]["kunde_id"] = kunde.id
    else:
        for feld in ("strasse", "ort", "telefon", "email"):
            if w.get(feld) and not getattr(kunde, feld):
                setattr(kunde, feld, w[feld])
    # externer Angebotseintrag (Status Angenommen, Badge TAIFUN, bestand)
    if angebot is None:
        angebot = Angebot(nummer="EXT-tmp", kunde_id=kunde.id, extern=True,
                          taifun_nummer=w["taifun_nummer"][:30],
                          extern_endbetrag_cent=z.betrag_cent, datum=z.datum,
                          vertriebler_id=z.ids.get("vertriebler"),
                          konfigurator_typ=sparte, bestand=True)
        session.add(angebot)
        session.flush()
        angebot.nummer = f"EXT-{angebot.id}"
        angebot_status_setzen(angebot, "Angenommen")
        angebot.versendet_am = z.datum
        angebot.angenommen_am = z.datum
        eintrag["neu"]["angebot_id"] = angebot.id
        from app import vorgaenge as vorgaenge_modul
        vorgaenge_vorher = {v.id for v in session.query(Vorgang)
                            .filter(Vorgang.kunde_id == kunde.id)}
        vorgang = vorgaenge_modul.vorgang_fuer_angebot(session, angebot)
        session.flush()
        if vorgang.id not in vorgaenge_vorher:
            eintrag["neu"]["vorgang_id"] = vorgang.id
        eintrag["aktion"] = "angelegt"
    else:
        angebot.extern_endbetrag_cent = z.betrag_cent
        angebot.datum = z.datum
        angebot.angenommen_am = z.datum
        if z.ids.get("vertriebler"):
            angebot.vertriebler_id = z.ids["vertriebler"]
        eintrag["aktion"] = "aktualisiert"
    eintrag["angebot_id"] = angebot.id
    # Projekt + Gewerk
    gewerk = session.get(Gewerk, angebot.projekt_gewerk_id) if angebot.projekt_gewerk_id else None
    if gewerk is None:
        from app import vorgaenge as vorgaenge_modul
        vorgang = vorgaenge_modul.vorgang_fuer_angebot(session, angebot)
        projekt = kern.offenes_projekt_fuer_vorgang(session, vorgang.id)
        # ein weiterer Auftrag derselben Sparte (andere TAIFUN-Nummer) bekommt
        # ein eigenes Projekt – je Projekt höchstens ein offenes Gewerk je Sparte
        if projekt is not None and session.query(Gewerk).filter(
                Gewerk.projekt_id == projekt.id, Gewerk.sparte == sparte,
                Gewerk.phase != "storniert").count():
            projekt = None
        if projekt is None:
            projekt = kern.projekt_anlegen(
                session, angebot, benutzer=benutzer,
                projektleiter_id=z.ids.get("projektleiter"),
                notiz_kopf=w.get("bemerkung", ""),
                adresse=_ausfuehrung(w.get("ausfuehrung", ""), kunde))
            if w.get("kanal"):
                projekt.kanal = w["kanal"][:100]
            eintrag["neu"]["projekt_id"] = projekt.id
            kern.verlauf(session, projekt.id, f"Projekt {projekt.nummer} angelegt "
                         "(Bestandsimport)", benutzer=benutzer)
        gewerk = kern.gewerk_anlegen(session, projekt, angebot, sparte,
                                     benutzer=benutzer, quelle="bestand",
                                     benachrichtigen_an=False)
        eintrag["neu"]["gewerk_id"] = gewerk.id
        gewerk.phase = z.phase
        gewerk.phase_geaendert_am = datetime.now()
        kern.verlauf(session, projekt.id,
                     f"Aus Bestandsimport angelegt ({imp.dateiname}, Zeile {z.nr})",
                     benutzer=benutzer, gewerk_id=gewerk.id)
    else:
        projekt = session.get(Projekt, gewerk.projekt_id)
        gewerk.auftragswert_aktuell = z.betrag_cent
        if gewerk.phase != z.phase and gewerk.phase not in ("abgeschlossen", "storniert"):
            kern.verlauf(session, projekt.id,
                         f"Bestandsimport: Phase {GEWERK_PHASEN_NAMEN.get(gewerk.phase)} → "
                         f"{GEWERK_PHASEN_NAMEN.get(z.phase)} ({imp.dateiname}, Zeile {z.nr})",
                         benutzer=benutzer, gewerk_id=gewerk.id)
            gewerk.phase = z.phase
            gewerk.phase_geaendert_am = datetime.now()
        if z.ids.get("projektleiter"):
            projekt.projektleiter_id = z.ids["projektleiter"]
    gewerk.bestand_import_id = imp.id
    eintrag["gewerk_id"] = gewerk.id
    eintrag["phase"] = gewerk.phase
    # Steckbrief aus den Spalten (quelle = auftragsdaten, FP darf überschreiben)
    werte = {}
    for spalte in STECKBRIEF_SPALTEN:
        if w.get(spalte):
            werte[spalte] = ({"J": "ja", "JA": "ja", "N": "nein", "NEIN": "nein"}
                             .get(w[spalte].upper(), w[spalte])
                             if spalte == "oeltank" else w[spalte])
    if werte:
        kern.auftragsdaten_speichern(session, gewerk, werte, benutzer=benutzer)
    gewerk.auftragsdaten_am = gewerk.auftragsdaten_am or datetime.now()
    # Pflichtaufgaben vor der Phase pauschal erledigen
    eintrag["pauschal_erledigt"] = _pflicht_vorher_erledigen(session, gewerk, benutzer)
    # Montagetermin mit Team und Bestätigung (nur wenn noch keiner existiert)
    if z.montage_von:
        vorhanden = (session.query(ProjektTermin)
                     .filter(ProjektTermin.gewerk_id == gewerk.id,
                             ProjektTermin.typ == "montage").first())
        bestaetigt = w.get("bestaetigt", "").upper() in ("J", "JA")
        if vorhanden is None:
            team_id = (z.ids.get("montageteam") or z.ids.get("elektroteam")
                       or z.ids.get("subteam"))
            zweck = ("wp" if z.ids.get("montageteam") else
                     "elektro" if z.ids.get("elektroteam") else "sub")
            if team_id:
                kern.team_termin_zuweisen(session, gewerk, zweck, team_id,
                                          z.montage_von, z.montage_bis, bestaetigt,
                                          benutzer=benutzer)
            else:
                session.add(ProjektTermin(
                    projekt_id=gewerk.projekt_id, gewerk_id=gewerk.id, typ="montage",
                    beginn=z.montage_von,
                    ende=z.montage_bis or kern.arbeitstage_addieren(z.montage_von, 4),
                    ganztaegig=True, kunde_bestaetigt=bestaetigt,
                    bestaetigt_am=datetime.now() if bestaetigt else None,
                    bestaetigt_quelle="import" if bestaetigt else "",
                    erstellt_von=benutzer.id if benutzer else None))
            session.flush()
            eintrag["neu"]["termin_ids"] = [t.id for t in session.query(ProjektTermin)
                                            .filter(ProjektTermin.gewerk_id == gewerk.id)]
        else:
            vorhanden.beginn = z.montage_von
            vorhanden.ende = z.montage_bis or kern.arbeitstage_addieren(z.montage_von, 4)
            vorhanden.kunde_bestaetigt = bestaetigt
    # übrige Teams nur zuordnen
    for schluessel, feld in (("montageteam", "wp_team_id"),
                             ("elektroteam", "elektro_team_id"),
                             ("subteam", "sub_team_id")):
        if z.ids.get(schluessel):
            setattr(gewerk, feld, z.ids[schluessel])
    kern.projektstatus_berechnen(session, projekt)
    session.flush()
    return eintrag


def importieren(session: Session, zeilen: list[Zeile], dateiname: str,
                benutzer=None) -> Bestandsimport:
    """Importiert NUR fehlerfreie Zeilen (Vorschau vorher). Protokoll am Import."""
    imp = Bestandsimport(dateiname=(dateiname or "")[:200],
                         benutzer_id=benutzer.id if benutzer else None)
    session.add(imp)
    session.flush()
    protokoll = []
    angelegt = aktualisiert = uebersprungen = 0
    for z in zeilen:
        if not z.ok:
            uebersprungen += 1
            protokoll.append({"zeile": z.nr, "aktion": "übersprungen",
                              "fehler": z.fehler})
            continue
        eintrag = _zeile_importieren(session, z, imp, benutzer)
        protokoll.append(eintrag)
        if eintrag["aktion"] == "angelegt":
            angelegt += 1
        else:
            aktualisiert += 1
    imp.angelegt, imp.aktualisiert, imp.uebersprungen = angelegt, aktualisiert, uebersprungen
    imp.protokoll_json = json.dumps(protokoll, ensure_ascii=False, default=str)
    imp.abgeschlossen_am = datetime.now()
    session.flush()
    return imp


# --- Rückgängig ---------------------------------------------------------------------

def _gewerk_unveraendert(session: Session, gewerk: Gewerk, imp: Bestandsimport,
                         phase: str) -> bool:
    """Seit dem Import nichts geändert: gleiche Phase, keine später erledigten
    Aufgaben, keine späteren Verlaufseinträge/Dokumente am Gewerk."""
    grenze = imp.abgeschlossen_am or imp.erstellt_am
    if gewerk.phase != phase or gewerk.bestand_import_id != imp.id:
        return False
    if (session.query(Aufgabe).filter(Aufgabe.gewerk_id == gewerk.id,
                                      Aufgabe.erledigt_am > grenze).count()):
        return False
    if (session.query(ProjektVerlauf).filter(ProjektVerlauf.gewerk_id == gewerk.id,
                                             ProjektVerlauf.erstellt_am > grenze).count()):
        return False
    if (session.query(ProjektDokument).filter(ProjektDokument.gewerk_id == gewerk.id,
                                              ProjektDokument.quelle != "auto").count()):
        return False
    return True


def rueckgaengig(session: Session, imp: Bestandsimport) -> tuple[int, list[str]]:
    """Löscht nur, was dieser Import ANGELEGT hat und seither unverändert ist.
    Liefert (Anzahl entfernter Gewerke, Hinweise zu behaltenen Zeilen)."""
    from app.models import GalerieDatei, Erfassung
    protokoll = json.loads(imp.protokoll_json or "[]")
    entfernt = 0
    behalten: list[str] = []
    for eintrag in reversed(protokoll):
        neu = eintrag.get("neu") or {}
        if not neu.get("gewerk_id"):
            if eintrag.get("aktion") == "aktualisiert":
                behalten.append(f"Zeile {eintrag['zeile']}: nur aktualisiert – bleibt")
            continue
        gewerk = session.get(Gewerk, neu["gewerk_id"])
        if gewerk is None:
            continue
        if not _gewerk_unveraendert(session, gewerk, imp, eintrag.get("phase", "")):
            behalten.append(f"Zeile {eintrag['zeile']}: seit dem Import geändert – bleibt")
            continue
        projekt_id = gewerk.projekt_id
        session.query(Aufgabe).filter(Aufgabe.gewerk_id == gewerk.id).delete()
        session.query(AufgabenpaketInstanz).filter(
            AufgabenpaketInstanz.gewerk_id == gewerk.id).delete()
        session.query(SteckbriefWert).filter(SteckbriefWert.gewerk_id == gewerk.id).delete()
        session.query(ProjektTermin).filter(ProjektTermin.gewerk_id == gewerk.id).delete()
        session.query(ProjektVerlauf).filter(ProjektVerlauf.gewerk_id == gewerk.id).delete()
        session.query(ProjektDokument).filter(ProjektDokument.gewerk_id == gewerk.id).delete()
        session.delete(gewerk)
        session.flush()
        entfernt += 1
        if neu.get("projekt_id") and not session.query(Gewerk).filter(
                Gewerk.projekt_id == projekt_id).count():
            session.query(ProjektVerlauf).filter(
                ProjektVerlauf.projekt_id == projekt_id).delete()
            session.query(ProjektTermin).filter(
                ProjektTermin.projekt_id == projekt_id).delete()
            session.query(ProjektDokument).filter(
                ProjektDokument.projekt_id == projekt_id).delete()
            projekt = session.get(Projekt, projekt_id)
            if projekt is not None:
                session.delete(projekt)
        if neu.get("angebot_id"):
            angebot = session.get(Angebot, neu["angebot_id"])
            if angebot is not None:
                session.delete(angebot)
        session.flush()
        if neu.get("vorgang_id"):
            vorgang = session.get(Vorgang, neu["vorgang_id"])
            if vorgang is not None and not session.query(Angebot).filter(
                    Angebot.vorgang_id == vorgang.id).count() and not session.query(
                    Projekt).filter(Projekt.vorgang_id == vorgang.id).count():
                from app.models import VorgangsNotiz
                session.query(VorgangsNotiz).filter(
                    VorgangsNotiz.vorgang_id == vorgang.id).delete()
                session.query(GalerieDatei).filter(
                    GalerieDatei.vorgang_id == vorgang.id).delete()
                session.delete(vorgang)
        session.flush()
        if neu.get("kunde_id"):
            kunde = session.get(Kunde, neu["kunde_id"])
            if kunde is not None and not any((
                    session.query(Angebot).filter(Angebot.kunde_id == kunde.id).count(),
                    session.query(Erfassung).filter(Erfassung.kunde_id == kunde.id).count(),
                    session.query(Vorgang).filter(Vorgang.kunde_id == kunde.id).count())):
                session.delete(kunde)
    imp.rueckgaengig_am = datetime.now()
    imp.rueckgaengig_hinweis = "\n".join(behalten)[:2000]
    session.flush()
    return entfernt, behalten
