# V4 (PLAN_PROJ_V4 Phase 93.2): Stücklisten-Pflege in der Oberfläche.
# Die Logik-Excel (Blatt "Stücklisten") bleibt Master – die UI schreibt
# zurück (Backup der alten Datei in config.BACKUP_ORDNER, danach Cache neu
# laden). CSV-Export/-Import für Massenpflege.
#
# Hinweis Server: projektierung_logik_v1.xlsx ist git-verfolgt; ein Rück-
# schreiben auf dem Server macht die Arbeitskopie "dirty" – update.bat
# (git pull --ff-only) kann dann blockieren, genau wie beim bestehenden
# Excel-Upload (docs/projektierung-entscheidungen.md, Phase 93).

import csv
import io
import shutil
from datetime import datetime

from app import config, projektierung_logik

BLATT = "Stücklisten"
KOPF = ["position", "lieferant_artnr", "menge_je_einheit", "bezeichnung",
        "lieferant", "mengeneinheit"]
GO_LIVE_QUOTE = 0.9

# v28 (PLAN_PROJ_V6 Phase 139): Konvention Spalte „lieferant“ – Standard-
# Lieferant (Parameter stueckliste_standard_lieferant, „Collin“) = bestellen ·
# Lager/LAGER = Lagerware · „–“/„-“/KEIN-MATERIAL = Leistung ohne Material ·
# anderer Name = Fremdlieferant. Nur „bestellen“ landet in der UGL.
ART_NAMEN = {"bestellen": "bestellen", "lager": "Lager", "leistung": "Leistung",
             "fremd": "Fremd"}
LAGER_WERTE = ("lager",)
LEISTUNG_WERTE = ("–", "-", "—", "kein-material", "kein material", "keinmaterial")


def standard_lieferant(session) -> str:
    from app import projektierung
    return (projektierung.parameter_holen(session, "stueckliste_standard_lieferant",
                                          "Collin") or "Collin").strip()


def art_fuer_lieferant(lieferant: str, standard: str = "Collin") -> str:
    """bestellen | lager | leistung | fremd nach der Konvention (Groß/Klein egal)."""
    wert = (lieferant or "").strip()
    klein = wert.lower()
    if not wert or klein == (standard or "Collin").strip().lower():
        return "bestellen"
    if klein in LAGER_WERTE:
        return "lager"
    if klein in LEISTUNG_WERTE:
        return "leistung"
    return "fremd"


def positionen(session, nur_ohne: bool = False, suche: str = "") -> list[dict]:
    """Angebotspositionen aus dem Artikelstamm (aktiv, mit Positionsnummer,
    ohne Z-Arbeitspakete) samt Stücklisten-Zeilen. v28: je Position
    `arten` = [(zeile, art)] nach der Lieferanten-Konvention; „ohne Zuordnung“
    = keine Zeile (auch Lager/Leistung zählen als zugeordnet)."""
    from app.models import Artikel
    logik = projektierung_logik.hole_logik(session)
    standard = standard_lieferant(session)
    ergebnis: dict[str, dict] = {}
    for artikel in (session.query(Artikel)
                    .filter(Artikel.aktiv.is_(True), Artikel.pos_nr != "")
                    .order_by(Artikel.pos_nr)):
        nr = artikel.pos_nr.strip()
        if not nr or nr.upper().startswith("Z") or nr in ergebnis:
            continue
        ergebnis[nr] = {"pos_nr": nr, "titel": artikel.titel,
                        "kategorie": artikel.kategorie, "einheit": artikel.einheit,
                        "zeilen": logik.stuecklisten.get(nr, [])}
    # Stücklisten zu Positionen, die (noch) nicht im Artikelstamm stehen
    for nr, zeilen in logik.stuecklisten.items():
        ergebnis.setdefault(nr, {"pos_nr": nr, "titel": "(nicht im Artikelstamm)",
                                 "kategorie": "", "einheit": "", "zeilen": zeilen})
    for p in ergebnis.values():
        p["arten"] = [(z, art_fuer_lieferant(z.lieferant, standard)) for z in p["zeilen"]]
        p["bestellbar"] = any(a == "bestellen" for _, a in p["arten"])
    liste = sorted(ergebnis.values(), key=lambda p: p["pos_nr"])
    if nur_ohne:
        liste = [p for p in liste if not p["zeilen"]]
    if suche:
        s = suche.lower()
        liste = [p for p in liste if s in p["pos_nr"].lower()
                 or s in (p["titel"] or "").lower()
                 or s in (p["kategorie"] or "").lower()
                 or any(s in z.lieferant_artnr.lower() for z in p["zeilen"])]
    return liste


def _zaehlt_fuer_quote(p: dict) -> bool:
    """v28 (Phase 139) [ANNAHME]: Go-live-Quote nur über WP-Positionen und
    spartenübergreifende – PV…/KL…-Positionen bleiben außen vor, solange sie
    keine Stücklisten-Zeile (Lieferant) haben."""
    if p["titel"] == "(nicht im Artikelstamm)":
        return False
    nr = (p["pos_nr"] or "").upper()
    if nr.startswith(("PV", "KL")) and not p["zeilen"]:
        return False
    return True


def fortschritt(session) -> tuple[int, int]:
    """(zugeordnet, gesamt) – eine Position zählt als zugeordnet, sobald sie
    mindestens eine Zeile hat (auch Lager/Leistung); Nenner laut
    _zaehlt_fuer_quote (v28)."""
    alle = [p for p in positionen(session) if _zaehlt_fuer_quote(p)]
    return sum(1 for p in alle if p["zeilen"]), len(alle)


def _zeilen_aus_datei(ws) -> list[list]:
    zeilen = []
    for zeile in ws.iter_rows(min_row=2, values_only=True):
        werte = list((tuple(zeile) + (None,) * 6)[:6])
        if not werte[0] and not werte[1]:
            continue
        zeilen.append(werte)
    return zeilen


def _schreiben(alle_zeilen: list[list]) -> str:
    """Blatt komplett neu schreiben (Kopf + Zeilen, sortiert nach Position),
    Backup vorher; gibt den Backup-Dateinamen zurück."""
    import openpyxl
    pfad = projektierung_logik.LOGIK_PFAD
    config.BACKUP_ORDNER.mkdir(parents=True, exist_ok=True)
    sicherung = config.BACKUP_ORDNER / f"projektierung_logik_{datetime.now():%Y%m%d_%H%M%S_%f}.xlsx"
    shutil.copyfile(pfad, sicherung)
    wb = openpyxl.load_workbook(pfad)
    if BLATT in wb.sheetnames:
        ws = wb[BLATT]
        ws.delete_rows(1, ws.max_row)
    else:
        ws = wb.create_sheet(BLATT)
    ws.append(KOPF)
    for werte in sorted(alle_zeilen, key=lambda z: (str(z[0]), str(z[1]))):
        ws.append(werte)
    wb.save(pfad)
    return sicherung.name


def _neu_laden(session) -> list[str]:
    logik = projektierung_logik.hole_logik(session, erzwingen=True)
    return logik.fehler


def _zeile(pos_nr, artnr, menge, bezeichnung, lieferant, einheit) -> list | None:
    artnr = (artnr or "").strip()
    if not artnr:
        return None
    try:
        menge_zahl = float(str(menge or "1").replace(",", "."))
    except ValueError:
        menge_zahl = 1.0
    return [pos_nr, artnr, menge_zahl, (bezeichnung or "").strip(),
            (lieferant or "").strip() or "Collin",
            ((einheit or "").strip() or "ST").upper()[:3]]


def position_speichern(session, pos_nr: str, zeilen: list[dict]) -> str:
    """Zeilen einer Position ersetzen (leere Artikelnummern fallen weg)."""
    import openpyxl
    pos_nr = pos_nr.strip()
    wb = openpyxl.load_workbook(projektierung_logik.LOGIK_PFAD, read_only=True)
    bestand = _zeilen_aus_datei(wb[BLATT]) if BLATT in wb.sheetnames else []
    wb.close()
    neu = [z for z in bestand if str(z[0]).strip() != pos_nr]
    for eingabe in zeilen:
        zeile = _zeile(pos_nr, eingabe.get("artnr"), eingabe.get("menge"),
                       eingabe.get("bezeichnung"), eingabe.get("lieferant"),
                       eingabe.get("einheit"))
        if zeile:
            neu.append(zeile)
    sicherung = _schreiben(neu)
    fehler = _neu_laden(session)
    return (f"Stückliste {pos_nr} gespeichert (Backup {sicherung})"
            + (" – Fehler: " + " · ".join(fehler) if fehler else "."))


def csv_export(session) -> str:
    puffer = io.StringIO()
    schreiber = csv.writer(puffer, delimiter=";")
    schreiber.writerow(KOPF)
    logik = projektierung_logik.hole_logik(session)
    for nr in sorted(logik.stuecklisten):
        for z in logik.stuecklisten[nr]:
            schreiber.writerow([nr, z.lieferant_artnr,
                                f"{z.menge_je_einheit:g}".replace(".", ","),
                                z.bezeichnung, z.lieferant, z.mengeneinheit])
    return "﻿" + puffer.getvalue()


def csv_import(session, inhalt: bytes) -> str:
    """Ersetzt das komplette Blatt durch die CSV (Kopfzeile wie Export)."""
    text = inhalt.decode("utf-8-sig", errors="replace")
    trenner = ";" if text.count(";") >= text.count(",") else ","
    leser = csv.reader(io.StringIO(text), delimiter=trenner)
    kopf = [k.strip().lower() for k in next(leser, [])]
    if kopf[:2] != KOPF[:2]:
        return "Import abgewiesen: Kopfzeile muss mit position;lieferant_artnr beginnen."
    neu = []
    for werte in leser:
        werte = (werte + [""] * 6)[:6]
        if not werte[0].strip():
            continue
        zeile = _zeile(werte[0].strip(), *werte[1:6])
        if zeile:
            neu.append(zeile)
    sicherung = _schreiben(neu)
    fehler = _neu_laden(session)
    return (f"{len(neu)} Stücklisten-Zeilen importiert (Backup {sicherung})"
            + (" – Fehler: " + " · ".join(fehler) if fehler else "."))
