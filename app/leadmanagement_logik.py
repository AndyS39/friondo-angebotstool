# Lead-Steuerdatei (v12, Phase 74): liest leadmanagement_logik_v1.xlsx
# (Blätter Qualifizierung, Scoring, Klassen, Kaskade, Gruende, Wunschzeiten).
# Muster wie app/projektierung_logik.py: Datei-basiert, mtime-Cache,
# Fehler/Warnungen statt Abbruch. Änderungen wirken auf NEUE Qualifizierungen.

import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from app import config

LOGIK_PFAD = Path(getattr(config, "LEADMANAGEMENT_LOGIK_PFAD",
                          Path(config.LOGIK_EXCEL_PFAD).parent
                          / "leadmanagement_logik_v1.xlsx"))

SPARTEN = ("WP", "PV", "KL", "WB")


@dataclass
class Frage:
    sparte: str
    key: str
    frage: str
    typ: str                      # janein | auswahl | mehrfach | zahl | text
    optionen: list[str]
    pflicht: bool
    reihenfolge: int
    erfassungs_frage: str = ""
    unqualifiziert_bei: str = ""


@dataclass
class ScoringRegel:
    frage_key: str
    bedingung: str                # =Wert, >N, >=N, <N, <=N, enthält Wert
    punkte: int
    bemerkung: str = ""


@dataclass
class KaskadenStufe:
    versuch_nr: int
    wiedervorlage_nach: str       # +2h, +1d 18:00, +3d, +7d
    aktion: str                   # keine | mail_nicht_erreicht
    letzter: bool


@dataclass
class Grund:
    phase: str                    # unqualifiziert | zurueckgestellt | no_show | verloren_vor_termin
    grund: str
    freitext_pflicht: bool


@dataclass
class Wunschzeit:
    key: str
    bezeichnung: str
    von: str
    bis: str
    wochentage: str               # mo-fr | sa


@dataclass
class LeadLogik:
    fragen: list[Frage] = field(default_factory=list)
    scoring: list[ScoringRegel] = field(default_factory=list)
    klassen: list[tuple[str, int]] = field(default_factory=list)   # sortiert absteigend
    kaskade: list[KaskadenStufe] = field(default_factory=list)
    gruende: list[Grund] = field(default_factory=list)
    wunschzeiten: list[Wunschzeit] = field(default_factory=list)
    fehler: list[str] = field(default_factory=list)
    warnungen: list[str] = field(default_factory=list)
    stand: str = ""

    def fragen_der_sparte(self, sparte: str) -> list[Frage]:
        return sorted([f for f in self.fragen if f.sparte == sparte],
                      key=lambda f: f.reihenfolge)

    def frage(self, key: str) -> Frage | None:
        for f in self.fragen:
            if f.key == key:
                return f
        return None

    def klasse_fuer(self, punkte: int) -> str:
        for klasse, ab in self.klassen:
            if punkte >= ab:
                return klasse
        return self.klassen[-1][0] if self.klassen else "C"

    def stufe(self, versuch_nr: int) -> KaskadenStufe | None:
        for s in self.kaskade:
            if s.versuch_nr == versuch_nr:
                return s
        return None

    def letzte_stufe(self) -> int:
        return max((s.versuch_nr for s in self.kaskade), default=0)

    def gruende_der_phase(self, phase: str) -> list[Grund]:
        return [g for g in self.gruende if g.phase == phase]


def _ja(wert) -> bool:
    return str(wert or "").strip().upper() in ("J", "JA", "X", "1", "TRUE")


def _erfassungs_keys() -> set[str]:
    """Gültige Frage-Keys des bestehenden Erfassungsbogens (konfigurator_logik,
    Blätter Fragen / Fragen PV / Fragen KL) – für die Warnungs-Prüfung."""
    keys: set[str] = set()
    try:
        from openpyxl import load_workbook
        wb = load_workbook(config.LOGIK_EXCEL_PFAD, read_only=True, data_only=True)
        for blatt in ("Fragen", "Fragen PV", "Fragen KL"):
            if blatt not in wb.sheetnames:
                continue
            for zeile in wb[blatt].iter_rows(min_row=2, values_only=True):
                if len(zeile) > 1 and zeile[1]:
                    keys.add(str(zeile[1]).strip())
        wb.close()
    except Exception:
        pass
    return keys


def einlesen(pfad: Path | None = None) -> LeadLogik:
    pfad = Path(pfad or LOGIK_PFAD)
    logik = LeadLogik()
    if not pfad.exists():
        logik.fehler.append(f"Steuerdatei fehlt: {pfad}")
        return logik
    try:
        from openpyxl import load_workbook
        wb = load_workbook(pfad, read_only=True, data_only=True)
    except Exception as problem:
        logik.fehler.append(f"Datei nicht lesbar: {problem}")
        return logik
    logik.stand = datetime.fromtimestamp(pfad.stat().st_mtime).strftime("%d.%m.%Y %H:%M")

    def zeilen(blatt):
        if blatt not in wb.sheetnames:
            logik.fehler.append(f"Blatt „{blatt}“ fehlt")
            return []
        return [z for z in wb[blatt].iter_rows(min_row=2, values_only=True)
                if any(x is not None and str(x).strip() != "" for x in z)]

    gesehen: set[str] = set()
    for z in zeilen("Qualifizierung"):
        sparte = str(z[0] or "").strip().upper()
        key = str(z[1] or "").strip()
        if not key:
            continue
        if sparte not in SPARTEN:
            logik.fehler.append(f"Qualifizierung {key}: unbekannte Sparte „{sparte}“")
            continue
        if key in gesehen:
            logik.warnungen.append(f"Qualifizierung: {key} doppelt – erste Zeile gilt")
            continue
        gesehen.add(key)
        typ = str(z[3] or "text").strip().lower()
        if typ not in ("janein", "auswahl", "mehrfach", "zahl", "text"):
            logik.fehler.append(f"Qualifizierung {key}: unbekannter Typ „{typ}“")
            continue
        try:
            reihe = int(z[6] or 0)
        except (TypeError, ValueError):
            reihe = 0
        logik.fragen.append(Frage(
            sparte=sparte, key=key, frage=str(z[2] or "").strip(), typ=typ,
            optionen=[o.strip() for o in str(z[4] or "").split("|") if o.strip()],
            pflicht=_ja(z[5]), reihenfolge=reihe,
            erfassungs_frage=str(z[7] or "").strip() if len(z) > 7 else "",
            unqualifiziert_bei=str(z[8] or "").strip() if len(z) > 8 else ""))

    frage_keys = {f.key for f in logik.fragen}
    for z in zeilen("Scoring"):
        key = str(z[0] or "").strip()
        if not key:
            continue
        try:
            punkte = int(z[2] or 0)
        except (TypeError, ValueError):
            logik.fehler.append(f"Scoring {key}: Punkte „{z[2]}“ keine Zahl")
            continue
        if key not in frage_keys:
            logik.warnungen.append(f"Scoring: Frage-Key {key} unbekannt")
        logik.scoring.append(ScoringRegel(
            frage_key=key, bedingung=str(z[1] or "").strip(), punkte=punkte,
            bemerkung=str(z[3] or "").strip() if len(z) > 3 else ""))

    for z in zeilen("Klassen"):
        try:
            logik.klassen.append((str(z[0] or "").strip().upper(), int(z[1] or 0)))
        except (TypeError, ValueError):
            logik.fehler.append(f"Klassen: ab_punkte „{z[1]}“ keine Zahl")
    logik.klassen.sort(key=lambda k: -k[1])

    for z in zeilen("Kaskade"):
        try:
            nr = int(z[0] or 0)
        except (TypeError, ValueError):
            continue
        logik.kaskade.append(KaskadenStufe(
            versuch_nr=nr, wiedervorlage_nach=str(z[1] or "").strip(),
            aktion=str(z[2] or "keine").strip(), letzter=_ja(z[3])))
    logik.kaskade.sort(key=lambda s: s.versuch_nr)

    for z in zeilen("Gruende"):
        logik.gruende.append(Grund(
            phase=str(z[0] or "").strip(), grund=str(z[1] or "").strip(),
            freitext_pflicht=_ja(z[2])))

    for z in zeilen("Wunschzeiten"):
        logik.wunschzeiten.append(Wunschzeit(
            key=str(z[0] or "").strip(), bezeichnung=str(z[1] or "").strip(),
            von=str(z[2] or "").strip(), bis=str(z[3] or "").strip(),
            wochentage=str(z[4] or "").strip()))

    # Warnung: erfassungs_frage-Keys gegen den bestehenden Erfassungsbogen
    gueltig = _erfassungs_keys()
    if gueltig:
        for f in logik.fragen:
            if f.erfassungs_frage and f.erfassungs_frage not in gueltig:
                logik.warnungen.append(
                    f"{f.key}: erfassungs_frage „{f.erfassungs_frage}“ existiert "
                    "nicht im Erfassungsbogen (Vorbelegung wird übersprungen)")
    wb.close()
    if not logik.fragen and not logik.fehler:
        logik.fehler.append("Blatt Qualifizierung ist leer")
    return logik


def bedingung_trifft(bedingung: str, wert) -> bool:
    """`=Wert`, `>N`, `>=N`, `<N`, `<=N`, `enthält Wert` gegen eine Antwort
    (mehrfach-Antworten kommen als Liste)."""
    bedingung = (bedingung or "").strip()
    if not bedingung:
        return False
    if bedingung.lower().startswith("enthält"):
        gesucht = bedingung[7:].strip()
        werte = wert if isinstance(wert, list) else [wert]
        return any(gesucht.lower() == str(w).strip().lower() for w in werte)
    if bedingung.startswith("="):
        soll = bedingung[1:].strip().lower()
        ist = str(wert if not isinstance(wert, list) else ",".join(map(str, wert)))
        return ist.strip().lower() == soll
    treffer = re.match(r"^(>=|<=|>|<)\s*(-?\d+(?:[.,]\d+)?)$", bedingung)
    if treffer:
        try:
            zahl = float(str(wert).replace(",", "."))
        except (TypeError, ValueError):
            return False
        grenze = float(treffer.group(2).replace(",", "."))
        return {"<": zahl < grenze, "<=": zahl <= grenze,
                ">": zahl > grenze, ">=": zahl >= grenze}[treffer.group(1)]
    return False


_cache: dict = {"logik": None, "mtime": None}


def hole_logik(erzwingen: bool = False) -> LeadLogik:
    """Eingelesene Logik mit mtime-Cache (wie projektierung_logik)."""
    mtime = LOGIK_PFAD.stat().st_mtime if LOGIK_PFAD.exists() else None
    if (not erzwingen and _cache["logik"] is not None
            and _cache["mtime"] == mtime):
        return _cache["logik"]
    logik = einlesen()
    _cache.update(logik=logik, mtime=mtime)
    return logik
