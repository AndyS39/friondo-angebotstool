# Klima-Artikelimport (v24, PLAN_V16 Phase 113): liest die TAIFUN-Positions-
# listen im Ordner Artikel-Preislisten/Klima/ (Layout wie die WP-Preisliste:
# GUID · Position · Menge · Einheit · Beschreibung · E-Preis · G-Preis ·
# Einkaufspreis Material) und legt sie als eigenen Klima-Artikelbereich an:
#
#   KL001 … KL048 = „Ersatzangebot KL.xlsx“ Pos. 001–048 (Artikelquelle),
#   KL049         = „Musterangebot KL.xlsx“ Pos. 006 (Systemgarantie, 0 €),
#   KL050         = Rollgerüst / Arbeitsgerüst (Zusatzartikel ohne TAIFUN-GUID,
#                   VK als Startwert aus dem Parameter „Rollgerüst VK“ des Blatts
#                   „KL-Parameter“, EK leer – Innendienst ergänzt ihn in der
#                   Parametrierung; danach gilt der Artikelpreis der Parametrierung,
#                   der Re-Import fasst Preis/Text von KL050 nicht mehr an).
#
# Nummernvergabe (wiederholbarer Import, Muster app/import_pv.py):
#   1. Blatt „KL-Artikel“ der Logik-Excel pinnt GUID → KL-Nr. (Entwicklungs-PC
#      und Server bekommen garantiert dieselben Nummern, auf die die Blätter
#      „Aktionen KL“ / „Paketmatrix KL“ / „Montagematrix KL“ verweisen);
#   2. sonst: Position in der Datei (Pos. 001 → KL001 …, Muster Pos. 006 → KL049).
#   Re-Import aktualisiert Texte/Preise über die Nummer bzw. den GUID-Anker und
#   legt nie doppelt an. Der WP-Preislisten-Import und der PV-Import fassen
#   KL-Artikel nie an (eigene Quelle „kl“) – und umgekehrt.
#
# Textregeln (Blatt „Textregeln“, Zeilen „Positionen KL020–KL025“, „Position
# KL013“, „Positionen KL026–KL029“) greifen bei jedem (Re-)Import; fehlen die
# Zeilen, gelten die Texte aus PLAN_V16 als Standard (mit Warnung im Bericht).

import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
from sqlalchemy.orm import Session

from app import config, models
from app.import_preisliste import (Diff, FELD_NAMEN, _header_indizes,
                                   euro_zu_cent, text_bereinigen)
from app.models import Artikel

# Agent C trägt QUELLE_KL in app/models.py ein – bis dahin derselbe Wert als Fallback
QUELLE_KL = getattr(models, "QUELLE_KL", "kl")

KL_ORDNER_STANDARD = config.PROJEKT_ORDNER / "Artikel-Preislisten" / "Klima"
ERSATZANGEBOT = ("Ersatzangebot KL.xlsx", "Ersatzangebot KL")
MUSTERANGEBOT = ("Musterangebot KL.xlsx", "Musterangebot KL")
ERSATZ_MAX_POS = 48                       # Pos. 001–048 → KL001–KL048
MUSTER_POSITIONEN = {"006": "KL049"}      # nur die Systemgarantie aus dem Muster
MUSTER_GUIDS = {"KL049": "{38FC5932-9389-46DE-A9D0-8855676E1FD8}"}

ROLLGERUEST_NR = "KL050"
ROLLGERUEST_BEZEICHNUNG = "Rollgerüst / Arbeitsgerüst, Auf- und Abbau"
ROLLGERUEST_EINHEIT = "psl."
ROLLGERUEST_VK_STANDARD = Decimal("499")
ROLLGERUEST_PARAMETER = "Rollgerüst VK"
ROLLGERUEST_EK_HINWEIS = ("KL050 „Rollgerüst / Arbeitsgerüst“ ohne EK – bitte EK in der "
                          "Parametrierung (Artikel → KL050) ergänzen.")

BLATT_ARTIKEL = "KL-Artikel"
BLATT_PARAMETER = "KL-Parameter"
BLATT_TEXTREGELN = "Textregeln"

FELDER_KL = ["kategorie", "bezeichnung", "beschreibung", "menge_standard", "einheit",
             "e_preis_cent", "ep_flag", "ek_cent"]

# v27 (PLAN_V17 Phase 128): Blockgröße – Commit je BLOCK geschriebene Zeilen
BLOCK = 200

# --- Textregeln (Standardtexte aus PLAN_V16 Phase 113 – Fallback) ----------

LEISTUNGSUMFANG_ANKER = "Diese Leistung besteht aus folgenden Positionen"
LEISTUNGSUMFANG_STANDARD = [
    "– Montage der Innengeräte (Wandmontage) und des Außengeräts auf Wand- oder "
    "Bodenkonsole (Konsole enthalten)",
    "– Kältemittelleitung isoliert, Kondensatleitung, Elektroverbindung Innen-/Außengerät "
    "und Montagekanal bis 5 m je Innengerät",
    "– eine Kernbohrung je Innengerät durch die Außenwand (bis 32 cm)",
    "– Evakuieren, Dichtheitsprüfung mit Dokumentation (F-Gase-Verordnung), Inbetriebnahme "
    "und Einweisung",
    "– An- und Abfahrt, Kleinmaterial, Entsorgung des Verpackungsmaterials",
]
LEISTUNGSUMFANG_NUMMERN = [f"KL{n:03d}" for n in range(20, 26)]
KL013_ALT = "für Split-Klimagerät CL3000i"
KL013_NEU = "für Bosch Climate Split-Klimageräte"
KURZ_NUMMERN = [f"KL{n:03d}" for n in range(26, 30)]
KURZ_BLEIBT = "Split Inneneinheit"
KURZ_MUSTER = "BOSCH CL3200iU W <26|35|53|70> E Inneneinheit <kW>"
KURZ_GRENZE = 60

_Q = "['‚‘’`´]"                       # gerade und typografische einfache Anführungszeichen
_QI = "[^'‚‘’`´]+"


@dataclass
class KlTextregeln:
    ergaenzen: dict[str, tuple[str, list[str]]] = field(default_factory=dict)   # Nr -> (Anker, Zeilen)
    ersetzen: dict[str, list[tuple[str, str]]] = field(default_factory=dict)    # Nr -> [(alt, neu)]
    kurz: dict[str, tuple[str, str, int]] = field(default_factory=dict)         # Nr -> (bleibt, Muster, Grenze)
    warnungen: list[str] = field(default_factory=list)


def kl_ordner() -> Path:
    return getattr(config, "KL_PREISLISTEN_ORDNER", KL_ORDNER_STANDARD)


def kl_nummer(zahl: int) -> str:
    return f"KL{zahl:03d}"


def _nummern_bereich(betrifft: str) -> list[str]:
    """„Position KL013“ → [KL013]; „Positionen KL020–KL025“ → [KL020 … KL025]."""
    m = re.fullmatch(r"Position(?:en)?\s+KL(\d{3})(?:\s*[–-]\s*KL(\d{3}))?", betrifft.strip())
    if not m:
        return []
    von, bis = int(m.group(1)), int(m.group(2) or m.group(1))
    return [kl_nummer(n) for n in range(von, bis + 1)]


def _logik_workbook():
    pfad = config.LOGIK_EXCEL_PFAD
    if not pfad.exists():
        return None
    return openpyxl.load_workbook(pfad, read_only=True, data_only=True)


def kl_textregeln_lesen(wb=None) -> KlTextregeln:
    """Blatt „Textregeln“ der Logik-Excel: KL-Zeilen (Betrifft „Position KL013“ /
    „Positionen KL020–KL025“). Drei Regelarten:
      · „nach der Zeile '<Anker>' den Leistungsumfang ergänzen:“ + Folgezeilen,
      · „'<alt>' ersetzen durch '<neu>'“,
      · „Zeile '<bleibt>' bleibt; Bezeichnung = '<Muster>' … länger als <n> Zeichen“.
    Fehlende Regeln werden mit den Standardtexten aus PLAN_V16 ersetzt (Warnung)."""
    regeln = KlTextregeln()
    if wb is None:
        wb = _logik_workbook()
    if wb is not None and BLATT_TEXTREGELN in wb.sheetnames:
        for row in wb[BLATT_TEXTREGELN].iter_rows(min_row=2, values_only=True):
            betrifft, regel = (str(v or "").strip() for v in (tuple(row) + (None,) * 2)[:2])
            nummern = _nummern_bereich(betrifft)
            if not nummern or not regel:
                continue
            m_erg = re.search(rf"nach der Zeile\s*{_Q}({_QI}){_Q}\s*(?:den\s+)?Leistungsumfang"
                              rf"\s+ergänzen:?\s*(.*)$", regel, re.S | re.I)
            m_ers = re.search(rf"{_Q}({_QI}){_Q}\s*ersetzen durch\s*{_Q}({_QI}){_Q}", regel)
            m_kurz = re.search(rf"Zeile\s*{_Q}({_QI}){_Q}\s*bleibt;\s*Bezeichnung\s*=\s*"
                               rf"{_Q}({_QI}){_Q}", regel)
            if m_erg:
                zeilen = [z.strip() for z in m_erg.group(2).splitlines() if z.strip()]
                if not zeilen:
                    regeln.warnungen.append(
                        f"Textregel „{betrifft}“: keine Ergänzungszeilen gefunden – ignoriert.")
                    continue
                for nr in nummern:
                    regeln.ergaenzen[nr] = (m_erg.group(1).strip(), zeilen)
            elif m_kurz:
                m_grenze = re.search(r"länger als\s*(\d+)\s*Zeichen", regel)
                grenze = int(m_grenze.group(1)) if m_grenze else KURZ_GRENZE
                for nr in nummern:
                    regeln.kurz[nr] = (m_kurz.group(1).strip(), m_kurz.group(2).strip(), grenze)
            elif m_ers:
                for nr in nummern:
                    regeln.ersetzen.setdefault(nr, []).append(
                        (m_ers.group(1).strip(), m_ers.group(2).strip()))
            else:
                regeln.warnungen.append(
                    f"Textregel „{betrifft}“ – „{regel[:60]}…“ nicht automatisch anwendbar.")
    elif wb is None:
        regeln.warnungen.append("Logik-Excel nicht gefunden – KL-Textregeln aus PLAN_V16 "
                                "als Standard verwendet.")
    else:
        regeln.warnungen.append("Blatt „Textregeln“ fehlt – KL-Textregeln aus PLAN_V16 "
                                "als Standard verwendet.")
    # Fallbacks (Standardtexte aus PLAN_V16 Phase 113)
    if not any(nr in regeln.ergaenzen for nr in LEISTUNGSUMFANG_NUMMERN):
        regeln.warnungen.append(
            "Textregel „Positionen KL020–KL025“ (Leistungsumfang) fehlt im Blatt "
            "„Textregeln“ – Standardtext aus PLAN_V16 verwendet.")
        for nr in LEISTUNGSUMFANG_NUMMERN:
            regeln.ergaenzen[nr] = (LEISTUNGSUMFANG_ANKER, list(LEISTUNGSUMFANG_STANDARD))
    if "KL013" not in regeln.ersetzen:
        regeln.warnungen.append(
            "Textregel „Position KL013“ fehlt im Blatt „Textregeln“ – Standardtext "
            "aus PLAN_V16 verwendet.")
        regeln.ersetzen["KL013"] = [(KL013_ALT, KL013_NEU)]
    if not any(nr in regeln.kurz for nr in KURZ_NUMMERN):
        regeln.warnungen.append(
            "Textregel „Positionen KL026–KL029“ (Kurzbezeichnung) fehlt im Blatt "
            "„Textregeln“ – Standardtext aus PLAN_V16 verwendet.")
        for nr in KURZ_NUMMERN:
            regeln.kurz[nr] = (KURZ_BLEIBT, KURZ_MUSTER, KURZ_GRENZE)
    return regeln


def kurzbezeichnung(erste_zeile: str, muster: str) -> str | None:
    """„BOSCH Klimagerät CL3200iU W 26 E, Split Inneneinheit, 292x729x200, 2,6 kW“
    + Muster „BOSCH CL3200iU W <26|35|53|70> E Inneneinheit <kW>“
    → „BOSCH CL3200iU W 26 E Inneneinheit 2,6 kW“ (None, wenn nicht ableitbar)."""
    m_code = re.search(r"\bW\s*(\d{2})\s*E\b", erste_zeile)
    m_kw = re.search(r"(\d+(?:[,.]\d+)?)\s*kW", erste_zeile)
    if not m_code or not m_kw:
        return None
    code = m_code.group(1)
    m_liste = re.search(r"<([^<>]*\|[^<>]*)>", muster)
    if m_liste and code not in [t.strip() for t in m_liste.group(1).split("|")]:
        return None
    text = re.sub(r"<[^<>]*\|[^<>]*>", code, muster)
    return text.replace("<kW>", f"{m_kw.group(1)} kW").strip()


def texte_anwenden(nr: str, text: str, regeln: KlTextregeln,
                   warnungen: list[str] | None = None) -> tuple[str, str]:
    """Textregeln auf den bereinigten Beschreibungstext anwenden; liefert
    (bezeichnung = erste Zeile, beschreibung = Rest)."""
    for alt, neu in regeln.ersetzen.get(nr, []):
        text = text.replace(alt, neu)
    zeilen = [z for z in text.splitlines() if z.strip()]
    if not zeilen:
        return "", ""
    bezeichnung, rest = zeilen[0].strip(), [z.strip() for z in zeilen[1:]]
    if nr in regeln.ergaenzen:
        anker, zusatz = regeln.ergaenzen[nr]
        if not any(z == zusatz[0] for z in rest):
            treffer = [i for i, z in enumerate(rest) if z.lower() == anker.lower()]
            if treffer:
                i = treffer[0] + 1
                rest = rest[:i] + list(zusatz) + rest[i:]
            else:
                if warnungen is not None:
                    warnungen.append(f"{nr}: Ankerzeile „{anker}“ nicht im Text – "
                                     "Leistungsumfang am Ende ergänzt.")
                rest = rest + list(zusatz)
    if nr in regeln.kurz:
        bleibt, muster, grenze = regeln.kurz[nr]
        if len(bezeichnung) > grenze and bleibt.lower() in bezeichnung.lower():
            kurz = kurzbezeichnung(bezeichnung, muster)
            if kurz:
                rest = [bezeichnung] + rest
                bezeichnung = kurz
            elif warnungen is not None:
                warnungen.append(f"{nr}: Kurzbezeichnung aus „{bezeichnung[:50]}…“ nicht "
                                 "ableitbar – Bezeichnung bleibt die erste Zeile.")
    return bezeichnung, "\n".join(rest)


# --- Pinning und Parameter aus der Logik-Excel ------------------------------

def pins_lesen(wb=None) -> dict[str, str]:
    """Blatt „KL-Artikel“ der Logik-Excel: GUID → KL-Nr. (leer = ohne Pins)."""
    if wb is None:
        wb = _logik_workbook()
    if wb is None or BLATT_ARTIKEL not in wb.sheetnames:
        return {}
    pins: dict[str, str] = {}
    zeilen = wb[BLATT_ARTIKEL].iter_rows(values_only=True)
    kopf = [str(z or "").strip() for z in next(zeilen, [])]
    try:
        i_nr, i_guid = kopf.index("KL-Nr."), kopf.index("GUID")
    except ValueError:
        return {}
    for row in zeilen:
        row = tuple(row) + (None,) * max(0, max(i_nr, i_guid) + 1 - len(row))
        nr = str(row[i_nr] or "").strip()
        guid = str(row[i_guid] or "").strip()
        if nr and guid.startswith("{"):
            pins[guid] = nr
    return pins


def rollgeruest_vk(wb=None, warnungen: list[str] | None = None) -> Decimal:
    """Parameter „Rollgerüst VK“ (Blatt „KL-Parameter“) als Startwert für KL050;
    fehlt die Zeile, gilt 499 € (Warnung)."""
    if wb is None:
        wb = _logik_workbook()
    if wb is not None and BLATT_PARAMETER in wb.sheetnames:
        for row in wb[BLATT_PARAMETER].iter_rows(min_row=2, values_only=True):
            name, wert = (tuple(row) + (None,) * 2)[:2]
            if str(name or "").strip() == ROLLGERUEST_PARAMETER:
                try:
                    return Decimal(str(wert).replace(".", "").replace(",", ".")
                                   if isinstance(wert, str) else str(wert))
                except (InvalidOperation, TypeError):
                    break
    if warnungen is not None:
        warnungen.append(f"KL-Parameter „{ROLLGERUEST_PARAMETER}“ fehlt oder ist keine Zahl – "
                         f"Standard {ROLLGERUEST_VK_STANDARD} € für KL050 verwendet.")
    return ROLLGERUEST_VK_STANDARD


# --- Dateien lesen -----------------------------------------------------------

@dataclass
class KlLeseErgebnis:
    artikel: list[dict] = field(default_factory=list)
    warnungen: list[str] = field(default_factory=list)


def _positionszeilen(pfad: Path, warnungen: list[str]):
    """Liefert (pos, guid, beschreibung_roh, einheit, e_preis, ek) je Positionszeile."""
    wb = openpyxl.load_workbook(pfad, data_only=True)
    ws = wb[wb.sheetnames[0]]
    zeilen = ws.iter_rows(values_only=True)
    # Spaltenhinweise (Multi/Artikelnummer/EK-Datum fehlen in den Angebots-
    # exporten immer) werden wie beim PV-Import nicht in den Bericht übernommen
    spalten_warnungen: list[str] = []
    indizes = _header_indizes(next(zeilen), spalten_warnungen)

    def zelle(row, name):
        i = indizes.get(name)
        return row[i] if i is not None and i < len(row) else None

    for row in zeilen:
        guid, beschreibung, pos = zelle(row, "GUID"), zelle(row, "Beschreibung"), zelle(row, "Position")
        if guid is None and beschreibung is None:
            continue
        if pos is None:
            continue   # Kategoriezeile (z. B. „Klimaanlage Bosch“ im Muster) – keine Position
        yield (str(pos).strip(), (str(guid).strip() if guid is not None else None),
               beschreibung, str(zelle(row, "Einheit") or "").strip(),
               zelle(row, "E-Preis"), zelle(row, "Einkaufspreis Material"))


def _artikel(nr: str, guid, text: str, einheit: str, e_preis, ek, kurz: str, pos: str,
             regeln: KlTextregeln, warnungen: list[str]) -> dict:
    bezeichnung, beschreibung = texte_anwenden(nr, text, regeln, warnungen)
    return {
        "guid": guid,
        "pos_nr": nr,
        "kategorie": f"KL · {kurz}",
        "bezeichnung": bezeichnung,
        "beschreibung": beschreibung,
        "menge_standard": 1.0,
        "einheit": einheit,
        "e_preis_cent": euro_zu_cent(e_preis or 0),
        "ep_flag": False,                      # EP nur über die Aktion (KL013 „als EP“)
        "ek_cent": euro_zu_cent(ek) if ek not in (None, "") else None,
        "quelle": QUELLE_KL,
        "_datei": kurz,
        "_position": pos,
    }


def rollgeruest_artikel(vk: Decimal | None = None) -> dict:
    """KL050 – Zusatzartikel ohne TAIFUN-GUID (Entscheidung 02.10.2026)."""
    return {
        "guid": None,
        "pos_nr": ROLLGERUEST_NR,
        "kategorie": "KL · Zusatz",
        "bezeichnung": ROLLGERUEST_BEZEICHNUNG,
        "beschreibung": "",
        "menge_standard": 1.0,
        "einheit": ROLLGERUEST_EINHEIT,
        "e_preis_cent": euro_zu_cent(vk if vk is not None else ROLLGERUEST_VK_STANDARD),
        "ep_flag": False,
        "ek_cent": None,
        "quelle": QUELLE_KL,
        "_datei": "Zusatz",
        "_position": "–",
    }


def lese_kl_dateien(pins: dict[str, str] | None = None,
                    regeln: KlTextregeln | None = None,
                    mit_rollgeruest: bool = True) -> KlLeseErgebnis:
    """Liest Ersatzangebot (Pos. 001–048 → KL001–KL048) und Musterangebot (nur
    Pos. 006 → KL049); gepinnte GUIDs (Blatt „KL-Artikel“) haben Vorrang vor
    der Position. KL050 wird als Zusatzartikel angehängt."""
    ergebnis = KlLeseErgebnis()
    wb_logik = None
    if pins is None or regeln is None:
        wb_logik = _logik_workbook()
    pins = pins_lesen(wb_logik) if pins is None else pins
    regeln = kl_textregeln_lesen(wb_logik) if regeln is None else regeln
    ergebnis.warnungen.extend(regeln.warnungen)
    vergeben: dict[str, str] = {}       # KL-Nr -> Herkunft

    def aufnehmen(nr, guid, roh, einheit, e_preis, ek, kurz, pos):
        text = text_bereinigen(roh)
        if not text:
            ergebnis.warnungen.append(f"{kurz} Pos. {pos}: leere Beschreibung – übersprungen.")
            return
        if nr in vergeben:
            ergebnis.warnungen.append(
                f"{kurz} Pos. {pos}: Nummer {nr} ist bereits durch {vergeben[nr]} belegt – "
                "übersprungen (Blatt „KL-Artikel“ prüfen).")
            return
        vergeben[nr] = f"{kurz} Pos. {pos}"
        ergebnis.artikel.append(_artikel(nr, guid, text, einheit, e_preis, ek, kurz, pos,
                                         regeln, ergebnis.warnungen))

    # 1) Ersatzangebot KL – Artikelquelle
    dateiname, kurz = ERSATZANGEBOT
    pfad = kl_ordner() / dateiname
    if not pfad.exists():
        ergebnis.warnungen.append(f"Klima-Positionsliste fehlt: {pfad}")
    else:
        for pos, guid, roh, einheit, e_preis, ek in _positionszeilen(pfad, ergebnis.warnungen):
            if guid is None:
                ergebnis.warnungen.append(f"{kurz} Pos. {pos}: ohne GUID – übersprungen.")
                continue
            nr = pins.get(guid)
            if nr is None:
                if not pos.isdigit() or not 1 <= int(pos) <= ERSATZ_MAX_POS:
                    ergebnis.warnungen.append(
                        f"{kurz} Pos. {pos}: außerhalb 001–{ERSATZ_MAX_POS:03d} und nicht im "
                        "Blatt „KL-Artikel“ gepinnt – übersprungen.")
                    continue
                nr = kl_nummer(int(pos))
                if pins:
                    ergebnis.warnungen.append(
                        f"{kurz} Pos. {pos}: GUID {guid} nicht im Blatt „KL-Artikel“ – "
                        f"Nummer nach Position vergeben ({nr}); Blatt bitte aktualisieren.")
            aufnehmen(nr, guid, roh, einheit, e_preis, ek, kurz, pos)

    # 2) Musterangebot KL – nur die Systemgarantie (Pos. 006 → KL049)
    dateiname, kurz = MUSTERANGEBOT
    pfad = kl_ordner() / dateiname
    if not pfad.exists():
        ergebnis.warnungen.append(f"Klima-Positionsliste fehlt: {pfad}")
    else:
        for pos, guid, roh, einheit, e_preis, ek in _positionszeilen(pfad, ergebnis.warnungen):
            nr = pins.get(guid) if guid else None
            if nr is None:
                nr = MUSTER_POSITIONEN.get(pos)
            if nr is None or nr not in MUSTER_POSITIONEN.values():
                continue
            if guid is None:
                ergebnis.warnungen.append(f"{kurz} Pos. {pos}: ohne GUID – übersprungen.")
                continue
            aufnehmen(nr, guid, roh, einheit, e_preis, ek, kurz, pos)
        for nr in MUSTER_POSITIONEN.values():
            if nr not in vergeben:
                ergebnis.warnungen.append(
                    f"{kurz}: Position für {nr} (Systemgarantie) nicht gefunden.")

    # 3) KL050 Rollgerüst (Zusatzartikel ohne GUID)
    if mit_rollgeruest:
        vk = rollgeruest_vk(wb_logik, ergebnis.warnungen)
        if ROLLGERUEST_NR in vergeben:
            ergebnis.warnungen.append(
                f"{ROLLGERUEST_NR} ist durch {vergeben[ROLLGERUEST_NR]} belegt – Rollgerüst "
                "nicht angelegt (Blatt „KL-Artikel“ prüfen).")
        else:
            ergebnis.artikel.append(rollgeruest_artikel(vk))
    return ergebnis


# --- Abgleich mit der Datenbank ---------------------------------------------

def _zuordnen(session: Session, ergebnis: KlLeseErgebnis):
    """Ordnet jede gelesene Zeile einem Bestandsartikel zu (Nummer, sonst GUID).
    Liefert [(daten, bestand_oder_None)] + Bestand."""
    bestand = session.query(Artikel).filter(Artikel.quelle == QUELLE_KL).all()
    nach_nr = {a.pos_nr: a for a in bestand}
    nach_guid = {a.guid: a for a in bestand if a.guid}
    paare = []
    belegt: set[int] = set()
    for daten in ergebnis.artikel:
        vorhanden = nach_nr.get(daten["pos_nr"])
        if vorhanden is None and daten["guid"]:
            vorhanden = nach_guid.get(daten["guid"])
        if vorhanden is not None and vorhanden.id in belegt:
            vorhanden = None
        if vorhanden is not None:
            belegt.add(vorhanden.id)
        paare.append((daten, vorhanden))
    return paare, bestand


def _felder_fuer(daten: dict, vorhanden: Artikel | None) -> list[str]:
    """Felder, die der Import setzt: alles – außer bei einem vorhandenen KL050,
    dessen Preis/EK/Text der Innendienst in der Parametrierung pflegt."""
    if vorhanden is not None and daten["pos_nr"] == ROLLGERUEST_NR:
        return ["kategorie", "menge_standard", "einheit"]
    return list(FELDER_KL)


def berechne_diff(session: Session, ergebnis: KlLeseErgebnis | None = None) -> Diff:
    ergebnis = ergebnis or lese_kl_dateien()
    diff = Diff(warnungen=list(ergebnis.warnungen))
    paare, bestand = _zuordnen(session, ergebnis)
    gefunden = set()
    for daten, vorhanden in paare:
        if vorhanden is None:
            diff.neu.append(daten)
            continue
        gefunden.add(vorhanden.id)
        felder = [f for f in _felder_fuer(daten, vorhanden)
                  if getattr(vorhanden, f) != daten[f]]
        if daten["guid"] and vorhanden.guid != daten["guid"]:
            felder.append("guid")
            if vorhanden.guid:
                diff.warnungen.append(
                    f"{daten['pos_nr']}: andere GUID als bisher ({vorhanden.guid} → "
                    f"{daten['guid']}) – Blatt „KL-Artikel“ bitte auf die neue GUID setzen.")
        if felder:
            diff.geaendert.append((vorhanden, daten,
                                   [FELD_NAMEN.get(f, "GUID") for f in felder]))
        else:
            diff.unveraendert += 1
        if not vorhanden.aktiv:
            diff.reaktiviert.append(vorhanden)
    diff.entfallen = [a for a in bestand if a.id not in gefunden and a.aktiv]
    for a in diff.entfallen:
        diff.warnungen.append(
            f"{a.pos_nr} „{a.titel[:60]}“ ist in keiner Klima-Positionsliste mehr – wird deaktiviert.")
    if any(d["pos_nr"] == ROLLGERUEST_NR and (v is None or v.ek_cent is None)
           for d, v in paare):
        diff.warnungen.append(ROLLGERUEST_EK_HINWEIS)
    return diff


def import_ausfuehren(session: Session) -> tuple[Diff, str]:
    """Wendet den Import an; Rückgabe wie import_pv.import_ausfuehren (Diff, Meldung).

    v27 (PLAN_V17 Phase 128): Schreiben in Blöcken mit Commit je BLOCK Zeilen.
    Anker je Zeile ist die KL-Nummer (Position bzw. Pin), sonst die GUID; jede Zeile
    ist ein Upsert, Verwaiste werden erst nach allen Zeilen deaktiviert. Abbruch
    nach Block n: die Blöcke 1…n sind gespeichert; ein bereits gelöster alter
    GUID-Anker (Schritt 1) wird beim nächsten Lauf über die KL-Nummer wieder aus
    der Datei gesetzt – keine Doppelanlage, kein Datenverlust."""
    from app import betrieb
    with betrieb.import_markieren("Klima-Positionslisten"):
        ergebnis = lese_kl_dateien()
        diff = berechne_diff(session, ergebnis)
        paare, _bestand = _zuordnen(session, ergebnis)
        fremde_guids = {g for (g,) in session.query(Artikel.guid)
                        .filter(Artikel.quelle != QUELLE_KL, Artikel.guid.isnot(None))}
        # GUID-Wechsel in zwei Schritten (guid ist eindeutig): erst alte Anker lösen
        neue_guids = {d["guid"] for d in ergebnis.artikel if d["guid"]}
        for _daten, vorhanden in paare:
            if vorhanden is not None and vorhanden.guid and vorhanden.guid in neue_guids:
                ziel = next((d for d in ergebnis.artikel if d["guid"] == vorhanden.guid), None)
                if ziel is not None and ziel["pos_nr"] != vorhanden.pos_nr:
                    vorhanden.guid = None
        session.flush()
        zaehler = 0
        for daten, vorhanden in paare:
            werte = {k: v for k, v in daten.items() if not k.startswith("_")}
            if werte["guid"] in fremde_guids:    # GUID gehört einem WP-/PV-Artikel
                diff.warnungen.append(
                    f"{werte['pos_nr']}: GUID {werte['guid']} gehört bereits einem anderen "
                    "Artikel – ohne GUID angelegt.")
                werte["guid"] = None
            if vorhanden is None:
                session.add(Artikel(**werte, aktiv=True))
            else:
                for feld in _felder_fuer(daten, vorhanden) + ["pos_nr", "quelle"]:
                    setattr(vorhanden, feld, werte[feld])
                if werte["guid"]:
                    vorhanden.guid = werte["guid"]
                vorhanden.aktiv = True
            zaehler += 1
            if zaehler % BLOCK == 0:
                session.commit()      # Block abschließen – Schreibsperre freigeben
        for artikel in diff.entfallen:
            artikel.aktiv = False
        session.commit()
    meldung = (f"KL: {len(diff.neu)} neu, {len(diff.geaendert)} geändert, "
               f"{diff.unveraendert} unverändert, {len(diff.entfallen)} deaktiviert")
    return diff, meldung


# --- Hilfen für migrate.py / Parametrierung ---------------------------------

def kl_referenzen() -> list[str]:
    """KL-Nummern, die die Logik-Blätter referenzieren (über logik.artikel_referenzen,
    sobald der Lader die KL-Blätter kennt); sonst alle Nummern des Blatts
    „KL-Artikel“ + KL050."""
    try:
        from app import logik as _logik
        l, _b = _logik.logik_einlesen()
        refs = sorted(r for r in _logik.artikel_referenzen(l) if re.fullmatch(r"KL\d{3}", r))
        if refs:
            return refs
    except Exception:
        pass
    try:
        nummern = set(pins_lesen().values())
    except Exception:
        nummern = set()
    if nummern:
        nummern.add(ROLLGERUEST_NR)
    return sorted(nummern)


def fehlende_kl_artikel(session: Session) -> list[str]:
    """Referenzierte KL-Nummern ohne aktiven Artikel (für den Auto-Import in migrate.py)."""
    refs = kl_referenzen()
    if not refs:
        return []
    vorhanden = {nr for (nr,) in session.query(Artikel.pos_nr)
                 .filter(Artikel.pos_nr.like("KL%"), Artikel.aktiv.is_(True))}
    return [r for r in refs if r not in vorhanden]
