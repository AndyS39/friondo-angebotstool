# Lead-Management V1 (v12) – Router unter /lead-management (Ziel der
# bestehenden Startportal-Karte). Der Plan nennt /leads; dort liegt aber die
# unveränderliche Leads-VOT-Liste (Leitplanke 2) – Entscheidung in
# docs/leadmanagement-entscheidungen.md. /api/leads (Phase 75) bleibt wie geplant.
# Demo-Schalter: Nicht-Sichtbare bekommen auf der Einstiegsseite die alte
# Platzhalterseite (Verhalten wie bisher), auf allen anderen Routen 404.

from datetime import datetime, timedelta
from urllib.parse import quote_plus

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import RedirectResponse
from sqlalchemy.orm import Session

from app import leadmanagement as kern
from app.db import get_session
from app.models import (Benutzer, Kampagne, Kunde, LeadAktivitaet,
                        LeadQualifizierung, LeadQuelle, Vorgang, VotTermin,
                        ANRUF_ERGEBNIS_NAMEN, LEAD_PHASEN, LEAD_PHASEN_NAMEN)
from app.templating import render

router = APIRouter(prefix="/lead-management")


def _gate(request: Request, session: Session) -> None:
    """404 statt 403 – im Demo-Modus soll das Modul unsichtbar sein."""
    if not kern.lead_modul_sichtbar(session, request.state.benutzer):
        raise HTTPException(status_code=404)


@router.get("")
async def startseite(request: Request, session: Session = Depends(get_session)):
    """Einstieg: Sichtbare landen auf der Anrufliste (Phase 76); alle anderen
    sehen die alte Platzhalterseite – exakt wie vor dem Modul."""
    if not kern.lead_modul_sichtbar(session, request.state.benutzer):
        return render(request, "platzhalter.html", aktiv=None,
                      titel="Lead-Management",
                      hinweis="Dieser Bereich ist im Aufbau (Coming soon). "
                              "Die Lead-Arbeit läuft bis dahin wie gewohnt über "
                              "das Angebotstool (Leads VOT).")
    return RedirectResponse("/lead-management/anrufliste", status_code=303)


# --- Phase 75: Schnellanlage, Import, Posteingang unklar ---------------------------

SPARTEN = ("WP", "PV", "KL", "WB")


def _quellen(session: Session) -> list[LeadQuelle]:
    return (session.query(LeadQuelle).filter(LeadQuelle.aktiv.is_(True))
            .order_by(LeadQuelle.name).all())


def _wunschzeiten_liste() -> list:
    from app import leadmanagement_logik
    return leadmanagement_logik.hole_logik().wunschzeiten


@router.get("/lead/{vorgang_id}")
async def lead_akte(request: Request, vorgang_id: int,
                    session: Session = Depends(get_session)):
    """Lead-Akte = Vorgangsakte (v10) mit Lead-Kopf (Phase 79) – bis dahin
    Weiterleitung auf die bestehende Akte."""
    _gate(request, session)
    return RedirectResponse(f"/vorgaenge/{vorgang_id}", status_code=303)


@router.get("/neu")
async def schnellanlage_formular(request: Request,
                                 session: Session = Depends(get_session)):
    _gate(request, session)
    return render(request, "leadmanagement/neu.html", aktiv="/lead-management",
                  mobil=True, quellen=_quellen(session),
                  kampagnen=session.query(Kampagne)
                  .filter(Kampagne.aktiv.is_(True)).order_by(Kampagne.name).all(),
                  wunschzeiten=_wunschzeiten_liste(), sparten=SPARTEN,
                  daten={}, duplikat=None,
                  meldung=request.query_params.get("meldung", ""))


@router.post("/neu")
async def schnellanlage(request: Request, session: Session = Depends(get_session)):
    _gate(request, session)
    form = await request.form()
    benutzer = request.state.benutzer
    daten = {feld: (form.get(feld) or "").strip()
             for feld in ("anrede", "vorname", "nachname", "telefon", "email",
                          "strasse", "plz", "ort", "nachricht")}
    daten["sparten"] = [s for s in form.getlist("sparten") if s in SPARTEN]
    daten["wunschzeiten"] = [w for w in form.getlist("wunschzeiten")]
    daten["einwilligung_werbung"] = form.get("einwilligung_werbung") == "on"
    daten["einwilligung_quelle"] = form.get("einwilligung_quelle") or "telefonisch"

    fehler = []
    if not daten["nachname"]:
        fehler.append("Nachname ist Pflicht.")
    if not daten["telefon"] and not daten["email"]:
        fehler.append("Telefon ist Pflicht (ersatzweise E-Mail).")
    if not daten["plz"]:
        fehler.append("PLZ ist Pflicht.")
    if not daten["sparten"]:
        fehler.append("Mindestens eine Sparte wählen.")
    quelle = None
    if str(form.get("quelle_id") or "").isdigit():
        quelle = session.get(LeadQuelle, int(form.get("quelle_id")))
    if quelle is None:
        fehler.append("Bitte eine Quelle wählen.")
    kampagne_id = (int(form.get("kampagne_id"))
                   if str(form.get("kampagne_id") or "").isdigit() else None)

    entscheidung = form.get("entscheidung") or ""
    duplikat = None
    if not fehler and not entscheidung:
        duplikat = kern.duplikat_pruefen(
            session, telefon=daten["telefon"], email=daten["email"],
            name=f"{daten['vorname']} {daten['nachname']}",
            plz=daten["plz"], strasse=daten["strasse"])
    if fehler or duplikat is not None:
        return render(request, "leadmanagement/neu.html",
                      aktiv="/lead-management", mobil=True,
                      quellen=_quellen(session),
                      kampagnen=session.query(Kampagne)
                      .filter(Kampagne.aktiv.is_(True)).order_by(Kampagne.name).all(),
                      wunschzeiten=_wunschzeiten_liste(), sparten=SPARTEN,
                      daten=dict(daten, quelle_id=form.get("quelle_id"),
                                 kampagne_id=form.get("kampagne_id")),
                      duplikat=duplikat, phasen_namen=LEAD_PHASEN_NAMEN,
                      meldung=" ".join(fehler))
    vorgang, status = kern.lead_anlegen(session, daten, quelle, "manuell",
                                        benutzer=benutzer,
                                        kampagne_id=kampagne_id,
                                        entscheidung=entscheidung)
    session.commit()
    from urllib.parse import quote_plus
    texte = {"neu": "Lead angelegt.",
             "angehaengt": "An den bestehenden Vorgang angehängt (neue Sparte).",
             "wiederkehrer": "Neuer Vorgang am bestehenden Kunden (Wiederkehrer)."}
    return RedirectResponse("/lead-management/anrufliste?meldung="
                            + quote_plus(texte.get(status, "Lead angelegt.")),
                            status_code=303)


# --- CSV/Excel-Import ---------------------------------------------------------------

_IMPORT_FELDER = ["anrede", "vorname", "nachname", "strasse", "plz", "ort",
                  "telefon", "email", "nachricht"]


def _import_ablage() -> "Path":
    from pathlib import Path

    from app import config
    ordner = config.DATA_ORDNER / "lead_import_tmp"
    ordner.mkdir(parents=True, exist_ok=True)
    return ordner


def _import_lesen(pfad) -> list[list[str]]:
    """CSV (;, , oder Tab) oder XLSX → Zeilenliste (max. 2000)."""
    from pathlib import Path
    pfad = Path(pfad)
    if pfad.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(pfad, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        zeilen = [[("" if z is None else str(z)).strip() for z in reihe]
                  for reihe in ws.iter_rows(values_only=True)]
        wb.close()
    else:
        import csv
        text = pfad.read_text(encoding="utf-8-sig", errors="replace")
        trenner = ";" if text.count(";") >= text.count(",") else ","
        if text.count("\t") > max(text.count(";"), text.count(",")):
            trenner = "\t"
        zeilen = [[(z or "").strip() for z in reihe]
                  for reihe in csv.reader(text.splitlines(), delimiter=trenner)]
    return [z for z in zeilen if any(z)][:2000]


@router.get("/import")
async def import_formular(request: Request,
                          session: Session = Depends(get_session)):
    _gate(request, session)
    return render(request, "leadmanagement/import.html",
                  aktiv="/lead-management", quellen=_quellen(session),
                  kampagnen=session.query(Kampagne)
                  .filter(Kampagne.aktiv.is_(True)).order_by(Kampagne.name).all(),
                  sparten=SPARTEN, schritt="upload", zeilen=[], kopf=[],
                  zuordnung={}, felder=_IMPORT_FELDER, token="", form_daten={},
                  duplikate={}, protokoll=None,
                  meldung=request.query_params.get("meldung", ""))


@router.post("/import")
async def import_verarbeiten(request: Request,
                             session: Session = Depends(get_session)):
    """Zweistufig: Upload → Vorschau (Spaltenzuordnung, Duplikat-Markierung)
    → Import mit Protokoll. Die Datei liegt kurz unter data/lead_import_tmp."""
    import json as json_modul
    import secrets
    from pathlib import Path
    from urllib.parse import quote_plus
    _gate(request, session)
    form = await request.form()
    benutzer = request.state.benutzer
    schritt = form.get("schritt") or "upload"
    kontext = dict(aktiv="/lead-management", quellen=_quellen(session),
                   kampagnen=session.query(Kampagne)
                   .filter(Kampagne.aktiv.is_(True)).order_by(Kampagne.name).all(),
                   sparten=SPARTEN, felder=_IMPORT_FELDER, protokoll=None,
                   duplikate={}, meldung="")

    if schritt == "upload":
        datei = form.get("datei")
        if datei is None or not getattr(datei, "filename", ""):
            return RedirectResponse("/lead-management/import?meldung="
                                    + quote_plus("Bitte eine Datei wählen."),
                                    status_code=303)
        inhalt = await datei.read()
        token = secrets.token_hex(8)
        endung = ".xlsx" if datei.filename.lower().endswith(".xlsx") else ".csv"
        (_import_ablage() / f"{token}{endung}").write_bytes(inhalt)
        pfad = _import_ablage() / f"{token}{endung}"
        zeilen = _import_lesen(pfad)
        if not zeilen:
            return RedirectResponse("/lead-management/import?meldung="
                                    + quote_plus("Datei ist leer oder unlesbar."),
                                    status_code=303)
        kopf = zeilen[0]
        # gespeicherte Zuordnung je Quelle vorschlagen, sonst Namensabgleich
        zuordnung = {}
        quelle_id = form.get("quelle_id") or ""
        if quelle_id.isdigit():
            quelle = session.get(LeadQuelle, int(quelle_id))
            if quelle is not None:
                try:
                    zuordnung = json_modul.loads(kern.parameter_holen(
                        session, f"import_mapping_{quelle.key}", "{}"))
                except ValueError:
                    zuordnung = {}
        if not zuordnung:
            for i, name in enumerate(kopf):
                schluessel = name.strip().lower().replace("-", "").replace("ß", "ss")
                for feld in _IMPORT_FELDER:
                    if feld in schluessel or schluessel in feld:
                        zuordnung[str(i)] = feld
                        break
        duplikate = {}
        for nr, zeile in enumerate(zeilen[1:21]):
            werte = {feld: (zeile[int(i)] if int(i) < len(zeile) else "")
                     for i, feld in zuordnung.items()}
            if kern.duplikat_pruefen(session, telefon=werte.get("telefon", ""),
                                     email=werte.get("email", ""),
                                     name=f"{werte.get('vorname', '')} "
                                          f"{werte.get('nachname', '')}",
                                     plz=werte.get("plz", ""),
                                     strasse=werte.get("strasse", "")):
                duplikate[nr] = True
        return render(request, "leadmanagement/import.html", schritt="vorschau",
                      zeilen=zeilen[1:21], kopf=kopf, zuordnung=zuordnung,
                      token=pfad.name, form_daten=dict(form), **{**kontext,
                      "duplikate": duplikate})

    # schritt == "ausfuehren"
    token = Path(form.get("token") or "").name
    pfad = _import_ablage() / token
    if not token or not pfad.exists():
        return RedirectResponse("/lead-management/import?meldung="
                                + quote_plus("Upload abgelaufen – bitte erneut "
                                             "hochladen."), status_code=303)
    zeilen = _import_lesen(pfad)
    kopf, datenzeilen = zeilen[0], zeilen[1:]
    zuordnung = {}
    for i in range(len(kopf)):
        feld = form.get(f"spalte_{i}") or ""
        if feld in _IMPORT_FELDER:
            zuordnung[str(i)] = feld
    quelle = (session.get(LeadQuelle, int(form.get("quelle_id")))
              if str(form.get("quelle_id") or "").isdigit() else None)
    kampagne_id = (int(form.get("kampagne_id"))
                   if str(form.get("kampagne_id") or "").isdigit() else None)
    sparten = [s for s in form.getlist("sparten") if s in SPARTEN]
    if quelle is not None:   # Zuordnung je Quelle merken
        kern.parameter_setzen(session, f"import_mapping_{quelle.key}",
                              json_modul.dumps(zuordnung))
    protokoll = {"angelegt": 0, "angehaengt": 0, "uebersprungen": []}
    for nr, zeile in enumerate(datenzeilen, start=2):
        try:
            werte = {feld: (zeile[int(i)] if int(i) < len(zeile) else "")
                     for i, feld in zuordnung.items()}
            werte["sparten"] = sparten
            werte["einwilligung_quelle"] = "portal"
            if not werte.get("nachname"):
                protokoll["uebersprungen"].append(f"Zeile {nr}: Nachname fehlt")
                continue
            if not werte.get("plz"):
                protokoll["uebersprungen"].append(f"Zeile {nr}: PLZ fehlt")
                continue
            if not werte.get("telefon") and not werte.get("email"):
                protokoll["uebersprungen"].append(
                    f"Zeile {nr}: weder Telefon noch E-Mail")
                continue
            _, status = kern.lead_anlegen(session, werte, quelle, "import",
                                          benutzer=benutzer,
                                          kampagne_id=kampagne_id)
            protokoll["angelegt" if status != "angehaengt"
                      else "angehaengt"] += 1
        except Exception as problem:
            protokoll["uebersprungen"].append(f"Zeile {nr}: {problem}")
    session.commit()
    try:
        pfad.unlink()
    except OSError:
        pass
    return render(request, "leadmanagement/import.html", schritt="fertig",
                  zeilen=[], kopf=[], zuordnung={}, token="", form_daten={},
                  **{**kontext, "protokoll": protokoll})


# --- Posteingang unklar ---------------------------------------------------------------

@router.get("/posteingang")
async def posteingang(request: Request, session: Session = Depends(get_session)):
    _gate(request, session)
    from app.models import LeadPosteingang
    eintraege = (session.query(LeadPosteingang)
                 .filter(LeadPosteingang.status == "offen")
                 .order_by(LeadPosteingang.erstellt_am.desc()).limit(100).all())
    return render(request, "leadmanagement/posteingang.html",
                  aktiv="/lead-management", eintraege=eintraege,
                  meldung=request.query_params.get("meldung", ""))


@router.post("/posteingang/{eintrag_id}/ignorieren")
async def posteingang_ignorieren(request: Request, eintrag_id: int,
                                 session: Session = Depends(get_session)):
    _gate(request, session)
    from app.models import LeadPosteingang
    eintrag = session.get(LeadPosteingang, eintrag_id)
    if eintrag is not None:
        eintrag.status = "ignoriert"
        session.commit()
    return RedirectResponse("/lead-management/posteingang", status_code=303)


@router.get("/posteingang/{eintrag_id}/anlegen")
async def posteingang_anlegen(request: Request, eintrag_id: int,
                              session: Session = Depends(get_session)):
    """„Als Lead anlegen“: Schnellanlage-Formular, soweit möglich vorbefüllt."""
    _gate(request, session)
    from app import lead_parser
    from app.models import LeadPosteingang, ParserRegel
    eintrag = session.get(LeadPosteingang, eintrag_id)
    if eintrag is None:
        return RedirectResponse("/lead-management/posteingang", status_code=303)
    hilfsregel = ParserRegel(format="zeilen", feldzuordnung="{}")
    daten = lead_parser.mail_parsen(hilfsregel, eintrag.betreff, eintrag.body)
    daten.setdefault("email", eintrag.absender)
    daten.setdefault("nachricht", eintrag.body[:1000])
    daten.pop("rohdaten", None)
    return render(request, "leadmanagement/neu.html", aktiv="/lead-management",
                  mobil=True, quellen=_quellen(session),
                  kampagnen=session.query(Kampagne)
                  .filter(Kampagne.aktiv.is_(True)).order_by(Kampagne.name).all(),
                  wunschzeiten=_wunschzeiten_liste(), sparten=SPARTEN,
                  daten=daten, duplikat=None,
                  meldung=f"Vorbefüllt aus Posteingang: {eintrag.betreff}")


# --- Phase 76: Anrufliste, Ergebnis-Buttons, Kaskade, Qualifizierung ----------------

def _anruf_zeilen(session: Session, benutzer, filter_werte: dict) -> list[dict]:
    """Priorisierte Arbeitsliste (Plan 76): (1) SLA gelb/rot, älteste zuerst ·
    (2) fällige nächste Aktionen/Rückrufe · (3) fällige Zurückgestellte ·
    (4) Rest nach Score-Klasse, dann Eingang."""
    from datetime import datetime as dt
    jetzt = dt.now()
    phase_filter = filter_werte.get("phase") or ""
    abfrage = session.query(Vorgang)
    if phase_filter:
        abfrage = abfrage.filter(Vorgang.lead_phase == phase_filter)
    else:
        abfrage = abfrage.filter(Vorgang.lead_phase.in_(
            ("neu", "in_kontaktierung", "zurueckgestellt", "nicht_erreicht")))
    if filter_werte.get("meine") and benutzer is not None:
        abfrage = abfrage.filter(Vorgang.leadmanager_id == benutzer.id)
    if filter_werte.get("quelle_id"):
        abfrage = abfrage.filter(Vorgang.quelle_id == int(filter_werte["quelle_id"]))
    if filter_werte.get("klasse"):
        abfrage = abfrage.filter(Vorgang.score_klasse == filter_werte["klasse"])
    vorgaenge = abfrage.all()

    kunden = {k.id: k for k in session.query(Kunde)
              .filter(Kunde.id.in_({v.kunde_id for v in vorgaenge} or {0}))}
    quellen = {q.id: q for q in session.query(LeadQuelle)}
    mehrfach = {}
    for v in session.query(Vorgang):
        mehrfach[v.kunde_id] = mehrfach.get(v.kunde_id, 0) + 1
    letzte_anrufe: dict[int, LeadAktivitaet] = {}
    for a in (session.query(LeadAktivitaet)
              .filter(LeadAktivitaet.typ == "anruf",
                      LeadAktivitaet.vorgang_id.in_({v.id for v in vorgaenge} or {0}))
              .order_by(LeadAktivitaet.zeitpunkt)):
        letzte_anrufe[a.vorgang_id] = a

    zeilen = []
    for v in vorgaenge:
        kunde = kunden.get(v.kunde_id)
        if kunde is None:
            continue
        if filter_werte.get("sparte") and \
                filter_werte["sparte"] not in (kunde.interesse or ""):
            continue
        if filter_werte.get("plz") and \
                not (kunde.plz or "").startswith(filter_werte["plz"]):
            continue
        suche = (filter_werte.get("q") or "").lower()
        if suche and suche not in " ".join(
                (kunde.vorname or "", kunde.nachname or "",
                 kunde.telefon or "", kunde.ort or "")).lower():
            continue
        # Phasen-Sonderfälle ohne expliziten Filter: nur fällige zeigen
        if not phase_filter and v.lead_phase == "zurueckgestellt" and (
                v.zurueckgestellt_bis is None or v.zurueckgestellt_bis > jetzt):
            continue
        if not phase_filter and v.lead_phase == "nicht_erreicht" and (
                v.naechste_aktion_am is None or v.naechste_aktion_am > jetzt):
            continue
        sla = kern.sla_status(session, v, jetzt)
        letzter = letzte_anrufe.get(v.id)
        if v.lead_phase == "neu" and sla["farbe"] in ("gelb", "rot"):
            gruppe, schluessel = 0, (v.eingang_am or v.angelegt_am).timestamp()
        elif v.naechste_aktion_am is not None and v.naechste_aktion_am <= jetzt:
            gruppe, schluessel = 1, v.naechste_aktion_am.timestamp()
        elif v.lead_phase == "zurueckgestellt":
            gruppe, schluessel = 2, (v.zurueckgestellt_bis or jetzt).timestamp()
        else:
            klassen_rang = {"A": 0, "B": 1, "C": 2}.get(v.score_klasse or "C", 2)
            gruppe, schluessel = 3, klassen_rang * 10 ** 12 + (
                v.eingang_am or v.angelegt_am).timestamp()
        zeilen.append({
            "vorgang": v, "kunde": kunde,
            "quelle": quellen.get(v.quelle_id),
            "sparten": [s for s in (kunde.interesse or "").split(",") if s.strip()],
            "sla": sla, "letzter": letzter,
            "wiederkehrer": mehrfach.get(v.kunde_id, 0) > 1,
            "monday": v.eingang_art == "monday",
            "nummer_pruefen": letzter is not None
                              and letzter.ergebnis == "falsche_nummer",
            "_sortierung": (gruppe, schluessel),
        })
    zeilen.sort(key=lambda z: z["_sortierung"])
    return zeilen


@router.get("/anrufliste")
async def anrufliste_voll(request: Request,
                          session: Session = Depends(get_session)):
    _gate(request, session)
    from app import leadmanagement_logik
    benutzer = request.state.benutzer
    filter_werte = {
        "meine": request.query_params.get("meine", "1") == "1",
        "quelle_id": request.query_params.get("quelle_id", ""),
        "sparte": request.query_params.get("sparte", ""),
        "klasse": request.query_params.get("klasse", ""),
        "plz": request.query_params.get("plz", ""),
        "phase": request.query_params.get("phase", ""),
        "q": request.query_params.get("q", ""),
    }
    zeilen = _anruf_zeilen(session, benutzer, filter_werte)
    if filter_werte["meine"] and not zeilen and not any(
            v for k, v in filter_werte.items() if k != "meine" and v):
        filter_werte["meine"] = False   # ohne eigene Leads direkt „Alle“
        zeilen = _anruf_zeilen(session, benutzer, filter_werte)
    logik = leadmanagement_logik.hole_logik()
    return render(request, "leadmanagement/anrufliste.html",
                  aktiv="/lead-management", zeilen=zeilen,
                  filter_werte=filter_werte, quellen=_quellen(session),
                  phasen_namen=LEAD_PHASEN_NAMEN,
                  ergebnis_namen=ANRUF_ERGEBNIS_NAMEN,
                  unq_gruende=logik.gruende_der_phase("unqualifiziert"),
                  zurueck_gruende=logik.gruende_der_phase("zurueckgestellt"),
                  demo_badge=kern.demo_aktiv(session),
                  meldung=request.query_params.get("meldung", ""))


@router.post("/anruf/{vorgang_id}")
async def anruf_ergebnis(request: Request, vorgang_id: int,
                         session: Session = Depends(get_session)):
    """Ein-Klick-Anrufergebnis (Plan 76): Aktivität, Zähler, Kaskade bzw.
    Folgedialog-Aktionen."""
    from datetime import datetime as dt
    from urllib.parse import quote_plus
    _gate(request, session)
    form = await request.form()
    benutzer = request.state.benutzer
    vorgang = session.get(Vorgang, vorgang_id)
    ergebnis = form.get("ergebnis") or ""
    if vorgang is None or ergebnis not in ANRUF_ERGEBNIS_NAMEN:
        return RedirectResponse("/lead-management/anrufliste", status_code=303)
    jetzt = dt.now()
    if vorgang.erstkontakt_am is None:
        vorgang.erstkontakt_am = jetzt
    vorgang.versuch_nr = (vorgang.versuch_nr or 0) + 1
    if vorgang.lead_phase in ("neu", "zurueckgestellt", "nicht_erreicht", None):
        vorgang.lead_phase = "in_kontaktierung"
        vorgang.zurueckgestellt_bis = None
    naechste = None
    meldung = f"{ANRUF_ERGEBNIS_NAMEN[ergebnis]} protokolliert."
    if ergebnis == "rueckruf_gewuenscht":
        roh = (form.get("rueckruf_am") or "").strip()
        try:
            naechste = dt.strptime(roh, "%Y-%m-%dT%H:%M")
        except ValueError:
            return RedirectResponse("/lead-management/anrufliste?meldung="
                                    + quote_plus("Rückruf: Datum/Uhrzeit ist "
                                                 "Pflicht."), status_code=303)
        vorgang.naechste_aktion_am = naechste
        meldung = f"Rückruf {naechste.strftime('%d.%m.%Y %H:%M')} gemerkt."
    kern.aktivitaet(session, vorgang.id, "anruf",
                    f"Anruf: {ANRUF_ERGEBNIS_NAMEN[ergebnis]}"
                    + (f" – {form.get('notiz')}" if form.get("notiz") else ""),
                    benutzer=benutzer, ergebnis=ergebnis,
                    naechste_aktion_am=naechste)
    if ergebnis == "erreicht":
        vorgang.erreicht_am = vorgang.erreicht_am or jetzt
        vorgang.naechste_aktion_am = None
        session.commit()
        kunde = session.get(Kunde, vorgang.kunde_id)
        sparte = next((s for s in (kunde.interesse or "").split(",")
                       if s.strip() in SPARTEN), "WP").strip()
        return RedirectResponse(
            f"/lead-management/lead/{vorgang.id}/qualifizierung/{sparte}",
            status_code=303)
    if ergebnis in ("nicht_erreicht", "besetzt", "mailbox"):
        meldung = (f"{ANRUF_ERGEBNIS_NAMEN[ergebnis]} – "
                   + kern.kaskade_anwenden(session, vorgang, benutzer))
    elif ergebnis == "falsche_nummer":
        meldung = "Falsche Nummer – Lead bleibt mit Kennzeichen „Nummer prüfen“."
    elif ergebnis == "kein_interesse":
        grund = (form.get("grund") or "").strip()
        text = (form.get("grund_text") or "").strip()
        fehler = _grund_pruefen(session, "unqualifiziert", grund, text)
        if fehler:
            return RedirectResponse("/lead-management/anrufliste?meldung="
                                    + quote_plus(fehler), status_code=303)
        vorgang.lead_phase = "unqualifiziert"
        vorgang.unqualifiziert_grund = grund
        vorgang.unqualifiziert_text = text
        vorgang.naechste_aktion_am = None
        kern.aktivitaet(session, vorgang.id, "status",
                        f"Unqualifiziert: {grund}"
                        + (f" – {text}" if text else ""), benutzer=benutzer)
        meldung = f"Kein Interesse – unqualifiziert ({grund})."
    session.commit()
    return RedirectResponse("/lead-management/anrufliste?meldung="
                            + quote_plus(meldung), status_code=303)


def _grund_pruefen(session: Session, phase: str, grund: str,
                   text: str) -> str | None:
    from app import leadmanagement_logik
    logik = leadmanagement_logik.hole_logik()
    passend = next((g for g in logik.gruende_der_phase(phase)
                    if g.grund == grund), None)
    if passend is None:
        return "Bitte einen Grund aus der Liste wählen."
    if passend.freitext_pflicht and not text:
        return f"Beim Grund „{grund}“ ist der Freitext Pflicht."
    return None


@router.post("/lead/{vorgang_id}/zurueckstellen")
async def zurueckstellen(request: Request, vorgang_id: int,
                         session: Session = Depends(get_session)):
    from datetime import datetime as dt
    from urllib.parse import quote_plus
    _gate(request, session)
    form = await request.form()
    benutzer = request.state.benutzer
    vorgang = session.get(Vorgang, vorgang_id)
    if vorgang is None:
        return RedirectResponse("/lead-management/anrufliste", status_code=303)
    try:
        bis = dt.strptime((form.get("bis") or "").strip(), "%Y-%m-%d")
    except ValueError:
        return RedirectResponse("/lead-management/anrufliste?meldung="
                                + quote_plus("Zurückstellen: Datum ist Pflicht."),
                                status_code=303)
    grund = (form.get("grund") or "").strip()
    text = (form.get("grund_text") or "").strip()
    fehler = _grund_pruefen(session, "zurueckgestellt", grund, text)
    if fehler:
        return RedirectResponse("/lead-management/anrufliste?meldung="
                                + quote_plus(fehler), status_code=303)
    vorgang.lead_phase = "zurueckgestellt"
    vorgang.zurueckgestellt_bis = bis
    vorgang.zurueckgestellt_grund = grund + (f" – {text}" if text else "")
    vorgang.naechste_aktion_am = None
    kern.aktivitaet(session, vorgang.id, "status",
                    f"Zurückgestellt bis {bis.strftime('%d.%m.%Y')}: "
                    f"{vorgang.zurueckgestellt_grund}", benutzer=benutzer)
    session.commit()
    zurueck = request.headers.get("referer", "/lead-management/anrufliste")
    return RedirectResponse(zurueck if zurueck.startswith(("/", str(request.base_url)))
                            else "/lead-management/anrufliste", status_code=303)


@router.post("/lead/{vorgang_id}/unqualifiziert")
async def unqualifiziert(request: Request, vorgang_id: int,
                         session: Session = Depends(get_session)):
    from urllib.parse import quote_plus
    _gate(request, session)
    form = await request.form()
    benutzer = request.state.benutzer
    vorgang = session.get(Vorgang, vorgang_id)
    if vorgang is None:
        return RedirectResponse("/lead-management/anrufliste", status_code=303)
    grund = (form.get("grund") or "").strip()
    text = (form.get("grund_text") or "").strip()
    fehler = _grund_pruefen(session, "unqualifiziert", grund, text)
    if fehler:
        return RedirectResponse("/lead-management/anrufliste?meldung="
                                + quote_plus(fehler), status_code=303)
    vorgang.lead_phase = "unqualifiziert"
    vorgang.unqualifiziert_grund = grund
    vorgang.unqualifiziert_text = text
    vorgang.naechste_aktion_am = None
    kern.aktivitaet(session, vorgang.id, "status",
                    f"Unqualifiziert: {grund}" + (f" – {text}" if text else ""),
                    benutzer=benutzer)
    session.commit()
    return RedirectResponse("/lead-management/anrufliste?meldung="
                            + quote_plus(f"Unqualifiziert ({grund})."),
                            status_code=303)


@router.post("/lead/{vorgang_id}/reaktivieren")
async def reaktivieren(request: Request, vorgang_id: int,
                       session: Session = Depends(get_session)):
    from urllib.parse import quote_plus
    _gate(request, session)
    form = await request.form()
    benutzer = request.state.benutzer
    vorgang = session.get(Vorgang, vorgang_id)
    if vorgang is None:
        return RedirectResponse("/lead-management/anrufliste", status_code=303)
    begruendung = (form.get("begruendung") or "").strip()
    if not begruendung:
        return RedirectResponse("/lead-management/anrufliste?meldung="
                                + quote_plus("Reaktivieren: Begründung ist "
                                             "Pflicht."), status_code=303)
    alte_phase = vorgang.lead_phase
    ziel = "in_kontaktierung" if alte_phase == "nicht_erreicht" else "neu"
    vorgang.lead_phase = ziel
    vorgang.zurueckgestellt_bis = None
    vorgang.unqualifiziert_grund = None
    vorgang.unqualifiziert_text = None
    vorgang.naechste_aktion_am = datetime.now()
    kern.aktivitaet(session, vorgang.id, "status",
                    f"Reaktiviert ({LEAD_PHASEN_NAMEN.get(alte_phase, alte_phase)}"
                    f" → {LEAD_PHASEN_NAMEN[ziel]}): {begruendung}",
                    benutzer=benutzer)
    session.commit()
    return RedirectResponse("/lead-management/anrufliste?meldung="
                            + quote_plus("Reaktiviert."), status_code=303)


# --- Qualifizierungsbogen -------------------------------------------------------------

@router.get("/lead/{vorgang_id}/qualifizierung/{sparte}")
async def qualifizierung_bogen(request: Request, vorgang_id: int, sparte: str,
                               session: Session = Depends(get_session)):
    import json as json_modul
    _gate(request, session)
    from app import leadmanagement_logik
    from app.models import LeadQualifizierung
    vorgang = session.get(Vorgang, vorgang_id)
    if vorgang is None or sparte not in SPARTEN:
        return RedirectResponse("/lead-management/anrufliste", status_code=303)
    kunde = session.get(Kunde, vorgang.kunde_id)
    logik = leadmanagement_logik.hole_logik()
    fragen = logik.fragen_der_sparte(sparte)
    interessen = [s.strip() for s in (kunde.interesse or "").split(",")
                  if s.strip() in SPARTEN] or [sparte]
    if sparte not in interessen:
        interessen.append(sparte)
    zeile = (session.query(LeadQualifizierung)
             .filter_by(vorgang_id=vorgang.id, sparte=sparte).first())
    try:
        antworten = json_modul.loads(zeile.antworten) if zeile else {}
    except ValueError:
        antworten = {}
    fertig = {q.sparte for q in session.query(LeadQualifizierung)
              .filter(LeadQualifizierung.vorgang_id == vorgang.id,
                      LeadQualifizierung.abgeschlossen_am.isnot(None))}
    scoring = [{"frage_key": r.frage_key, "bedingung": r.bedingung,
                "punkte": r.punkte} for r in logik.scoring]
    return render(request, "leadmanagement/qualifizierung.html",
                  aktiv="/lead-management", vorgang=vorgang, kunde=kunde,
                  sparte=sparte, fragen=fragen, antworten=antworten,
                  interessen=interessen, fertig=fertig,
                  basis_punkte=(vorgang.score_punkte or 0),
                  scoring_json=json_modul.dumps(scoring, ensure_ascii=False),
                  klassen_json=json_modul.dumps(logik.klassen),
                  antworten_json=json_modul.dumps(antworten, ensure_ascii=False),
                  meldung=request.query_params.get("meldung", ""))


@router.post("/lead/{vorgang_id}/qualifizierung/{sparte}")
async def qualifizierung_speichern(request: Request, vorgang_id: int,
                                   sparte: str,
                                   session: Session = Depends(get_session)):
    from urllib.parse import quote_plus

    from app import leadmanagement_logik
    _gate(request, session)
    vorgang = session.get(Vorgang, vorgang_id)
    if vorgang is None or sparte not in SPARTEN:
        return RedirectResponse("/lead-management/anrufliste", status_code=303)
    form = await request.form()
    benutzer = request.state.benutzer
    logik = leadmanagement_logik.hole_logik()
    antworten: dict = {}
    fehlend = []
    for frage in logik.fragen_der_sparte(sparte):
        if frage.typ == "mehrfach":
            werte = form.getlist(f"f_{frage.key}")
            if werte:
                antworten[frage.key] = werte
            elif frage.pflicht:
                fehlend.append(frage.frage)
        else:
            wert = (form.get(f"f_{frage.key}") or "").strip()
            if wert:
                antworten[frage.key] = wert
            elif frage.pflicht:
                fehlend.append(frage.frage)
    if form.get("aktion") == "spaeter":
        # Zwischenstand ohne Abschluss speichern
        from app.models import LeadQualifizierung
        zeile = (session.query(LeadQualifizierung)
                 .filter_by(vorgang_id=vorgang.id, sparte=sparte).first())
        if zeile is None:
            zeile = LeadQualifizierung(vorgang_id=vorgang.id, sparte=sparte)
            session.add(zeile)
        import json as json_modul
        zeile.antworten = json_modul.dumps(antworten, ensure_ascii=False)
        session.commit()
        return RedirectResponse("/lead-management/anrufliste?meldung="
                                + quote_plus("Zwischenstand gespeichert."),
                                status_code=303)
    if fehlend:
        return RedirectResponse(
            f"/lead-management/lead/{vorgang.id}/qualifizierung/{sparte}"
            f"?meldung=" + quote_plus("Pflichtfragen offen: "
                                      + " · ".join(fehlend[:3])),
            status_code=303)
    zeile = kern.qualifizierung_abschliessen(session, vorgang, sparte,
                                             antworten, benutzer=benutzer)
    session.commit()
    kunde = session.get(Kunde, vorgang.kunde_id)
    offene = [s.strip() for s in (kunde.interesse or "").split(",")
              if s.strip() in SPARTEN and s.strip() != sparte
              and not (session.query(type(zeile))
                       .filter_by(vorgang_id=vorgang.id, sparte=s.strip())
                       .filter(type(zeile).abgeschlossen_am.isnot(None)).count())]
    return render(request, "leadmanagement/qualifizierung_fertig.html",
                  aktiv="/lead-management", vorgang=vorgang, kunde=kunde,
                  sparte=sparte, punkte=zeile.score_punkte,
                  klasse=zeile.score_klasse, offene_sparten=offene,
                  meldung="")
