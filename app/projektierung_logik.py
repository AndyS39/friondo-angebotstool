# Projektierungs-Steuerdatei (v11, Phase 65): liest projektierung_logik_v1.xlsx
# (Blätter Aufgabenpakete, Paketregeln, Ordnerstruktur, Sub-Typen) – bewusst
# getrennt vom Konfigurator, damit beide unabhängig gepflegt werden.
# Änderungen an der Datei wirken nur auf NEUE Paket-Aktivierungen; bestehende
# Aufgaben-Instanzen bleiben unverändert (deshalb kein DB-Abbild der Pakete).

from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from app import config

LOGIK_PFAD = Path(getattr(config, "PROJEKTIERUNG_LOGIK_PFAD",
                          Path(config.LOGIK_EXCEL_PFAD).parent
                          / "projektierung_logik_v1.xlsx"))

SPARTEN = ("WP", "PV", "KL", "WB")
ROLLEN = ("projektierer", "elektroplaner", "feinplaner", "innendienst",
          "buchhaltung", "montage")


@dataclass
class PaketSchritt:
    nr: int
    titel: str
    beschreibung: str
    rolle: str
    pflicht: bool
    faellig_regel: str
    wartet_frist_tage: int | None


@dataclass
class Paket:
    key: str
    name: str
    sparte: str                       # WP/PV/KL/WB/ALLE
    schritte: list[PaketSchritt] = field(default_factory=list)


@dataclass
class PaketRegel:
    sparte: str
    bedingung: str                    # IMMER oder <frage_key>=<wert> (ab V2)
    paket_key: str


@dataclass
class ProjektierungsLogik:
    pakete: dict[str, Paket] = field(default_factory=dict)
    regeln: list[PaketRegel] = field(default_factory=list)
    ordner: list[dict] = field(default_factory=list)
    sub_typen: list[str] = field(default_factory=list)
    fehler: list[str] = field(default_factory=list)
    warnungen: list[str] = field(default_factory=list)
    stand: str = ""

    def pakete_fuer_sparte(self, sparte: str) -> list[Paket]:
        return [p for p in self.pakete.values()
                if p.sparte in ("ALLE", sparte)]

    def immer_pakete(self, sparte: str) -> list[Paket]:
        """Pakete mit IMMER-Regel für diese Sparte (inkl. ALLE) – V1 wertet
        nur IMMER aus, Frage-Regeln (FP-…) folgen in V2."""
        ergebnis = []
        gesehen: set[str] = set()
        for regel in self.regeln:
            if regel.bedingung.strip().upper() != "IMMER":
                continue
            if regel.sparte not in ("ALLE", sparte):
                continue
            paket = self.pakete.get(regel.paket_key)
            if paket is None or paket.key in gesehen:
                continue
            if paket.sparte not in ("ALLE", sparte):
                continue
            gesehen.add(paket.key)
            ergebnis.append(paket)
        return ergebnis


def _text(wert) -> str:
    return str(wert).strip() if wert is not None else ""


def einlesen(pfad: Path | None = None) -> ProjektierungsLogik:
    pfad = Path(pfad or LOGIK_PFAD)
    logik = ProjektierungsLogik()
    if not pfad.exists():
        logik.fehler.append(f"Steuerdatei fehlt: {pfad}")
        return logik
    try:
        wb = openpyxl.load_workbook(pfad, data_only=True)
    except Exception as problem:
        logik.fehler.append(f"Steuerdatei nicht lesbar: {problem}")
        return logik

    # --- Aufgabenpakete ---
    if "Aufgabenpakete" not in wb.sheetnames:
        logik.fehler.append("Blatt „Aufgabenpakete“ fehlt.")
    else:
        for zeile in wb["Aufgabenpakete"].iter_rows(min_row=2, values_only=True):
            key = _text(zeile[0])
            if not key:
                continue
            name = _text(zeile[1]) or key
            sparte = (_text(zeile[2]) or "ALLE").upper()
            if sparte not in SPARTEN + ("ALLE",):
                logik.warnungen.append(f"Paket {key}: unbekannte Sparte „{sparte}“.")
            try:
                nr = int(float(zeile[3])) if zeile[3] is not None else 0
            except (TypeError, ValueError):
                nr = 0
            rolle = (_text(zeile[6]) or "projektierer").lower()
            if rolle not in ROLLEN:
                logik.warnungen.append(
                    f"Paket {key} Schritt {nr}: unbekannte Rolle „{rolle}“ – "
                    "Verantwortlicher wird der Projektleiter.")
            wartet = None
            if zeile[9] is not None and _text(zeile[9]):
                try:
                    wartet = int(float(zeile[9]))
                except (TypeError, ValueError):
                    logik.warnungen.append(
                        f"Paket {key} Schritt {nr}: wartet_frist_tage "
                        f"„{zeile[9]}“ nicht lesbar.")
            paket = logik.pakete.setdefault(key, Paket(key=key, name=name,
                                                       sparte=sparte))
            paket.schritte.append(PaketSchritt(
                nr=nr, titel=_text(zeile[4]), beschreibung=_text(zeile[5]),
                rolle=rolle if rolle in ROLLEN else "projektierer",
                pflicht=_text(zeile[7]).upper() in ("J", "JA", "X", "1"),
                faellig_regel=_text(zeile[8]).upper().replace(" ", ""),
                wartet_frist_tage=wartet))
        for paket in logik.pakete.values():
            paket.schritte.sort(key=lambda s: s.nr)

    # --- Paketregeln ---
    if "Paketregeln" in wb.sheetnames:
        for zeile in wb["Paketregeln"].iter_rows(min_row=2, values_only=True):
            paket_key = _text(zeile[2])
            if not paket_key:
                continue
            if paket_key not in logik.pakete:
                logik.warnungen.append(
                    f"Paketregel verweist auf unbekanntes Paket „{paket_key}“.")
                continue
            logik.regeln.append(PaketRegel(
                sparte=(_text(zeile[0]) or "ALLE").upper(),
                bedingung=_text(zeile[1]) or "IMMER",
                paket_key=paket_key))
    else:
        logik.fehler.append("Blatt „Paketregeln“ fehlt.")

    # --- Ordnerstruktur ---
    if "Ordnerstruktur" in wb.sheetnames:
        for zeile in wb["Ordnerstruktur"].iter_rows(min_row=2, values_only=True):
            pfad_text = _text(zeile[0])
            if pfad_text:
                logik.ordner.append({"pfad": pfad_text,
                                     "ebene": (_text(zeile[1]) or "projekt").lower()})

    # --- Sub-Typen ---
    if "Sub-Typen" in wb.sheetnames:
        for zeile in wb["Sub-Typen"].iter_rows(min_row=2, values_only=True):
            typ = _text(zeile[0])
            if typ:
                logik.sub_typen.append(typ)

    import datetime
    logik.stand = (f"{datetime.datetime.fromtimestamp(pfad.stat().st_mtime):%d.%m.%Y %H:%M}"
                   f" · {len(logik.pakete)} Pakete · {len(logik.regeln)} Regeln")
    return logik


_cache: dict = {"mtime": None, "logik": None}


def hole_logik(session=None, erzwingen: bool = False) -> ProjektierungsLogik:
    """Gecachter Zugriff (mtime-basiert); session-Parameter für Symmetrie zum
    Konfigurator-Muster. Beim Neu-Einlesen werden Ordnervorlage und Sub-Typen
    in die projektierung_parameter übernommen (wirken auf neue Anlagen)."""
    mtime = LOGIK_PFAD.stat().st_mtime if LOGIK_PFAD.exists() else None
    if not erzwingen and _cache["logik"] is not None and _cache["mtime"] == mtime:
        return _cache["logik"]
    logik = einlesen()
    _cache["mtime"] = mtime
    _cache["logik"] = logik
    if session is not None and not logik.fehler:
        _parameter_uebernehmen(session, logik)
    return logik


def _parameter_uebernehmen(session, logik: ProjektierungsLogik) -> None:
    import json

    from app import projektierung
    if logik.ordner:
        projektierung.parameter_setzen(
            session, "ordnervorlage", json.dumps(logik.ordner, ensure_ascii=False))
    if logik.sub_typen:
        projektierung.parameter_setzen(
            session, "sub_typen", json.dumps(logik.sub_typen, ensure_ascii=False))


def sub_typen(session) -> list[str]:
    import json

    from app import projektierung
    roh = projektierung.parameter_holen(session, "sub_typen", "")
    if roh:
        try:
            return [str(t) for t in json.loads(roh)]
        except (ValueError, TypeError):
            pass
    from app.projektierung import SUB_TYPEN_STANDARD
    return list(SUB_TYPEN_STANDARD)
