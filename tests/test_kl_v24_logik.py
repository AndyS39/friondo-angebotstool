# Tests PLAN_V16 Phase 113 – Lader und Parametrierung des Klimakonfigurators
# (Agent B): app/logik.py (_kl_einlesen/_kl_pruefen, refs_extrahieren mit KL,
# logik_fuer_sparte(…, "KL"), artikel_referenzen/artikel_pruefen), Routen
# /parametrierung (KL-Parameter, Import-Button), /parametrierung/kl-logik und
# /parametrierung/artikel/kl-import. Laufen gegen die Entwicklungs-DB; der
# KL-Import wird bei Bedarf ausgeführt (idempotent, Muster tests/test_pv_v13.py).
import copy
import re
import unittest
import warnings
from datetime import datetime
from types import SimpleNamespace

warnings.filterwarnings("ignore")

import openpyxl
from fastapi.testclient import TestClient

from app import config, models
from app import logik as logik_modul
from app.logik import (Aktion, ArtikelRef, Logik, Pruefbericht, bedingung_parsen,
                       refs_extrahieren)
from app.db import SessionLocal, init_db
from app.main import app
from app.models import Artikel

QUELLE_KL = getattr(models, "QUELLE_KL", "kl")
GUID_SYSTEMGARANTIE = "{38FC5932-9389-46DE-A9D0-8855676E1FD8}"
RAUMFRAGEN = ["KR01", "KR08", "KR09", "KR10", "KR02", "KR03", "KR04", "KR05", "KR06",
              "KR11", "KR07"]


def kl_import_sicherstellen(session) -> None:
    aktiv = (session.query(Artikel).filter(Artikel.quelle == QUELLE_KL,
                                           Artikel.aktiv.is_(True)).count())
    if aktiv < 50:
        from app import import_klima
        import_klima.import_ausfuehren(session)


class Basis(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        init_db()
        cls.s = SessionLocal()
        kl_import_sicherstellen(cls.s)
        cls.logik, cls.bericht = logik_modul.neu_einlesen(cls.s)
        cls.kl = logik_modul.logik_fuer_sparte(cls.logik, "KL")

    @classmethod
    def tearDownClass(cls):
        cls.s.close()


# --- Lader ---------------------------------------------------------------------------

class Lader(Basis):
    def test_validierung_gruen(self):
        self.assertTrue(self.bericht.ok, self.bericht.fehler)
        # keine neuen Hinweise durch die KL-Blätter (Entwurf nennt nur Optionen mit Wirkung)
        self.assertFalse([w for w in self.bericht.warnungen if "KL" in w or "Klima" in w],
                         self.bericht.warnungen)
        self.assertEqual(logik_modul.kl_parameter_fehlende(self.logik), [])

    def test_fragen_kl(self):
        fragen = self.logik.sparten_fragen["KL"]
        self.assertEqual(len(fragen), 28)
        reihenfolge = [f.id for f in sorted(fragen.values(), key=lambda f: f.reihenfolge)]
        self.assertEqual([f for f in reihenfolge if f.startswith("KR")], RAUMFRAGEN)
        self.assertLess(reihenfolge.index("KO06"), reihenfolge.index("KO04"))
        # „nur wenn KE01 = Nein oder Unklar“ wird als Antwort-Bedingung mit zwei Werten gelesen
        for fid in ("KE02", "KE03"):
            b = fragen[fid].bedingung
            self.assertEqual((b.art, b.frage_id, b.werte), ("antwort", "KE01", ["Nein", "Unklar"]))
        self.assertEqual(fragen["KR08"].bedingung.art, "wiederholgruppe")
        self.assertEqual(fragen["KR08"].bedingung.frage_id, "KO05")
        self.assertIn("Optionen wie Außengeräte (KO04)", fragen["KR07"].hinweis)
        self.assertEqual(fragen["KO08"].antworten, ["Ja", "Nein", "als Eventualposition"])
        self.assertEqual(fragen["KR05"].antworten, ["bis 5 m", "6–10 m", "11–15 m", "über 15 m"])

    def test_zahlen_der_blaetter(self):
        l = self.logik
        self.assertEqual(len(l.kl_aktionen), 45)
        self.assertGreaterEqual(len(l.kl_paket), 20)
        self.assertEqual(len(l.kl_bloecke), 3)
        self.assertEqual(len(l.kl_parameter), 17)
        self.assertEqual(len(l.kl_artikel_nummern), 50)
        self.assertEqual(len(l.kl_artikel), 49)                  # KL050 ohne GUID
        self.assertEqual(l.kl_artikel[GUID_SYSTEMGARANTIE], "KL049")
        self.assertIn("KL050", l.kl_artikel_nummern)
        # Kombinationen: 215 Zeilen, Codes aufsteigend normalisiert
        self.assertIn("9+9+9", l.kl_kombis["CL5000M 79/3 E"][3])
        self.assertIn("12+18", l.kl_kombis["CL5000M 53/2 E"][2])
        self.assertNotIn(3, l.kl_kombis["CL5000M 53/2 E"])
        gesamt = sum(len(s) for je in l.kl_kombis.values() for s in je.values())
        self.assertEqual(gesamt, 215)
        for je_anzahl in l.kl_kombis.values():
            for anzahl, kombis in je_anzahl.items():
                for k in kombis:
                    codes = k.split("+")
                    self.assertEqual(len(codes), anzahl)
                    self.assertEqual(codes, sorted(codes, key=int))
                    self.assertTrue(set(codes) <= set(logik_modul.KL_CODES))
        self.assertEqual(l.kl_kombi_artikel["CL5000M 79/3 E"], "KL038")
        self.assertEqual(l.kl_kombi_artikel["CL7000M 53/2 E"], "KL045")
        # Montagematrix
        self.assertEqual(l.kl_montage[(3, 1)], "KL023")
        self.assertEqual(l.kl_montage, {(1, 1): "KL020", (2, 1): "KL021", (2, 2): "KL022",
                                        (3, 1): "KL023", (3, 2): "KL024", (4, 1): "KL025"})
        self.assertIn("Montagekombination <Innengeräte> Innengeräte", l.kl_montage_ampel)
        # Parameter mit Einheit
        self.assertEqual(logik_modul._kl_zahl(l.kl_parameter["Klasse 24 bis"][0]), 7.0)
        self.assertEqual(l.kl_parameter_einheit["Klasse 24 bis"], "kW")
        self.assertEqual(l.kl_parameter["Standardserie"][0], "Climate 3200i (Standard)")
        self.assertEqual(l.kl_parameter["Meterposition Zusatzleitung"][0], "KL017")
        self.assertEqual(l.kl_parameter["§14a-Außengeräte"][0], "CL5000M 105/4 E; CL5000M 125/5 E")

    def test_paketmatrix(self):
        zeilen = {(z.serie, z.typ, z.klasse): z for z in self.logik.kl_paket}
        self.assertEqual(zeilen[("Climate 3200i (Standard)", "single", "9")].artikel, ["KL034"])
        self.assertEqual(zeilen[("Climate 3200i (Standard)", "single", "24")].artikel,
                         ["KL033", "KL029"])
        self.assertEqual(zeilen[("Climate 3200i (Standard)", "multi_innen", "9")].artikel, ["KL026"])
        self.assertEqual(zeilen[("Climate 3200i (Standard)", "multi_aussen", "79/3")].artikel,
                         ["KL038"])
        self.assertEqual(zeilen[("Climate 7000i", "single", "12")].artikel, ["KL044", "KL042"])
        nicht = zeilen[("Climate 7000i", "single", "18 / 24")]
        self.assertTrue(nicht.nicht_im_sortiment)
        self.assertEqual(nicht.artikel, [])
        self.assertIn("nicht verfügbar", nicht.ampel_grund)
        multi8000 = zeilen[("Climate 8000i", "multi_aussen", "–")]
        self.assertTrue(multi8000.nicht_im_sortiment)
        self.assertEqual(multi8000.ampel_grund,
                         "Multi-Split in Serie Climate 8000i nicht verfügbar")
        self.assertEqual(multi8000.typ_roh, "Multi")
        self.assertEqual({z.typ for z in self.logik.kl_paket},
                         {"single", "multi_innen", "multi_aussen"})
        # Multi-Außengeräte in Blattreihenfolge klein → groß
        aussen = [z.artikel[0] for z in self.logik.kl_paket
                  if z.serie.startswith("Climate 3200i") and z.typ == "multi_aussen"]
        self.assertEqual(aussen, ["KL037", "KL038", "KL039", "KL040"])

    def test_aktionen(self):
        nach = {}
        for a in self.logik.kl_aktionen:
            nach.setdefault((a.frage, a.antwort), []).append(a)
        refs = lambda a: [(r.ref, r.menge, r.ep) for r in a.artikel]   # noqa: E731
        self.assertEqual(refs(nach[("Grundpaket", "immer")][0]), [("KL049", "1", False)])
        self.assertEqual(refs(nach[("KO08", "Ja")][0]), [("KL013", "Innengeräte", False)])
        self.assertEqual(refs(nach[("KO08", "als Eventualposition")][0]),
                         [("KL013", "Innengeräte", True)])
        self.assertEqual(refs(nach[("KE01", "Ja")][0]), [("KL001", "Außengeräte", False)])
        self.assertEqual(refs(nach[("KA01", "auf einem Schrägdach")][0]),
                         [("KL004", "1", False), ("KL005", "1", False), ("KL006", "1", False)])
        self.assertEqual(refs(nach[("KR05", "6–10 m")][0]), [("KL017", "5", False)])
        self.assertEqual(refs(nach[("KR05", "11–15 m")][0]), [("KL017", "10", False)])
        self.assertEqual(refs(nach[("KR11", "Beton")][0]), [("KL014", "1", False)])
        self.assertEqual(refs(nach[("KZ01", "Gerüst erforderlich")][0]), [("KL050", "1", False)])
        self.assertEqual(refs(nach[("§14a", "Außengerät in Liste „§14a-Außengeräte“ (KL-Parameter)")][0]),
                         [("KL003", "1", False)])
        # Dokumentationszeilen ohne Referenzen
        self.assertEqual(nach[("Händisch", "nie automatisch")][0].artikel, [])
        for a in nach[("Grundpaket", "immer")][1:]:
            self.assertEqual(a.artikel, [])
        # Hinweise: typ „hinweis“, Grund = Text, keine Position
        hinweise = {a.frage: a for a in self.logik.kl_aktionen if a.typ == "hinweis"}
        self.assertEqual(set(hinweise), {"KO03", "KO01", "KR04", "KR03", "KA01", "KA02", "KE01"})
        for a in hinweise.values():
            self.assertEqual(a.artikel, [])
        self.assertEqual(hinweise["KO03"].ampel_grund, "Zustimmung des Eigentümers einholen")
        self.assertEqual(hinweise["KO01"].ampel_grund, "Gewerbeobjekt – Auslegung prüfen")
        self.assertEqual(hinweise["KE01"].ampel_grund, "Stromversorgung am Außengerät prüfen")
        self.assertEqual(hinweise["KA02"].ampel_grund, "Montagehöhe über 4 m ohne Gerüst/Bühne prüfen")
        # Ampeln: Gründe wörtlich (G7/G8), Zusatzbedingungen geparst
        ampeln = {(a.frage, a.antwort): a for a in self.logik.kl_aktionen if a.typ == "ampel"}
        self.assertEqual(ampeln[("KE02", "über 15 m")].ampel_grund, "Elektrozuleitung über 15 m")
        self.assertEqual(ampeln[("KR05", "über 15 m")].ampel_grund,
                         "Leitungslänge über 15 m: Raum <Nr> <Name> – max. Leitungslänge der Serie prüfen")
        b = ampeln[("KE02", "über 15 m")].zusatz_bedingung
        self.assertEqual((b.art, b.frage_id, b.werte), ("antwort", "KE01", ["Nein", "Unklar"]))
        b = nach[("KZ01", "Gerüst erforderlich")][0].zusatz_bedingung
        self.assertEqual(b.frage_id, "KA01")
        self.assertEqual(len(b.werte), 6)
        self.assertNotIn("auf einem Schrägdach", b.werte)
        b = nach[("KA02", "über 4 m")][0].zusatz_bedingung
        self.assertEqual((b.frage_id, b.werte), ("KZ01", ["vom Boden/Leiter"]))
        self.assertFalse([a for a in self.logik.kl_aktionen
                          if a.frage in ("KA02", "KR04", "KR03") and a.typ == "ampel"])

    def test_bloecke(self):
        bloecke = {b.nr: b for b in self.logik.kl_bloecke}
        self.assertEqual(bloecke[3].ueberschrift, "Klimaanlage Bosch")
        self.assertEqual(bloecke[1].ueberschrift, "(ohne Überschrift)")
        refs3 = [r.ref for r in bloecke[3].refs]
        self.assertEqual(refs3[0], "KL034")
        self.assertEqual(refs3[-1], "KL049")
        self.assertLess(refs3.index("KL038"), refs3.index("KL026"))
        self.assertLess(refs3.index("KL026"), refs3.index("KL013"))
        self.assertEqual([r.ref for r in bloecke[1].refs][:6],
                         ["KL020", "KL021", "KL022", "KL023", "KL024", "KL025"])
        self.assertEqual([r.ref for r in bloecke[2].refs], ["KL001", "KL008", "KL009", "KL003"])
        for b in bloecke.values():
            self.assertEqual(b.wann.art, "immer")

    def test_refs_extrahieren_kl(self):
        def r(text):
            return [(x.ref, x.menge, x.ep) for x in refs_extrahieren(text)]
        self.assertEqual(r("Artikel: KL017 × 5"), [("KL017", "5", False)])
        self.assertEqual(r("Artikel: KL013 × Innengeräte"), [("KL013", "Innengeräte", False)])
        self.assertEqual(r("Artikel: KL013 × Innengeräte als EP"), [("KL013", "Innengeräte", True)])
        self.assertEqual(r("Artikel: KL013 (EP) × Innengeräte"), [("KL013", "Innengeräte", True)])
        self.assertEqual(r("Artikel: KL001 × Außengeräte"), [("KL001", "Außengeräte", False)])
        self.assertEqual(r("Artikel: KL004 ×1 · KL005 ×1 · KL006 ×1 je Außengerät"),
                         [("KL004", "1", False), ("KL005", "1", False), ("KL006", "1", False)])
        self.assertEqual(r("Artikel: KL003 ×1 je betroffenem Außengerät"), [("KL003", "1", False)])
        self.assertEqual(r("Artikel: KL016 ×1 je Raum"), [("KL016", "1", False)])
        self.assertEqual(r("KL033 (CL3000i 70 E Außen) + KL029 (CL3200iU W 70 E Innen)"),
                         [("KL033", "1", False), ("KL029", "1", False)])
        self.assertEqual(r("–"), [])
        # WP/PV unverändert
        self.assertEqual(r("Pos. 045 ×1 · PV013 × Modulanzahl"),
                         [("045", "1", False), ("PV013", "Modulanzahl", False)])

    def test_sicht_kl(self):
        kl = self.kl
        self.assertEqual(kl.sparte, "KL")
        self.assertIs(kl.aktionen, self.logik.kl_aktionen)
        self.assertIs(kl.bloecke, self.logik.kl_bloecke)
        self.assertIs(kl.kl_paket, self.logik.kl_paket)
        self.assertEqual(kl.kl_montage[(3, 1)], "KL023")
        self.assertEqual(kl.kl_kombi_artikel["CL5000M 79/3 E"], "KL038")
        self.assertEqual(kl.anhaenge, self.logik.anhaenge)
        self.assertEqual(kl.kfw, self.logik.kfw)
        self.assertEqual(kl.seiten[0], "Objekt & Anlage")
        self.assertEqual(kl.seiten[-1], "Einschätzung")
        self.assertIn("KR08", kl.fragen)
        self.assertEqual(kl.pakete, [])
        # WP/PV-Sichten unverändert
        self.assertIs(logik_modul.logik_fuer_sparte(self.logik, "WP"), self.logik)
        pv = logik_modul.logik_fuer_sparte(self.logik, "PV")
        self.assertEqual(pv.sparte, "PV")
        self.assertEqual(pv.kl_aktionen, [])
        self.assertIsNone(logik_modul.logik_fuer_sparte(self.logik, "WB"))

    def test_artikel_referenzen_und_pruefung(self):
        refs = logik_modul.artikel_referenzen(self.logik)
        kl_refs = sorted(r for r in refs if r.startswith("KL"))
        self.assertEqual(len(kl_refs), 41)
        self.assertTrue(set(kl_refs) <= set(self.logik.kl_artikel_nummern))
        for nr in ("KL002", "KL007", "KL010", "KL011", "KL012", "KL018", "KL030", "KL031", "KL032"):
            self.assertNotIn(nr, refs)                 # nur händisch / Editor-Artikel
        self.assertIn("Kombinationen KL CL5000M 79/3 E", refs["KL038"])
        self.assertIn("Montagematrix KL 3/1", refs["KL023"])
        self.assertIn("KL-Parameter „Meterposition Zusatzleitung“", refs["KL017"])
        self.assertTrue(any(q.startswith("Aktionen KL KO08") for q in refs["KL013"]))
        self.assertTrue(any(r.startswith("PV") for r in refs))   # PV-Referenzen bleiben
        # fehlender KL-Artikel im Stamm → Fehler mit Import-Hinweis
        kopie = copy.copy(self.logik)
        kopie.kl_aktionen = list(self.logik.kl_aktionen) + [
            Aktion("KO07", "Ja", "Artikel: KL099 ×1", "normal", "",
                   [ArtikelRef("KL099", "1", False)], "")]
        bericht = Pruefbericht()
        logik_modul.artikel_pruefen(kopie, self.s, bericht)
        treffer = [f for f in bericht.fehler if "KL099" in f]
        self.assertEqual(len(treffer), 1)
        self.assertIn("Artikel KL099 fehlt im Artikelstamm", treffer[0])
        self.assertIn("bitte Artikel → Klima-Positionslisten importieren", treffer[0])

    def test_anhaenge_regel_sparte_kl(self):
        # Platzhalterzeile „(Bosch Climate Broschüre – Zulieferung)“ wird übersprungen …
        self.assertFalse([a for a in self.logik.anhaenge if a.datei.startswith("(")])
        # … mit Dateinamen greift „wenn Sparte = KL“ für KL-Angebote (konfigurator_typ)
        wb = openpyxl.load_workbook(config.LOGIK_EXCEL_PFAD, data_only=True)
        ws = wb["Anhänge"]
        zeile = next(r for r in range(2, ws.max_row + 1)
                     if str(ws.cell(r, 2).value or "").strip() == "wenn Sparte = KL")
        ws.cell(zeile, 1).value = "Bosch Climate Broschüre.pdf"
        bericht = Pruefbericht()
        anhaenge = logik_modul._anhaenge_einlesen(wb, bericht)
        eintrag = next(a for a in anhaenge if a.datei == "Bosch Climate Broschüre.pdf")
        self.assertEqual((eintrag.art, eintrag.antwort), ("sparte", "KL"))
        self.assertFalse([w for w in bericht.warnungen if "Sparte = KL" in w])
        from app import anhaenge as anhaenge_modul
        kopie = copy.copy(self.kl)
        kopie.anhaenge = anhaenge
        kl_angebot = SimpleNamespace(konfigurator_typ="KL", protokoll_json="[]", positionen=[])
        wp_angebot = SimpleNamespace(konfigurator_typ="WP", protokoll_json="[]", positionen=[])
        self.assertIn("Bosch Climate Broschüre.pdf",
                      [a.datei for a in anhaenge_modul.fuer_angebot(kopie, kl_angebot)])
        self.assertNotIn("Bosch Climate Broschüre.pdf",
                         [a.datei for a in anhaenge_modul.fuer_angebot(kopie, wp_angebot)])

    def test_bedingung_ke01(self):
        b = bedingung_parsen("nur wenn KE01 = Nein oder Unklar")
        self.assertEqual((b.art, b.frage_id, b.werte), ("antwort", "KE01", ["Nein", "Unklar"]))
        from app import konfigurator as engine
        fragen = self.kl.fragen
        self.assertTrue(engine.ist_sichtbar(fragen["KE02"], {"KE01": "Nein"}, fragen))
        self.assertTrue(engine.ist_sichtbar(fragen["KE02"], {"KE01": "Unklar"}, fragen))
        self.assertFalse(engine.ist_sichtbar(fragen["KE02"], {"KE01": "Ja"}, fragen))


# --- Fehlerfälle (Excel im Speicher verändert, nichts gespeichert) ------------------------

def _wb():
    return openpyxl.load_workbook(config.LOGIK_EXCEL_PFAD, data_only=True)


def _zeile_finden(ws, spalte: int, wert: str) -> int:
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, spalte).value or "").strip() == wert:
            return r
    raise AssertionError(f"{ws.title}: „{wert}“ nicht gefunden")


def _einlesen(wb) -> tuple[Logik, Pruefbericht]:
    bericht = Pruefbericht()
    logik = Logik({}, [], [], [], {}, datetime.now())
    logik.sparten_fragen["KL"] = logik_modul._fragen_einlesen(wb, bericht, "Fragen KL")
    logik_modul._kl_einlesen(wb, logik, bericht)
    return logik, bericht


class Fehlerfaelle(unittest.TestCase):
    def test_unveraendert_gruen(self):
        logik, bericht = _einlesen(_wb())
        self.assertTrue(bericht.ok, bericht.fehler)
        self.assertEqual(bericht.warnungen, [])
        self.assertEqual(len(logik.kl_aktionen), 45)

    def test_fehlende_referenz(self):
        wb = _wb()
        ws = wb["Montagematrix KL"]
        ws.cell(_zeile_finden(ws, 3, "KL023"), 3).value = "KL099"
        logik, bericht = _einlesen(wb)
        self.assertEqual(logik.kl_montage[(3, 1)], "KL099")
        treffer = [f for f in bericht.fehler if f.startswith("KL-Referenz KL099 fehlt im Blatt „KL-Artikel“")]
        self.assertEqual(len(treffer), 1, bericht.fehler)
        self.assertIn("Montagematrix KL 3/1", treffer[0])

    def test_kombination_unzulaessiger_code(self):
        wb = _wb()
        ws = wb["Kombinationen KL"]
        ws.cell(2, 4).value = "10+12"
        ws.cell(2, 3).value = 2
        logik, bericht = _einlesen(wb)
        self.assertTrue(any("unzulässige Codes (10)" in f and "Kombinationen KL Zeile 2" in f
                            for f in bericht.fehler), bericht.fehler)
        self.assertNotIn("10+12", logik.kl_kombis["CL5000M 53/2 E"].get(2, set()))
        # Anzahl passt nicht zur Kombination
        wb = _wb()
        wb["Kombinationen KL"].cell(3, 3).value = 2
        _logik, bericht = _einlesen(wb)
        self.assertTrue(any("passt nicht zur Kombination" in f for f in bericht.fehler), bericht.fehler)
        # Normalisierung: „12+9“ → „9+12“
        wb = _wb()
        wb["Kombinationen KL"].cell(9, 4).value = "12+9"
        logik, bericht = _einlesen(wb)
        self.assertTrue(bericht.ok, bericht.fehler)
        self.assertIn("9+12", logik.kl_kombis["CL5000M 53/2 E"][2])

    def test_montagematrix_doppelzeile(self):
        wb = _wb()
        ws = wb["Montagematrix KL"]
        zeile = _zeile_finden(ws, 3, "KL024")      # 3/2 → 3/1 (doppelt zu KL023)
        ws.cell(zeile, 2).value = 1
        logik, bericht = _einlesen(wb)
        self.assertTrue(any("Montagematrix KL: Kombination 3 Innengeräte / 1 Außengeräte doppelt" in f
                            for f in bericht.fehler), bericht.fehler)
        self.assertEqual(logik.kl_montage[(3, 1)], "KL023")   # erste Zeile gilt

    def test_kr07_hinweis(self):
        wb = _wb()
        ws = wb["Fragen KL"]
        ws.cell(_zeile_finden(ws, 2, "KR07"), 7).value = "Zuordnung"
        _logik, bericht = _einlesen(wb)
        self.assertTrue(any(f.startswith("Fragen KL KR07: Hinweis muss „es erscheinen nur so viele Optionen wie")
                            for f in bericht.fehler), bericht.fehler)

    def test_parameter_fehlt_oder_unlesbar(self):
        wb = _wb()
        ws = wb["KL-Parameter"]
        zeile = _zeile_finden(ws, 1, "Klasse 24 bis")
        ws.cell(zeile, 1).value = None
        ws.cell(zeile, 2).value = None
        logik, bericht = _einlesen(wb)
        self.assertTrue(bericht.ok, bericht.fehler)           # fehlend = Hinweis, nicht Fehler
        self.assertIn("KL-Parameter „Klasse 24 bis“ fehlt – Standardwert 7,0 wird verwendet (Klima v24).",
                      bericht.warnungen)
        self.assertEqual(logik_modul.kl_parameter_fehlende(logik), ["Klasse 24 bis"])
        self.assertEqual(logik_modul.kl_standardwert_text("Klasse 24 bis"), "7,0")
        wb = _wb()
        ws = wb["KL-Parameter"]
        ws.cell(_zeile_finden(ws, 1, "W/m² stark"), 2).value = "neunzig"
        _logik, bericht = _einlesen(wb)
        self.assertIn("KL-Parameter „W/m² stark“: „neunzig“ ist keine Zahl.", bericht.fehler)
        wb = _wb()
        ws = wb["KL-Parameter"]
        ws.cell(_zeile_finden(ws, 1, "Gewerbe-Verhalten"), 2).value = "Egal"
        _logik, bericht = _einlesen(wb)
        self.assertTrue(any("Gewerbe-Verhalten" in f and "Hinweis“ oder „AMPEL" in f
                            for f in bericht.fehler), bericht.fehler)
        wb = _wb()
        ws = wb["KL-Parameter"]
        ws.cell(_zeile_finden(ws, 1, "Standardserie"), 2).value = "Climate 9000i"
        _logik, bericht = _einlesen(wb)
        self.assertTrue(any("Standardserie" in f and "keine Option von KO06" in f
                            for f in bericht.fehler), bericht.fehler)

    def test_aktion_unbekannte_frage_und_antwort(self):
        wb = _wb()
        ws = wb["Aktionen KL"]
        ws.cell(_zeile_finden(ws, 1, "KO07"), 1).value = "KX99"
        _logik, bericht = _einlesen(wb)
        self.assertIn("Aktionen KL: unbekannte Frage „KX99“ (Antwort „Ja“).", bericht.fehler)
        wb = _wb()
        ws = wb["Aktionen KL"]
        zeile = _zeile_finden(ws, 2, "Gerüst erforderlich")
        ws.cell(zeile, 2).value = "Kran erforderlich"
        ws.cell(zeile, 5).value = "nur wenn KA01 = auf dem Mond"
        _logik, bericht = _einlesen(wb)
        self.assertTrue(any(f.startswith("Aktionen KL KZ01: Antwort „Kran erforderlich“ ist keine bekannte Option")
                            for f in bericht.fehler), bericht.fehler)
        self.assertIn("Aktionen KL KZ01: Zusatzbedingungswert „auf dem Mond“ ist keine Option von KA01.",
                      bericht.fehler)
        # Positionszeile ohne Abdeckung aller Optionen → Hinweis
        wb = _wb()
        ws = wb["Aktionen KL"]
        ws.cell(_zeile_finden(ws, 2, "11–15 m"), 2).value = "6–10 m"
        _logik, bericht = _einlesen(wb)
        self.assertIn("Aktionen KL KR05: keine Aktionszeile für Option(en) 11–15 m.", bericht.warnungen)

    def test_blatt_fehlt(self):
        wb = _wb()
        wb.remove(wb["Paketmatrix KL"])
        logik, bericht = _einlesen(wb)
        self.assertIn("Blatt „Paketmatrix KL“ fehlt (Pflicht zu „Aktionen KL“).", bericht.fehler)
        self.assertEqual(logik.kl_aktionen, [])
        # ohne „Aktionen KL“ bleibt KL ein reiner Bogen – keine Fehler, keine kl_*-Daten
        wb = _wb()
        wb.remove(wb["Aktionen KL"])
        logik, bericht = _einlesen(wb)
        self.assertTrue(bericht.ok, bericht.fehler)
        self.assertEqual((logik.kl_aktionen, logik.kl_paket, logik.kl_montage), ([], [], {}))
        sicht = logik_modul.logik_fuer_sparte(logik, "KL")
        self.assertEqual((sicht.sparte, sicht.aktionen, sicht.bloecke), ("KL", [], []))

    def test_paketmatrix_ohne_artikel(self):
        wb = _wb()
        ws = wb["Paketmatrix KL"]
        ws.cell(2, 5).value = "folgt"
        _logik, bericht = _einlesen(wb)
        self.assertTrue(any(f.startswith("Paketmatrix KL Climate 3200i (Standard) / Single / 9: Spalte „Artikel“ ohne")
                            for f in bericht.fehler), bericht.fehler)
        wb = _wb()
        ws = wb["Paketmatrix KL"]
        ws.cell(2, 2).value = "Dual"
        _logik, bericht = _einlesen(wb)
        self.assertTrue(any("unbekannter Typ „Dual“" in f for f in bericht.fehler), bericht.fehler)


# --- Parametrierung (Routen) ----------------------------------------------------------

class Parametrierung(Basis):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.client = TestClient(app)
        cls.client.post("/login", data={"benutzer_id": "1", "pin": "1234"})

    def test_uebersicht(self):
        r = self.client.get("/parametrierung")
        self.assertEqual(r.status_code, 200)
        for text in ("KL-Parameter", "Klima-Positionslisten importieren",
                     'action="/parametrierung/artikel/kl-import"', "Klasse 24 bis",
                     'name="db_ampel_rot_unter_KL"', 'name="db_ampel_gruen_ueber_KL"',
                     "/parametrierung/kl-logik", "KL-Aktionen"):
            self.assertIn(text, r.text, text)
        self.assertNotIn("Fehlende KL-Parameter", r.text)

    def test_kl_logik_lesesicht(self):
        r = self.client.get("/parametrierung/kl-logik")
        self.assertEqual(r.status_code, 200)
        for text in ("Paketmatrix KL", "Montagematrix KL", "Kombinationen KL", "CL5000M 79/3 E",
                     "9+9+9", "KL023", "Klimaanlage Bosch", "KL-Artikel (Sparte KL)", "KL050",
                     "Klima-Prüfungen ohne Fehler", "Angebotsaufbau KL", "Aktionen KL"):
            self.assertIn(text, r.text, text)
        self.assertIn("EK Material", r.text)          # Admin sieht EK

    def test_kl_import_vorschau_und_ausfuehren(self):
        r = self.client.get("/parametrierung/artikel/kl-import")
        self.assertEqual(r.status_code, 200)
        self.assertIn("Klima-Positionslisten-Import", r.text)
        self.assertIn("Import ausführen", r.text)
        r = self.client.post("/parametrierung/artikel/kl-import", follow_redirects=False)
        self.assertEqual(r.status_code, 303)
        self.assertIn("/parametrierung?meldung=", r.headers["location"])
        self.assertIn("Klima-Import+abgeschlossen", r.headers["location"])
        self.assertIn("0+neu", r.headers["location"])          # Re-Import ohne Doppel
        aktiv = (self.s.query(Artikel).filter(Artikel.quelle == QUELLE_KL,
                                              Artikel.aktiv.is_(True)).count())
        self.assertEqual(aktiv, 50)
        # Logik danach neu eingelesen und weiterhin grün
        _logik, bericht = logik_modul.hole_logik(self.s)
        self.assertTrue(bericht.ok, bericht.fehler)

    def test_db_ampel_kl_speichern(self):
        from app.models import db_schwellen, einstellung_holen
        vorher = (einstellung_holen(self.s, "db_ampel_rot_unter_KL", ""),
                  einstellung_holen(self.s, "db_ampel_gruen_ueber_KL", ""))
        try:
            r = self.client.post("/parametrierung/einstellungen",
                                 data={"db_sparten_formular": "1",
                                       "db_ampel_rot_unter_KL": "1500",
                                       "db_ampel_gruen_ueber_KL": "2500"},
                                 follow_redirects=False)
            self.assertEqual(r.status_code, 303)
            self.s.expire_all()
            self.assertEqual(db_schwellen(self.s, "KL"), (1500, 2500))   # Euro
        finally:
            self.client.post("/parametrierung/einstellungen",
                             data={"db_sparten_formular": "1",
                                   "db_ampel_rot_unter_KL": vorher[0],
                                   "db_ampel_gruen_ueber_KL": vorher[1]},
                             follow_redirects=False)


if __name__ == "__main__":
    unittest.main()
