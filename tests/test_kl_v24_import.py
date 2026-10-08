# Tests PLAN_V16 Phase 113 (v24, Agent A): KL-Blätter der Live-Excel und
# Import der Klima-Positionslisten (app/import_klima.py). Laufen wie
# tests/test_pv_v13.py gegen die Entwicklungs-DB; der Import legt die
# KL-Artikel an, falls sie fehlen (idempotent, keine Testdaten zum Aufräumen).
import re
import unittest
import warnings

warnings.filterwarnings("ignore")

import openpyxl

from app import config, import_klima, import_preisliste, import_pv
from app import logik as logik_modul
from app.db import SessionLocal, init_db
from app.models import Artikel

QUELLE_KL = import_klima.QUELLE_KL
KL_BLAETTER = ("Fragen KL", "KL-Artikel", "Paketmatrix KL", "Kombinationen KL",
               "Montagematrix KL", "Aktionen KL", "Angebotsaufbau KL", "KL-Parameter")
MONTAGE_PLAN = {(1, 1): "KL020", (2, 1): "KL021", (2, 2): "KL022",
                (3, 1): "KL023", (3, 2): "KL024", (4, 1): "KL025"}
FRAGEN_REIHENFOLGE = ["KO01", "KO02", "KO03", "KO06", "KO04", "KO05", "KO07", "KO08", "KO09",
                      "KR01", "KR08", "KR09", "KR10", "KR02", "KR03", "KR04", "KR05", "KR06",
                      "KR11", "KR07", "KA01", "KA02", "KE01", "KE02", "KE03", "KZ01",
                      "S01", "S02"]
LEISTUNGSUMFANG = [
    "– Montage der Innengeräte (Wandmontage) und des Außengeräts auf Wand- oder "
    "Bodenkonsole (Konsole enthalten)",
    "– Kältemittelleitung isoliert, Kondensatleitung, Elektroverbindung Innen-/Außengerät "
    "und Montagekanal bis 5 m je Innengerät",
    "– eine Kernbohrung je Innengerät durch die Außenwand (bis 32 cm)",
    "– Evakuieren, Dichtheitsprüfung mit Dokumentation (F-Gase-Verordnung), Inbetriebnahme "
    "und Einweisung",
    "– An- und Abfahrt, Kleinmaterial, Entsorgung des Verpackungsmaterials",
]


def _zeilen(wb, blatt):
    """(Kopf, Datenzeilen ohne Leerzeilen) eines Blatts."""
    rows = [tuple(r) for r in wb[blatt].iter_rows(values_only=True)
            if any(v is not None and str(v).strip() for v in r)]
    return [str(v or "").strip() for v in rows[0]], rows[1:]


def _text(v) -> str:
    return str(v or "").strip()


# --- Blätter der Live-Excel ----------------------------------------------------

class Blaetter(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.wb = openpyxl.load_workbook(config.LOGIK_EXCEL_PFAD, read_only=True,
                                        data_only=True)

    @classmethod
    def tearDownClass(cls):
        cls.wb.close()

    def test_blaetter_vorhanden(self):
        for name in KL_BLAETTER + ("Lesehilfe", "Textregeln", "Anhänge"):
            self.assertIn(name, self.wb.sheetnames)

    def test_fragen_kl(self):
        kopf, zeilen = _zeilen(self.wb, "Fragen KL")
        self.assertEqual(kopf[:7], ["Seite", "ID", "Fragetext", "Typ", "Antwortmöglichkeiten",
                                    "Anzeigen wenn", "Hinweis"])
        self.assertEqual(len(zeilen), 28)
        self.assertEqual([_text(z[1]) for z in zeilen], FRAGEN_REIHENFOLGE)
        nach_id = {_text(z[1]): z for z in zeilen}
        self.assertEqual(_text(nach_id["KE02"][5]), "nur wenn KE01 = Nein oder Unklar")
        self.assertEqual(_text(nach_id["KE03"][5]), "nur wenn KE01 = Nein oder Unklar")
        self.assertEqual(_text(nach_id["KR05"][4]), "bis 5 m | 6–10 m | 11–15 m | über 15 m")
        self.assertEqual(_text(nach_id["KA02"][4]), "unter 2 m | 2–4 m | über 4 m")
        self.assertEqual(_text(nach_id["KE02"][4]), "0–5 m | 6–10 m | 11–15 m | über 15 m")
        self.assertEqual(_text(nach_id["KR07"][6]),
                         "es erscheinen nur so viele Optionen wie Außengeräte (KO04)")
        self.assertEqual(_text(nach_id["KO05"][6]),
                         "öffnet je Raum einen eigenen Fragenblock (höchstens 12 Räume)")
        self.assertEqual(_text(nach_id["KR08"][6]),
                         "Grundfläche des Raums (Dezimalzahl möglich, z. B. 22,5)")
        self.assertEqual(_text(nach_id["KO06"][6]),
                         "Vorbelegung Climate 3200i. Climate 7000i: nur bis 3,4 kW je Raum · "
                         "Climate 8000i: nur Single-Split bis 3,5 kW")
        for fid in ("KO01", "KO02", "KO03", "KO04", "KO07", "KR09", "KR10", "KR02", "KR03",
                    "KR04", "KR06", "KR11", "KA02", "KE01", "KE02", "KE03"):
            self.assertEqual(_text(nach_id[fid][6]), "", fid)
        for z in zeilen:
            self.assertIn(_text(z[5]).split(" ")[0], ("immer", "je", "nur"))
            self.assertEqual(_text(z[5]) if _text(z[5]).startswith("je") else "je Raum (KO05)",
                             "je Raum (KO05)")
        alles = " ".join(_text(v) for z in zeilen for v in z)
        for verboten in ("NEU ·", "GEÄNDERT", "[ANNAHME]"):
            self.assertNotIn(verboten, alles)
        # je-Raum-Block: KR01 … KR07 in der Plan-Reihenfolge
        raum = [_text(z[1]) for z in zeilen if _text(z[5]) == "je Raum (KO05)"]
        self.assertEqual(raum, ["KR01", "KR08", "KR09", "KR10", "KR02", "KR03", "KR04",
                                "KR05", "KR06", "KR11", "KR07"])

    def test_fragen_kl_lader(self):
        logik, bericht = logik_modul.logik_einlesen()
        self.assertFalse([f for f in bericht.fehler if "Fragen KL" in f], bericht.fehler)
        fragen = logik.sparten_fragen["KL"]
        self.assertEqual(list(fragen), FRAGEN_REIHENFOLGE)
        ke02 = fragen["KE02"].bedingung
        self.assertEqual((ke02.art, ke02.frage_id, ke02.werte), ("antwort", "KE01",
                                                                   ["Nein", "Unklar"]))
        self.assertEqual(fragen["KR07"].bedingung.art, "wiederholgruppe")
        self.assertTrue(re.search(r"Optionen wie .*\(([A-Z]{1,2}\d{2})\)",
                                  fragen["KR07"].hinweis))
        self.assertEqual(fragen["KO09"].typ, "Freitext groß")
        self.assertEqual(fragen["KR08"].typ, "Zahleneingabe")

    def test_kl_artikel(self):
        kopf, zeilen = _zeilen(self.wb, "KL-Artikel")
        self.assertEqual(kopf[:8], ["KL-Nr.", "Datei", "Pos. in Datei", "GUID",
                                    "Bezeichnung (1. Zeile)", "VK netto (€)", "EK (€)", "EP"])
        self.assertEqual(len(zeilen), 50)
        self.assertEqual([_text(z[0]) for z in zeilen], [f"KL{n:03d}" for n in range(1, 51)])
        nach_nr = {_text(z[0]): z for z in zeilen}
        # GUIDs aus den Dateien (Entwurf war ab KL030 versetzt)
        self.assertEqual(_text(nach_nr["KL001"][3]), "{B23F65ED-EC7E-4948-AB14-274241F0E1C5}")
        self.assertEqual(_text(nach_nr["KL020"][3]), "{FF5ECB53-BC6C-4DEA-998F-0A0AF2FD2F40}")
        self.assertEqual(_text(nach_nr["KL030"][3]), "{0418B42B-7D2A-4EC6-885B-89A0E7667935}")
        self.assertEqual(_text(nach_nr["KL048"][3]), "{715A1E99-746B-4094-9F2A-90C16455FFA6}")
        self.assertEqual(_text(nach_nr["KL049"][3]), "{38FC5932-9389-46DE-A9D0-8855676E1FD8}")
        self.assertEqual(_text(nach_nr["KL049"][1]), "Musterangebot KL")
        self.assertEqual(_text(nach_nr["KL049"][2]), "006")
        self.assertFalse(_text(nach_nr["KL050"][3]).startswith("{"))
        guids = [_text(z[3]) for z in zeilen if _text(z[3]).startswith("{")]
        self.assertEqual(len(guids), 49)
        self.assertEqual(len(set(guids)), 49)
        for nr in range(1, 49):
            self.assertEqual(_text(nach_nr[f"KL{nr:03d}"][2]), f"{nr:03d}")
            self.assertEqual(_text(nach_nr[f"KL{nr:03d}"][1]), "Ersatzangebot KL")
        self.assertEqual(float(nach_nr["KL023"][5]), 3178.0)
        self.assertEqual(float(nach_nr["KL038"][5]), 1966.24)
        self.assertEqual(float(nach_nr["KL034"][5]), 1341.6)
        self.assertEqual(float(nach_nr["KL049"][5]), 0.0)
        self.assertEqual(float(nach_nr["KL050"][5]), 499.0)
        self.assertIn(nach_nr["KL050"][6], (None, ""))
        self.assertTrue(all(_text(z[7]) == "" for z in zeilen))   # EP nur über die Aktion
        pins = import_klima.pins_lesen()
        self.assertEqual(len(pins), 49)
        self.assertEqual(pins["{38FC5932-9389-46DE-A9D0-8855676E1FD8}"], "KL049")

    def test_paketmatrix(self):
        kopf, zeilen = _zeilen(self.wb, "Paketmatrix KL")
        self.assertEqual(kopf[:6], ["Serie (KO06)", "Typ", "Klasse / Außengerät",
                                    "Kühllast je Raum", "Artikel", "Bemerkung"])
        self.assertEqual(len(zeilen), 23)
        serien = {_text(z[0]) for z in zeilen}
        self.assertEqual(serien, {"Climate 3200i (Standard)", "Climate 7000i", "Climate 8000i"})
        single_3200 = {_text(z[2]): _text(z[4]) for z in zeilen
                       if _text(z[0]).startswith("Climate 3200i") and _text(z[1]) == "Single"}
        self.assertTrue(single_3200["9"].startswith("KL034"))
        self.assertTrue(single_3200["12"].startswith("KL035"))
        self.assertTrue(single_3200["18"].startswith("KL036"))
        self.assertEqual(re.findall(r"KL\d{3}", single_3200["24"]), ["KL033", "KL029"])
        multi_aussen = [z for z in zeilen if _text(z[1]) == "Multi-Außengerät"]
        self.assertEqual([re.findall(r"KL\d{3}", _text(z[4]))[0] for z in multi_aussen],
                         ["KL037", "KL038", "KL039", "KL040", "KL045", "KL046"])
        nicht = [z for z in zeilen if _text(z[4]) == "nicht im Sortiment"]
        self.assertEqual(len(nicht), 3)
        for z in nicht:
            self.assertTrue(_text(z[5]).startswith("→ AMPEL: "), z)
        alle_refs = {r for z in zeilen for r in re.findall(r"KL\d{3}", _text(z[4]))}
        self.assertTrue(alle_refs <= {f"KL{n:03d}" for n in range(1, 51)})

    def test_kombinationen(self):
        kopf, zeilen = _zeilen(self.wb, "Kombinationen KL")
        self.assertEqual(kopf[:5], ["Außengerät", "Artikel", "Anzahl Innengeräte",
                                    "Kombination (Klassen-Codes, aufsteigend)", "Quelle"])
        self.assertEqual(len(zeilen), 215)
        erlaubt = {"7", "9", "12", "18", "24"}
        kombis: dict[str, dict[int, set[str]]] = {}
        for z in zeilen:
            codes = _text(z[3]).split("+")
            self.assertTrue(set(codes) <= erlaubt, z)
            self.assertEqual(codes, sorted(codes, key=int), z)
            self.assertEqual(len(codes), int(z[2]), z)
            kombis.setdefault(_text(z[0]), {}).setdefault(int(z[2]), set()).add(_text(z[3]))
        self.assertEqual(sum(1 for z in zeilen if _text(z[0]).startswith("CL5000M")), 205)
        self.assertEqual(sum(1 for z in zeilen if _text(z[0]).startswith("CL7000M")), 10)
        self.assertIn("9+9+9", kombis["CL5000M 79/3 E"][3])         # Fall A
        self.assertNotIn(3, kombis["CL5000M 53/2 E"])              # max. 2 Innengeräte
        self.assertIn("12+18", kombis["CL5000M 53/2 E"][2])         # Fall C
        self.assertIn("9+9+9+9", kombis["CL5000M 105/4 E"][4])     # Fall H
        self.assertNotIn("24+24", kombis["CL5000M 53/2 E"].get(2, set()))   # Fall F → G4
        for odu, nr in (("CL5000M 53/2 E", "KL037"), ("CL5000M 79/3 E", "KL038"),
                        ("CL5000M 105/4 E", "KL039"), ("CL5000M 125/5 E", "KL040"),
                        ("CL7000M 53/2 E", "KL045"), ("CL7000M 79/3 E", "KL046")):
            self.assertEqual({_text(z[1]) for z in zeilen if _text(z[0]) == odu}, {nr})

    def test_montagematrix(self):
        kopf, zeilen = _zeilen(self.wb, "Montagematrix KL")
        self.assertEqual(kopf[:5], ["Innengeräte gesamt", "Außengeräte (KO04)", "Position",
                                    "VK netto (€)", "Bemerkung"])
        self.assertEqual(len(zeilen), 7)
        matrix = {(int(z[0]), int(z[1])): _text(z[2]) for z in zeilen
                  if _text(z[0]).isdigit()}
        self.assertEqual(matrix, MONTAGE_PLAN)
        ampel = [z for z in zeilen if _text(z[2]) == "AMPEL"]
        self.assertEqual(len(ampel), 1)
        self.assertEqual(_text(ampel[0][0]), "alle anderen")
        self.assertIn("individuell – Grund: Montagekombination", _text(ampel[0][4]))

    def test_aktionen_kl(self):
        kopf, zeilen = _zeilen(self.wb, "Aktionen KL")
        self.assertEqual(kopf[:5], ["Frage", "Antwort / Bedingung", "Aktion", "Bemerkung",
                                    "Zusatzbedingung"])
        aktionen = [_text(z[2]) for z in zeilen]
        for soll in ("Artikel: KL049 ×1", "Artikel: KL017 × 5", "Artikel: KL017 × 10",
                     "Artikel: KL013 × Innengeräte", "Artikel: KL013 × Innengeräte als EP",
                     "Artikel: KL001 × Außengeräte", "Artikel: KL008 × Außengeräte",
                     "Artikel: KL004 ×1 · KL005 ×1 · KL006 ×1 je Außengerät",
                     "Artikel: KL050 ×1", "Artikel: KL015 ×1", "Artikel: KL019 ×1",
                     "Artikel: KL016 ×1", "Artikel: KL014 ×1", "Artikel: KL009 ×1",
                     "Hinweis: Zustimmung des Eigentümers einholen",
                     "Hinweis: Gewerbeobjekt – Auslegung prüfen",
                     "Hinweis: Flachdach – keine Dachdurchführung, Leitungsführung über "
                     "Attika/Fassade prüfen",
                     "Hinweis: Montageort Außengerät klären",
                     "Hinweis: Montagehöhe über 4 m ohne Gerüst/Bühne prüfen",
                     "Hinweis: Stromversorgung am Außengerät prüfen",
                     "AMPEL: individuell – Grund: Elektrozuleitung über 15 m"):
            self.assertIn(soll, aktionen)
        for z in zeilen:
            self.assertNotIn("≠", _text(z[4]), z)
            self.assertFalse(_text(z[2]).startswith("Fachlicher Hinweis"), z)
            if _text(z[4]):
                self.assertIsNotNone(logik_modul.bedingung_parsen(_text(z[4])), z)
        nach = {(_text(z[0]), _text(z[1])): z for z in zeilen}
        self.assertEqual(
            _text(nach[("KZ01", "Gerüst erforderlich")][4]),
            "nur wenn KA01 = auf dem Boden oder an der Außenwand oder auf einer Terrasse "
            "oder auf einem Balkon oder auf einem Flachdach oder noch unklar")
        self.assertEqual(_text(nach[("KE03", "Nein")][4]), "nur wenn KE01 = Nein oder Unklar")
        self.assertEqual(_text(nach[("KA02", "über 4 m")][4]), "nur wenn KZ01 = vom Boden/Leiter")
        self.assertEqual(_text(nach[("KA01", "auf einem Schrägdach")][4]), "")
        # Dokumentationszeilen tragen keine Artikelreferenz in der Spalte Aktion
        for frage in ("Auslegung", "Händisch"):
            for z in zeilen:
                if _text(z[0]) == frage:
                    self.assertFalse(re.findall(r"KL\d{3}", _text(z[2])), z)
        grundpaket = [_text(z[2]) for z in zeilen if _text(z[0]) == "Grundpaket"]
        self.assertEqual([re.findall(r"KL\d{3}", a) for a in grundpaket], [["KL049"], []])
        # Antworten der Fragezeilen sind Optionen der Fragen KL
        logik, _b = logik_modul.logik_einlesen()
        fragen = logik.sparten_fragen["KL"]
        for z in zeilen:
            frage = _text(z[0])
            if frage not in fragen:
                self.assertIn(frage, ("Grundpaket", "Auslegung", "§14a", "Händisch"), z)
                continue
            problem = logik_modul._antwort_pruefen(fragen[frage], _text(z[1]), fragen)
            self.assertIsNone(problem, (z, problem))

    def test_angebotsaufbau(self):
        kopf, zeilen = _zeilen(self.wb, "Angebotsaufbau KL")
        self.assertEqual(kopf[:4], ["Block", "Überschrift im Angebot", "Inhalt", "Wann"])
        bloecke = {int(float(_text(z[0]))): z for z in zeilen
                   if _text(z[0]).replace(".0", "").isdigit()}
        self.assertEqual(sorted(bloecke), [1, 2, 3])
        self.assertEqual(_text(bloecke[3][1]), "Klimaanlage Bosch")
        self.assertEqual(re.findall(r"KL\d{3}", _text(bloecke[1][2])),
                         ["KL020", "KL021", "KL022", "KL023", "KL024", "KL025", "KL019",
                          "KL016", "KL014", "KL017", "KL004", "KL005", "KL006", "KL050",
                          "KL015"])
        self.assertEqual(re.findall(r"KL\d{3}", _text(bloecke[2][2])),
                         ["KL001", "KL008", "KL009", "KL003"])
        self.assertEqual(re.findall(r"KL\d{3}", _text(bloecke[3][2])),
                         ["KL034", "KL035", "KL036", "KL033", "KL037", "KL038", "KL039",
                          "KL040", "KL043", "KL044", "KL045", "KL046", "KL047", "KL048",
                          "KL026", "KL027", "KL028", "KL029", "KL041", "KL042", "KL013",
                          "KL049"])
        self.assertTrue(_text(bloecke[3][2]).endswith("· Auslegungszeile"))
        for nr in (1, 2, 3):
            self.assertEqual(_text(bloecke[nr][3]), "immer")
        sonstige = [_text(z[0]) for z in zeilen if _text(z[0]) not in ("1", "2", "3")]
        for name in ("Summenblock", "Nachtexte", "Anhänge", "Protokoll"):
            self.assertIn(name, sonstige)
        summen = next(z for z in zeilen if _text(z[0]) == "Summenblock")
        self.assertEqual(_text(summen[1]),
                         "Netto · Umsatzsteuer 19 % · Gesamt-Betrag · ggf. Rabatt · Endbetrag")

    def test_kl_parameter(self):
        kopf, zeilen = _zeilen(self.wb, "KL-Parameter")
        self.assertEqual(kopf[:4], ["Parameter", "Wert", "Einheit", "Bemerkung"])
        self.assertEqual(len(zeilen), 17)
        werte = {_text(z[0]): _text(z[1]) for z in zeilen}
        self.assertEqual(list(werte), [
            "W/m² normal", "W/m² stark", "Höhenfaktor bis 2,5 m", "Höhenfaktor 2,5–3 m",
            "Höhenfaktor über 3 m", "Klasse 9 bis", "Klasse 12 bis", "Klasse 18 bis",
            "Klasse 24 bis", "Leitung je Innengerät inklusive", "Meterposition Zusatzleitung",
            "Standardserie", "Max. Innengeräte je Außengerät", "Max. Außengeräte",
            "§14a-Außengeräte", "Gewerbe-Verhalten", "Rollgerüst VK"])
        self.assertEqual((float(werte["W/m² normal"]), float(werte["W/m² stark"])), (60, 90))
        self.assertEqual([float(werte[k]) for k in ("Klasse 9 bis", "Klasse 12 bis",
                                                     "Klasse 18 bis", "Klasse 24 bis")],
                         [2.6, 3.5, 5.3, 7.0])
        self.assertEqual(werte["Standardserie"], "Climate 3200i (Standard)")
        self.assertEqual(werte["Meterposition Zusatzleitung"], "KL017")
        self.assertEqual(werte["§14a-Außengeräte"], "CL5000M 105/4 E; CL5000M 125/5 E")
        self.assertEqual(werte["Gewerbe-Verhalten"], "Hinweis")
        self.assertEqual(float(werte["Rollgerüst VK"]), 499)
        self.assertEqual(import_klima.rollgeruest_vk(), 499)

    def test_textregeln(self):
        kopf, zeilen = _zeilen(self.wb, "Textregeln")
        betrifft = [_text(z[0]) for z in zeilen]
        for b in ("Positionen KL020–KL025", "Position KL013", "Positionen KL026–KL029"):
            self.assertIn(b, betrifft)
        regeln = import_klima.kl_textregeln_lesen()
        self.assertEqual(regeln.warnungen, [])          # alle drei Zeilen erkannt
        for nr in ("KL020", "KL021", "KL022", "KL023", "KL024", "KL025"):
            self.assertEqual(regeln.ergaenzen[nr],
                             ("Diese Leistung besteht aus folgenden Positionen", LEISTUNGSUMFANG))
        self.assertEqual(regeln.ersetzen["KL013"],
                         [("für Split-Klimagerät CL3000i", "für Bosch Climate Split-Klimageräte")])
        for nr in ("KL026", "KL027", "KL028", "KL029"):
            self.assertEqual(regeln.kurz[nr], ("Split Inneneinheit",
                                               "BOSCH CL3200iU W <26|35|53|70> E Inneneinheit <kW>",
                                               60))
        # WP-Import stolpert nicht über die KL-Zeilen (drei Spalten, kein Absturz)
        wp_regeln = import_preisliste.lade_textregeln(self.wb)
        self.assertEqual(wp_regeln.zeile_entfernen.get("045"), "Pufferspeicher BST 50")

    def test_anhaenge_und_lesehilfe(self):
        kopf, zeilen = _zeilen(self.wb, "Anhänge")
        # v27-Nachtrag 2 (08.10.2026): die v24-Platzhalterzeile „(Bosch Climate Broschüre –
        # Zulieferung)“ trägt jetzt den Dateinamen der gelieferten Broschüre
        zeile = next((z for z in zeilen if _text(z[0]) == "Bosch Climate 3200i.pdf"), None)
        self.assertIsNotNone(zeile)
        self.assertEqual(_text(zeile[1]), "wenn Sparte = KL")
        self.assertIn("Bosch Climate 3200i", _text(zeile[2]))
        self.assertIn("(v27)", _text(zeile[2]))
        self.assertFalse([z for z in zeilen if _text(z[0]).startswith("(Bosch Climate Broschüre")])
        self.assertTrue(re.match(r"wenn\s+Sparte\s*=\s*(\w+)$", _text(zeile[1])))
        logik, bericht = logik_modul.logik_einlesen()
        self.assertFalse([w for w in bericht.warnungen if "Bosch Climate" in w])
        texte = " ".join(_text(v) for row in self.wb["Lesehilfe"].iter_rows(values_only=True)
                         for v in row)
        self.assertIn("Klima (v24)", texte)
        self.assertIn("Kombinationen KL", texte)


# --- Import auf der Entwicklungs-DB --------------------------------------------

class Import(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        init_db()
        cls.s = SessionLocal()
        if cls.s.query(Artikel).filter(Artikel.quelle == QUELLE_KL,
                                       Artikel.aktiv.is_(True)).count() < 50:
            cls.diff, cls.meldung = import_klima.import_ausfuehren(cls.s)
        else:
            cls.diff, cls.meldung = import_klima.berechne_diff(cls.s), ""

    @classmethod
    def tearDownClass(cls):
        cls.s.close()

    def aktive(self):
        return {a.pos_nr: a for a in self.s.query(Artikel)
                .filter(Artikel.quelle == QUELLE_KL, Artikel.aktiv.is_(True))}

    def test_50_artikel_aktiv(self):
        kl = self.aktive()
        self.assertEqual(sorted(kl), [f"KL{n:03d}" for n in range(1, 51)])
        self.assertTrue(all(a.quelle == "kl" for a in kl.values()))
        self.assertTrue(all(not a.ep_flag for a in kl.values()))   # EP nur über die Aktion
        self.assertTrue(all(a.kategorie.startswith("KL · ") for a in kl.values()))
        self.assertEqual(import_klima.fehlende_kl_artikel(self.s), [])

    def test_stichproben(self):
        kl = self.aktive()
        self.assertEqual((kl["KL023"].e_preis_cent, kl["KL023"].einheit), (317800, "psl."))
        self.assertEqual((kl["KL038"].e_preis_cent, kl["KL038"].einheit), (196624, "Stück"))
        self.assertEqual((kl["KL034"].e_preis_cent, kl["KL034"].einheit), (134160, "Set"))
        self.assertEqual((kl["KL049"].e_preis_cent, kl["KL049"].ek_cent), (0, 0))
        self.assertEqual(kl["KL049"].guid, "{38FC5932-9389-46DE-A9D0-8855676E1FD8}")
        self.assertEqual(kl["KL049"].bezeichnung, "Erweiterung Systemgarantie")
        self.assertEqual((kl["KL050"].e_preis_cent, kl["KL050"].einheit, kl["KL050"].guid,
                          kl["KL050"].ek_cent), (49900, "psl.", None, None))
        self.assertEqual(kl["KL050"].bezeichnung, "Rollgerüst / Arbeitsgerüst, Auf- und Abbau")
        self.assertEqual(kl["KL001"].guid, "{B23F65ED-EC7E-4948-AB14-274241F0E1C5}")
        self.assertEqual(kl["KL001"].ek_cent, 1000)
        self.assertEqual(kl["KL017"].einheit, "m")
        # Bezeichnung = erste Zeile, Beschreibung = Rest, _x000D_ bereinigt
        self.assertEqual(kl["KL020"].bezeichnung, "Montagepauschale der Klimaanlage")
        self.assertTrue(kl["KL020"].beschreibung.startswith("Klimaanlage mit 1 Innengerät"))
        for a in kl.values():
            self.assertNotIn("_x000D_", a.bezeichnung + a.beschreibung)
            self.assertNotIn("\n", a.bezeichnung)

    def test_textregeln_wirksam(self):
        kl = self.aktive()
        for nr in ("KL020", "KL021", "KL022", "KL023", "KL024", "KL025"):
            zeilen = kl[nr].beschreibung.splitlines()
            i = zeilen.index("Diese Leistung besteht aus folgenden Positionen")
            self.assertEqual(zeilen[i + 1:i + 6], LEISTUNGSUMFANG, nr)
            self.assertEqual(zeilen.count(LEISTUNGSUMFANG[0]), 1, nr)   # Re-Import: nicht doppelt
        self.assertEqual(kl["KL013"].beschreibung, "für Bosch Climate Split-Klimageräte")
        self.assertNotIn("CL3000i", kl["KL013"].beschreibung)
        self.assertEqual(kl["KL026"].bezeichnung, "BOSCH CL3200iU W 26 E Inneneinheit 2,6 kW")
        self.assertEqual(kl["KL027"].bezeichnung, "BOSCH CL3200iU W 35 E Inneneinheit 3,5 kW")
        self.assertEqual(kl["KL028"].bezeichnung, "BOSCH CL3200iU W 53 E Inneneinheit 5,3 kW")
        self.assertEqual(kl["KL029"].bezeichnung, "BOSCH CL3200iU W 70 E Inneneinheit 7 kW")
        self.assertTrue(kl["KL026"].beschreibung.startswith(
            "BOSCH Klimagerät CL3200iU W 26 E, Split Inneneinheit"))
        # Regel (3) greift nur bei „Split Inneneinheit“ + über 60 Zeichen
        self.assertTrue(kl["KL030"].bezeichnung.startswith("BOSCH Klimagerät CL3000i 26 E"))
        self.assertTrue(kl["KL041"].bezeichnung.startswith("BOSCH Split-Klimagerät CL7000iU"))

    def test_reimport_ohne_doppel(self):
        vorher = {a.pos_nr: a.id for a in self.aktive().values()}
        diff, meldung = import_klima.import_ausfuehren(self.s)
        self.assertTrue(meldung.startswith("KL: "))
        self.assertEqual((len(diff.neu), len(diff.entfallen)), (0, 0))
        nachher = self.aktive()
        self.assertEqual({nr: a.id for nr, a in nachher.items()}, vorher)
        self.assertEqual(self.s.query(Artikel).filter(Artikel.pos_nr.like("KL%")).count(), 50)
        guids = [a.guid for a in nachher.values() if a.guid]
        self.assertEqual(len(guids), len(set(guids)))
        diff2 = import_klima.berechne_diff(self.s)
        self.assertEqual((len(diff2.neu), len(diff2.geaendert), len(diff2.entfallen)),
                         (0, 0, 0))
        self.assertIn(import_klima.ROLLGERUEST_EK_HINWEIS, diff2.warnungen)

    def test_wp_und_pv_import_fassen_kl_nicht_an(self):
        ergebnis = import_preisliste.lese_dateien()
        diff = import_preisliste.berechne_diff(self.s, ergebnis)
        self.assertFalse([a for a in diff.entfallen if a.quelle == QUELLE_KL])
        self.assertFalse([a for a in diff.entfallen if a.pos_nr.startswith("KL")])
        pv_diff = import_pv.berechne_diff(self.s)
        self.assertFalse([a for a in pv_diff.entfallen if a.quelle == QUELLE_KL])
        # und umgekehrt: der KL-Import sieht nur seine Quelle
        kl_diff = import_klima.berechne_diff(self.s)
        self.assertTrue(all(a.quelle == QUELLE_KL for a in kl_diff.entfallen))

    def test_kl_referenzen(self):
        refs = import_klima.kl_referenzen()
        self.assertTrue(refs)
        self.assertTrue(all(re.fullmatch(r"KL\d{3}", r) for r in refs))
        for nr in ("KL023", "KL038", "KL034", "KL049", "KL050", "KL013"):
            self.assertIn(nr, refs)


if __name__ == "__main__":
    unittest.main()
