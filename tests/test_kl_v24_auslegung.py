# Tests PLAN_V16 Phase 114 – Klimakonfigurator, Rechenkern app/kl_auslegung.py
# (Agent C). Kontrollfälle A–H aus PLAN_V16 Phase 116 auf Ebene auslegen() /
# artikel_ermitteln() / positionen_bauen(), Validierung (Fall G), Protokoll,
# Einbindung in konfigurator.* und angebot_aufbau.angebot_anlegen.
#
# Logik-Quelle: solange Agent A (KL-Blätter in der Live-Excel) und Agent B
# (Lader `kl_*`) nicht fertig sind, baut `attrappe()` die Logik-Sicht KL aus
# den Werten des Entwurfs docs/KL-Logik-Entwurf.xlsx nach der vereinbarten
# Schnittstelle (SimpleNamespace). Liefert logik_fuer_sparte(…, "KL") bereits
# die kl_*-Felder, laufen dieselben Tests zusätzlich gegen die echte Logik
# (Klasse RechenwegEcht). Artikelpreise für die Netto-Kontrolle kommen aus dem
# Blatt „KL-Artikel“ des Entwurfs (= TAIFUN-Listen).
import copy
import json
import re
import unittest
import warnings
from decimal import Decimal
from types import SimpleNamespace

warnings.filterwarnings("ignore")

import openpyxl

from app import config, kl_auslegung
from app import konfigurator as engine
from app import logik as logik_modul
from app.logik import Aktion, AngebotsBlock, ArtikelRef, Bedingung, bedingung_parsen

ENTWURF = config.PROJEKT_ORDNER / "docs" / "KL-Logik-Entwurf.xlsx"
TEST_EMAIL = "kl-v24-auslegung@test.local"
STARK = "stark (Süd-/Westseite, große Fenster, Dachgeschoss)"

# Inhalt-Spalte des Blatts „Angebotsaufbau KL“ laut PLAN_V16 Phase 113
# (jede KL-Nummer ausdrücklich – der Entwurf nennt sie nur beschreibend)
BLOCK_INHALT = {
    1: ("KL020 · KL021 · KL022 · KL023 · KL024 · KL025 Montagepauschale lt. Montagematrix · "
        "KL019 Demontage · KL016 Kondensatpumpe · KL014 Kernbohrung Beton · KL017 Zusatzmeter · "
        "KL004 · KL005 · KL006 Dacharbeiten · KL050 Rollgerüst · KL015 Arbeitsbühne"),
    2: "KL001 · KL008 Elektro · KL009 Gehäuse · KL003 §14a-Anmeldung",
    3: ("KL034 · KL035 · KL036 · KL033 · KL037 · KL038 · KL039 · KL040 · KL043 · KL044 · "
        "KL045 · KL046 · KL047 · KL048 Außengeräte/Sets · KL026 · KL027 · KL028 · KL029 · "
        "KL041 · KL042 Innengeräte · KL013 WLAN-Gateway · KL049 Systemgarantie · Auslegungszeile"),
}
BLOCK_UEBERSCHRIFT = {1: "(ohne Überschrift)", 2: "(ohne Überschrift)", 3: "Klimaanlage Bosch"}

# Zusatzbedingungen des Entwurfs mit „≠“ → Schreibweise laut PLAN_V16 (Parser kennt kein ≠)
ZUSATZ_UMSCHREIBEN = {
    "nur wenn KA01 ≠ Schrägdach": ("nur wenn KA01 = auf dem Boden oder an der Außenwand oder "
                                   "auf einer Terrasse oder auf einem Balkon oder auf einem "
                                   "Flachdach oder noch unklar"),
    "nur wenn KE01 ≠ Ja": "nur wenn KE01 = Nein oder Unklar",
    "nur wenn KA01 = Schrägdach": "nur wenn KA01 = auf einem Schrägdach",
}

_MENGE_MUSTER = r"\s*×\s*(\([^)]*\)|[\wäöüÄÖÜß. ]+)"


def _z(wert) -> str:
    return str(wert).strip() if wert is not None else ""


def kl_refs(text: str) -> list[ArtikelRef]:
    """KL-Referenzen einer Aktionszeile (Attrappe des Laders): „KL017 × 5“,
    „KL013 × Innengeräte als EP“, „KL004 ×1 · KL005 ×1 · KL006 ×1 je Außengerät“."""
    refs = []
    for m in re.finditer(r"\bKL(\d{3})\b(\s*\(EP\))?", text or ""):
        rest = text[m.end():]
        mm = re.match(_MENGE_MUSTER, rest)
        menge = mm.group(1).strip() if mm else "1"
        ep_nach = bool(re.match(r"[^+·]*als EP", rest))
        menge = re.sub(r"\s*als EP.*$", "", menge).strip() or "1"
        refs.append(ArtikelRef(f"KL{m.group(1)}", menge, bool(m.group(2)) or ep_nach))
    return refs


_ATTRAPPE = None


def attrappe():
    """Logik-Sicht KL nach der vereinbarten Schnittstelle, Werte aus dem Entwurf."""
    global _ATTRAPPE
    if _ATTRAPPE is not None:
        return copy.copy(_ATTRAPPE)
    wb = openpyxl.load_workbook(ENTWURF, data_only=True)
    bericht = logik_modul.Pruefbericht()
    fragen = logik_modul._fragen_einlesen(wb, bericht, "Fragen KL")
    for fid in ("KE02", "KE03"):          # „nur wenn KE01 ≠ Ja“ → ausgeschrieben
        fragen[fid].bedingung = bedingung_parsen("nur wenn KE01 = Nein oder Unklar")

    aktionen = []
    for row in wb["Aktionen KL"].iter_rows(min_row=2, values_only=True):
        frage, antwort, aktion_roh, bemerkung, zusatz = (
            _z(v) for v in (tuple(row) + (None,) * 5)[:5])
        if not frage:
            continue
        if aktion_roh.startswith("AMPEL"):
            m = re.search(r"Grund:\s*(.+)$", aktion_roh)
            typ, grund = "ampel", (m.group(1).strip() if m else aktion_roh)
        elif re.match(r"(Fachlicher Hinweis|Hinweis|Protokollhinweis)", aktion_roh):
            typ, grund = "hinweis", ""
        else:
            typ, grund = "normal", ""
        zusatz_b = None
        if zusatz:
            zusatz_b = bedingung_parsen(ZUSATZ_UMSCHREIBEN.get(zusatz, zusatz))
        aktionen.append(Aktion(frage, antwort, aktion_roh, typ, grund,
                               kl_refs(aktion_roh), bemerkung, zusatz_b))

    bloecke = [AngebotsBlock(nr, BLOCK_UEBERSCHRIFT[nr], inhalt, Bedingung("immer", "immer"),
                             kl_refs(inhalt)) for nr, inhalt in BLOCK_INHALT.items()]

    typen = {"Single": "single", "Multi-Innengerät": "multi_innen",
             "Multi-Außengerät": "multi_aussen", "Multi": "multi_aussen"}
    paket = []
    for row in wb["Paketmatrix KL"].iter_rows(min_row=2, values_only=True):
        serie, typ, klasse, kuehllast, artikel, bemerkung = (
            _z(v) for v in (tuple(row) + (None,) * 6)[:6])
        if not serie:
            continue
        paket.append(SimpleNamespace(
            serie=serie, typ=typen.get(typ, typ.lower()), klasse=klasse,
            kuehllast=kuehllast, artikel=re.findall(r"KL\d{3}", artikel),
            bemerkung=bemerkung, nicht_im_sortiment="nicht im Sortiment" in artikel))

    kombis: dict = {}
    kombi_artikel: dict = {}
    for row in wb["Kombinationen KL"].iter_rows(min_row=2, values_only=True):
        name, art, anzahl, kombi = (_z(v) for v in (tuple(row) + (None,) * 4)[:4])
        if not name or not kombi:
            continue
        codes = sorted(str(kombi).replace(" ", "").split("+"), key=int)
        kombis.setdefault(name, {}).setdefault(int(float(anzahl)), set()).add("+".join(codes))
        kombi_artikel[name] = art

    montage = {}
    for row in wb["Montagematrix KL"].iter_rows(min_row=2, values_only=True):
        innen, aussen, pos = (_z(v) for v in (tuple(row) + (None,) * 3)[:3])
        if innen.isdigit() and aussen.isdigit() and pos.startswith("KL"):
            montage[(int(innen), int(aussen))] = pos

    parameter, einheit = {}, {}
    for row in wb["KL-Parameter"].iter_rows(min_row=2, values_only=True):
        name, wert, einh, bemerkung = (_z(v) for v in (tuple(row) + (None,) * 4)[:4])
        if name:
            parameter[name] = (wert, bemerkung)
            einheit[name] = einh

    seiten = []
    for f in sorted(fragen.values(), key=lambda f: f.reihenfolge):
        if f.seite and f.seite not in seiten:
            seiten.append(f.seite)
    _ATTRAPPE = SimpleNamespace(
        fragen=fragen, aktionen=aktionen, pakete=[], bloecke=bloecke, kfw={},
        geladen_am=None, anhaenge=[], sparten_fragen={}, vermerke=[], pv_aktionen=[],
        pv_bloecke=[], pv_parameter={}, pv_parameter_einheit={}, pv_kombis=[],
        bafa_anlagen=[], sparte="KL", seiten=seiten,
        kl_aktionen=aktionen, kl_bloecke=bloecke, kl_paket=paket, kl_kombis=kombis,
        kl_kombi_artikel=kombi_artikel, kl_montage=montage, kl_parameter=parameter,
        kl_parameter_einheit=einheit, kl_artikel={})
    return copy.copy(_ATTRAPPE)


_ARTIKEL = None


def artikel_attrappe() -> dict:
    """Artikelstamm KL001–KL050 aus dem Blatt „KL-Artikel“ des Entwurfs (Preise =
    TAIFUN-Listen) – für positionen_bauen ohne Datenbank."""
    global _ARTIKEL
    if _ARTIKEL is not None:
        return dict(_ARTIKEL)
    wb = openpyxl.load_workbook(ENTWURF, data_only=True)
    artikel = {}
    for row in wb["KL-Artikel"].iter_rows(min_row=2, values_only=True):
        nr, _datei, _pos, guid, bez, einheit, vk, ek, _ep = (tuple(row) + (None,) * 9)[:9]
        if not nr or not str(nr).startswith("KL"):
            continue
        try:
            ek_cent = int((Decimal(str(ek)) * 100).quantize(Decimal(1)))
        except Exception:
            ek_cent = None
        guid_text = _z(guid)
        artikel[str(nr)] = SimpleNamespace(
            pos_nr=str(nr), bezeichnung=_z(bez), beschreibung=_z(bez),
            einheit=_z(einheit) or "Stück",
            e_preis_cent=int((Decimal(str(vk)) * 100).quantize(Decimal(1))),
            ek_cent=ek_cent, ep_flag=False,
            guid=guid_text if guid_text.startswith("{") and "?" not in guid_text else None,
            aktiv=True, quelle="kl", kategorie="Klima")
    _ARTIKEL = artikel
    return dict(artikel)


def echte_logik():
    """KL-Sicht der Live-Logik, sobald Agent A/B fertig sind (sonst None)."""
    try:
        from app.db import SessionLocal, init_db
        init_db()
        s = SessionLocal()
        try:
            voll, _bericht = logik_modul.neu_einlesen(s)
        finally:
            s.close()
        kl = logik_modul.logik_fuer_sparte(voll, "KL")
        if (kl is not None and getattr(kl, "kl_aktionen", None)
                and getattr(kl, "kl_paket", None) and "KR08" in kl.fragen):
            return kl
    except Exception:
        return None
    return None


# --- Antworten der Kontrollfälle -------------------------------------------------

def raum(name="Raum", flaeche=20, hoehe="bis 2,5 m", last="normal", lage="Erdgeschoss",
         zweck="hauptsächlich kühlen", ort="Außenwand", entfernung="bis 5 m", pumpe="Nein",
         durchbruch="Außenwand bis 32 cm (in Pauschale)", zuordnung="1") -> dict:
    return {"KR01": name, "KR08": flaeche, "KR09": hoehe, "KR10": last, "KR02": lage,
            "KR03": zweck, "KR04": ort, "KR05": entfernung, "KR06": pumpe,
            "KR11": durchbruch, "KR07": zuordnung}


def kl_basis(raeume: list[dict], **zusatz) -> dict:
    """Vollständiger KL-Bogen: KO01 EFH, KO02 1995, KO03 Ja, 3200i, 1 Außengerät,
    KA01 Außenwand, KA02 unter 2 m, KE01 Nein / KE02 0–5 m / KE03 Ja, vom Boden."""
    antworten = {"KO01": "EFH", "KO02": 1995, "KO03": "Ja",
                 "KO06": "Climate 3200i (Standard)", "KO04": "1", "KO05": float(len(raeume)),
                 "KO07": "Nein", "KO08": "als Eventualposition", "KO09": "",
                 "KA01": "an der Außenwand", "KA02": "unter 2 m",
                 "KE01": "Nein", "KE02": "0–5 m", "KE03": "Ja",
                 "KZ01": "vom Boden/Leiter", "S01": "warm", "S02": ""}
    for i, r in enumerate(raeume, 1):
        for key, wert in r.items():
            antworten[f"{key}#{i}"] = wert
    antworten.update(zusatz)
    if antworten.get("KE01") == "Ja":
        antworten.pop("KE02", None)
        antworten.pop("KE03", None)
    return antworten


def fall_a() -> dict:
    return kl_basis([raum("Wohnzimmer", 25), raum("Schlafzimmer", 18), raum("Büro", 28)])


def fall_b() -> dict:
    return kl_basis([raum("Wohnküche", 30, "2,5–3 m", STARK, entfernung="6–10 m", pumpe="Ja",
                          durchbruch="Beton oder dicker als 32 cm")],
                    KO07="Ja", KO08="Nein", KA02="2–4 m", KE01="Ja",
                    KZ01="Gerüst erforderlich")


def fall_c() -> dict:
    return kl_basis([raum("Wohnen", 50), raum("Essen", 55, last=STARK),
                     raum("Kind 1", 20, zuordnung="2"), raum("Kind 2", 20, zuordnung="2"),
                     raum("Büro", 20, zuordnung="2")], KO04="2")


def fall_h() -> dict:
    return kl_basis([raum(f"Raum {i}", 40) for i in range(1, 5)],
                    KA01="auf einem Schrägdach", KZ01="Gerüst erforderlich")


def mengen(positionen) -> dict:
    ergebnis: dict = {}
    for p in positionen:
        if p["pos_nr"]:
            ergebnis[p["pos_nr"]] = ergebnis.get(p["pos_nr"], 0) + p["menge"]
    return ergebnis


def refs(gewaehlt) -> dict:
    ergebnis: dict = {}
    for g in gewaehlt:
        ergebnis[g.ref] = ergebnis.get(g.ref, 0) + g.menge
    return ergebnis


def netto_cent(positionen) -> int:
    return int(round(sum(p["menge"] * p["e_preis_cent"] for p in positionen
                         if not p["ep_flag"])))


# --- Rechenweg (Attrappe) ------------------------------------------------------------

class RechenwegAttrappe(unittest.TestCase):
    echt = False

    @classmethod
    def setUpClass(cls):
        cls.logik = attrappe()
        cls.artikel = artikel_attrappe()

    def positionen(self, antworten):
        return kl_auslegung.positionen_bauen(self.logik, antworten, self.artikel)

    # -- Grundlagen ---------------------------------------------------------------
    def test_bogen_vollstaendig(self):
        for antworten in (fall_a(), fall_b(), fall_c(), fall_h()):
            self.assertIsNone(engine.naechste_frage(self.logik, antworten))

    def test_parameter(self):
        p = kl_auslegung.parameter(self.logik)
        self.assertEqual(p["fehlende"], [])
        self.assertEqual((p["w_qm_normal"], p["w_qm_stark"]), (60.0, 90.0))
        self.assertEqual(p["hoehenfaktor"], {"bis 2,5 m": 1.0, "2,5–3 m": 1.1, "über 3 m": 1.2})
        self.assertEqual(p["klassen"], [("9", 2.6), ("12", 3.5), ("18", 5.3), ("24", 7.0)])
        self.assertEqual(p["standardserie"], "Climate 3200i (Standard)")
        self.assertEqual(p["p14a_geraete"], ["CL5000M 105/4 E", "CL5000M 125/5 E"])
        self.assertEqual((p["max_innen_je_aussen"], p["max_aussen"]), (5, 3))
        self.assertEqual((p["leitung_inklusive_m"], p["meterposition"]), (5.0, "KL017"))
        self.assertEqual(p["gewerbe_verhalten"], "Hinweis")
        self.assertEqual(p["rollgeruest_vk"], 499.0)
        # ohne Logik: Standardwerte + alle Pflichtzeilen als fehlend
        leer = kl_auslegung.parameter(None)
        self.assertEqual(len(leer["fehlende"]), len(kl_auslegung.PFLICHT_KL_PARAMETER))
        self.assertEqual(leer["klassen"], kl_auslegung.STANDARD_KLASSEN)

    def test_kuehllast_und_klasse(self):
        self.assertEqual(kl_auslegung.kuehllast(25, 1.0, 60), 1.5)
        self.assertEqual(kl_auslegung.kuehllast(18, 1.0, 60), 1.08)
        self.assertEqual(kl_auslegung.kuehllast(28, 1.0, 60), 1.68)
        self.assertEqual(kl_auslegung.kuehllast(30, 1.1, 90), 2.97)
        self.assertEqual(kl_auslegung.kuehllast(100, 1.2, 90), 10.8)
        klassen = kl_auslegung.STANDARD_KLASSEN
        self.assertEqual([kl_auslegung.klasse_fuer(kw, klassen)
                          for kw in (1.5, 2.6, 2.97, 3.5, 4.95, 5.3, 7.0, 7.01)],
                         ["9", "9", "12", "12", "18", "18", "24", None])

    # -- Fall A -------------------------------------------------------------------
    def test_fall_a_auslegung(self):
        a = kl_auslegung.auslegen(self.logik, fall_a())
        self.assertEqual([r.kuehllast_kw for r in a.raeume], [1.5, 1.08, 1.68])
        self.assertEqual([r.klasse for r in a.raeume], ["9", "9", "9"])
        self.assertEqual([r.artikel for r in a.raeume], [["KL026"]] * 3)
        g = a.aussengeraete[0]
        self.assertEqual((g.typ, g.kombination, g.artikel, g.geraet),
                         ("multi", "9+9+9", ["KL038"], "CL5000M 79/3 E"))
        self.assertEqual((a.innengeraete_gesamt, a.aussengeraete_gesamt, a.montage_ref),
                         (3, 1, "KL023"))
        self.assertEqual(a.ampeln, [])
        self.assertEqual(a.hinweise, [])
        d = a.als_dict()
        self.assertEqual(d["aussengeraete"][0]["kombination"], "9+9+9")
        self.assertEqual(d["raeume"][0]["kuehllast_kw"], 1.5)
        json.dumps(d)   # JSON-fähig (kl_json)

    def test_fall_a_artikel(self):
        a = kl_auslegung.auslegen(self.logik, fall_a())
        gewaehlt = kl_auslegung.artikel_ermitteln(self.logik, fall_a(), a)
        self.assertEqual(refs(gewaehlt), {"KL008": 1, "KL049": 1, "KL038": 1, "KL026": 3,
                                          "KL013": 3, "KL023": 1})
        self.assertTrue(all(g.ep for g in gewaehlt if g.ref == "KL013"))
        self.assertFalse(any(g.ep for g in gewaehlt if g.ref != "KL013"))

    def test_fall_a_positionen(self):
        pos = self.positionen(fall_a())
        self.assertEqual([p["pos_nr"] for p in pos],
                         ["KL023", "KL008", "KL038", "KL026", "KL013", "KL049", ""])
        self.assertEqual(mengen(pos), {"KL023": 1, "KL008": 1, "KL038": 1, "KL026": 3,
                                       "KL013": 3, "KL049": 1})
        nach_nr = {p["pos_nr"]: p for p in pos}
        self.assertEqual((nach_nr["KL023"]["e_preis_cent"], nach_nr["KL008"]["e_preis_cent"],
                          nach_nr["KL038"]["e_preis_cent"], nach_nr["KL026"]["e_preis_cent"],
                          nach_nr["KL013"]["e_preis_cent"], nach_nr["KL049"]["e_preis_cent"]),
                         (317800, 49800, 196624, 29214, 8250, 0))
        self.assertTrue(nach_nr["KL013"]["ep_flag"])
        self.assertFalse(nach_nr["KL026"]["ep_flag"])
        self.assertEqual(netto_cent(pos), 651866)                    # 6.518,66 €
        self.assertEqual(int(round(651866 * 19 / 100)), 123855)       # USt 1.238,55 €
        # Blöcke und Gruppe „Klimaanlage Bosch“ (Block 3)
        self.assertEqual([p["block_nr"] for p in pos], [1, 2, 3, 3, 3, 3, 3])
        self.assertEqual({p["gruppe"] for p in pos if p["block_nr"] == 3}, {"Klimaanlage Bosch"})
        self.assertEqual({p["gruppe"] for p in pos if p["block_nr"] < 3}, {""})
        # Auslegungszeile als LETZTE Zeile von Block 3
        letzte = pos[-1]
        self.assertEqual((letzte["bezeichnung"], letzte["e_preis_cent"], letzte["menge"],
                          letzte["einheit"], letzte["ep_flag"], letzte["block_nr"]),
                         ("Auslegung der Klimaanlage", 0, 1.0, "psl.", False, 3))
        text = letzte["beschreibung"]
        self.assertEqual(text.splitlines()[0], "Auslegung Klimaanlage – Serie Bosch Climate 3200i")
        self.assertIn("Raum 1 „Wohnzimmer“: 25,0 m² × 1,0 × 60 W/m² = 1,50 kW → Innengerät "
                      "2,6 kW (Klasse 9), Außengerät 1", text)
        self.assertIn("Raum 2 „Schlafzimmer“: 18,0 m² × 1,0 × 60 W/m² = 1,08 kW", text)
        self.assertIn("Außengerät 1: BOSCH CL5000M 79/3 E Multisplit-Außeneinheit 7,9 kW, "
                      "Multi-Split, Kombination 9+9+9, 3 Innengeräte", text)
        self.assertIn("Montage: 3 Innengeräte / 1 Außengerät (Pauschale). Leitungslänge bis "
                      "5 m je Innengerät enthalten.", text)
        self.assertTrue(text.endswith("Auslegung nach Raumfläche; verbindliche Festlegung in "
                                      "der technischen Feinplanung vor Ort."))
        self.assertEqual([p["sort"] for p in pos], list(range(1, 8)))

    # -- Fall B -------------------------------------------------------------------
    def test_fall_b(self):
        a = kl_auslegung.auslegen(self.logik, fall_b())
        r = a.raeume[0]
        self.assertEqual((r.kuehllast_kw, r.klasse, r.hoehenfaktor, r.w_qm), (2.97, "12", 1.1, 90.0))
        g = a.aussengeraete[0]
        self.assertEqual((g.typ, g.artikel, g.kombination), ("single", ["KL035"], ""))
        self.assertEqual((a.montage_ref, a.ampeln), ("KL020", []))
        pos = self.positionen(fall_b())
        self.assertEqual([p["pos_nr"] for p in pos],
                         ["KL020", "KL019", "KL016", "KL014", "KL017", "KL050", "KL001",
                          "KL035", "KL049", ""])
        m = mengen(pos)
        self.assertEqual(m["KL017"], 5)
        self.assertNotIn("KL013", m)
        self.assertNotIn("KL008", m)
        self.assertEqual(netto_cent(pos), 549500)                    # 5.495,00 €
        self.assertEqual(int(round(549500 * 19 / 100)), 104405)
        self.assertIn("Außengerät 1: Bosch CL3200i-Set 35 WE Single-Split, Single-Split, "
                      "1 Innengerät", pos[-1]["beschreibung"])
        # KE02/KE03 unsichtbar (KE01 = Ja)
        sichtbar = {f.id for f in engine.sichtbare_fragen(self.logik, fall_b())}
        self.assertNotIn("KE02", sichtbar)
        self.assertNotIn("KE03", sichtbar)

    # -- Fall C -------------------------------------------------------------------
    def test_fall_c_montagematrix_ampel(self):
        a = kl_auslegung.auslegen(self.logik, fall_c())
        self.assertEqual([r.klasse for r in a.raeume], ["12", "18", "9", "9", "9"])
        self.assertEqual([r.kuehllast_kw for r in a.raeume], [3.0, 4.95, 1.2, 1.2, 1.2])
        g1, g2 = a.aussengeraete
        self.assertEqual((g1.kombination, g1.artikel, g1.geraet),
                         ("12+18", ["KL037"], "CL5000M 53/2 E"))
        self.assertEqual((g2.kombination, g2.artikel, g2.geraet),
                         ("9+9+9", ["KL038"], "CL5000M 79/3 E"))
        self.assertEqual(a.montage_ref, "")
        self.assertEqual(a.ampeln, ["Montagekombination 5 Innengeräte / 2 Außengeräte "
                                    "nicht hinterlegt"])
        self.assertEqual(engine.ampel_gruende(self.logik, fall_c()), a.ampeln)
        prot = kl_auslegung.protokoll_zeilen(self.logik, fall_c())
        self.assertEqual({z["seite"] for z in prot}, {"Auslegung"})
        texte = " | ".join(z["antwort"] for z in prot)
        self.assertIn("Kombination 12+18", texte)
        self.assertIn("Kombination 9+9+9", texte)
        montage = [z for z in prot if z["frage"] == "Montage"][0]
        self.assertEqual(montage["ampel_grund"], a.ampeln[0])
        self.assertIn("nicht hinterlegt", montage["antwort"])

    # -- Fall D -------------------------------------------------------------------
    def test_fall_d_serien(self):
        d1 = kl_basis([raum("A", 20), raum("B", 20)], KO06="Climate 8000i")
        self.assertEqual(kl_auslegung.auslegen(self.logik, d1).ampeln,
                         ["Multi-Split in Serie Climate 8000i nicht verfügbar"])
        d2 = kl_basis([raum("Wohnen", 50)], KO06="Climate 7000i", KO08="Ja")
        a = kl_auslegung.auslegen(self.logik, d2)
        self.assertEqual((a.raeume[0].kuehllast_kw, a.raeume[0].klasse), (3.0, "12"))
        self.assertEqual((a.aussengeraete[0].artikel, a.ampeln), (["KL044", "KL042"], []))
        pos = self.positionen(d2)
        self.assertEqual(mengen(pos), {"KL020": 1, "KL008": 1, "KL044": 1, "KL042": 1,
                                       "KL013": 1, "KL049": 1})
        gateway = [p for p in pos if p["pos_nr"] == "KL013"][0]
        self.assertFalse(gateway["ep_flag"])                        # KO08 = Ja: kein EP
        self.assertEqual(gateway["e_preis_cent"], 8250)
        reihenfolge = [p["pos_nr"] for p in pos]
        self.assertLess(reihenfolge.index("KL044"), reihenfolge.index("KL042"))
        self.assertLess(reihenfolge.index("KL042"), reihenfolge.index("KL013"))
        d3 = kl_basis([raum("Wohnen", 70)], KO06="Climate 7000i")
        a = kl_auslegung.auslegen(self.logik, d3)
        self.assertEqual((a.raeume[0].kuehllast_kw, a.raeume[0].klasse), (4.2, "18"))
        self.assertEqual(a.ampeln, ["Klasse 18 in Serie Climate 7000i nicht verfügbar – "
                                    "Serie wechseln"])

    # -- Fall E / F ---------------------------------------------------------------
    def test_fall_e_kuehllast(self):
        e = kl_basis([raum("Dachstudio", 100, "über 3 m", STARK, lage="Dachgeschoss")])
        a = kl_auslegung.auslegen(self.logik, e)
        self.assertEqual((a.raeume[0].kuehllast_kw, a.raeume[0].klasse), (10.8, None))
        self.assertEqual(a.ampeln, ["Kühllast über 7,0 kW: Raum 1 Dachstudio (10,80 kW) – "
                                    "Raum teilen oder Sonderlösung"])
        self.assertEqual(engine.ampel_gruende(self.logik, e), a.ampeln)

    def test_fall_f_kombination(self):
        f = kl_basis([raum("Wohnen", 60, last=STARK), raum("Essen", 60, last=STARK)])
        a = kl_auslegung.auslegen(self.logik, f)
        self.assertEqual([r.klasse for r in a.raeume], ["24", "24"])
        self.assertEqual(a.aussengeraete[0].kombination, "24+24")
        self.assertEqual(a.ampeln, ["Multi-Split-Kombination 24+24 an Außengerät 1 nicht "
                                    "freigegeben (Bosch Tab. 7)"])

    def test_g5_mehr_als_fuenf(self):
        a = kl_auslegung.auslegen(self.logik, kl_basis([raum(f"R{i}", 15) for i in range(6)]))
        self.assertIn("Mehr als 5 Innengeräte an Außengerät 1", a.ampeln)

    # -- Fall G -------------------------------------------------------------------
    def test_fall_g_validierung(self):
        g = kl_basis([raum("A", 20), raum("B", 20), raum("C", 20)], KO04="2")
        self.assertEqual(kl_auslegung.validierung(self.logik, g),
                         ["Außengerät 2 hat keinen zugeordneten Raum – Zuordnung (KR07) "
                          "prüfen oder Anzahl Außengeräte anpassen."])
        self.assertIn(kl_auslegung.validierung(self.logik, g)[0],
                      engine.ampel_gruende(self.logik, g))
        g["KR07#3"] = "2"
        self.assertEqual(kl_auslegung.validierung(self.logik, g), [])
        a = kl_auslegung.auslegen(self.logik, g)
        self.assertEqual((a.montage_ref, a.ampeln), ("KL024", []))      # 3/2 → KL024
        # Raumfläche fehlt / 0 / Komma-Eingabe
        g["KR08#2"] = 0
        self.assertEqual(kl_auslegung.validierung(self.logik, g), ["Raum 2: Raumfläche fehlt"])
        g["KR08#2"] = "22,5"
        self.assertEqual(kl_auslegung.validierung(self.logik, g), [])
        self.assertEqual(kl_auslegung.auslegen(self.logik, g).raeume[1].kuehllast_kw, 1.35)

    # -- Fall H -------------------------------------------------------------------
    def test_fall_h_schraegdach_14a(self):
        a = kl_auslegung.auslegen(self.logik, fall_h())
        g = a.aussengeraete[0]
        self.assertEqual((g.kombination, g.artikel, g.geraet, g.p14a),
                         ("9+9+9+9", ["KL039"], "CL5000M 105/4 E", True))
        self.assertEqual((a.montage_ref, a.ampeln), ("KL025", []))
        m = mengen(self.positionen(fall_h()))
        self.assertEqual((m.get("KL004"), m.get("KL005"), m.get("KL006")), (1, 1, 1))
        self.assertNotIn("KL050", m)
        self.assertEqual((m.get("KL003"), m.get("KL025"), m.get("KL039"), m.get("KL026")),
                         (1, 1, 1, 4))
        # ohne Schrägdach: Rollgerüst statt Dacharbeiten
        m2 = mengen(self.positionen(kl_basis([raum(f"Raum {i}", 40) for i in range(1, 5)],
                                             KZ01="Gerüst erforderlich")))
        self.assertEqual(m2.get("KL050"), 1)
        self.assertNotIn("KL004", m2)
        # §14a-Liste ohne 105/4 → kein KL003
        logik2 = copy.copy(self.logik)
        logik2.kl_parameter = dict(self.logik.kl_parameter)
        logik2.kl_parameter["§14a-Außengeräte"] = ("CL5000M 125/5 E", "")
        m3 = mengen(kl_auslegung.positionen_bauen(logik2, fall_h(), self.artikel))
        self.assertNotIn("KL003", m3)
        self.assertEqual(m3.get("KL039"), 1)

    # -- G7 / G8, Hinweise ------------------------------------------------------------
    def test_g7_g8_leitungslaengen(self):
        antworten = fall_a()
        antworten["KR05#2"] = "über 15 m"
        antworten["KE02"] = "über 15 m"
        gruende = engine.ampel_gruende(self.logik, antworten)
        self.assertIn("Leitungslänge über 15 m: Raum 2 Schlafzimmer – max. Leitungslänge der "
                      "Serie prüfen", gruende)
        self.assertIn("Elektrozuleitung über 15 m", gruende)
        # Zusatzmeter: 6–10 m → KL017 × 5 je Raum, 11–15 m → × 10; Elektro ebenso
        antworten = fall_a()
        antworten["KR05#1"] = "6–10 m"
        antworten["KR05#3"] = "11–15 m"
        antworten["KE02"] = "6–10 m"
        self.assertEqual(mengen(self.positionen(antworten)).get("KL017"), 20)

    def test_fachliche_hinweise(self):
        antworten = kl_basis([raum("Wohnen", 20, ort="Innenwand"),
                              raum("Schlafen", 20, zweck="regelmäßig heizen und kühlen")],
                             KO01="Gewerbe", KO03="Nein, Zustimmung noch nicht vorhanden",
                             KA01="auf einem Flachdach", KE01="Unklar")
        hinweise = kl_auslegung.fachliche_hinweise(self.logik, antworten)
        self.assertEqual(hinweise, [
            "Zustimmung des Eigentümers einholen",
            "Gewerbeobjekt – Auslegung prüfen",
            "Raum 1 Wohnen: Innengerät an Innenwand – Leitungsführung und Kondensatablauf prüfen",
            "Raum 2 Schlafen: Heizbetrieb als Hauptzweck – Serie 7000i/8000i empfohlen",
            "Flachdach – keine Dachdurchführung, Leitungsführung über Attika/Fassade prüfen",
            "Stromversorgung am Außengerät prüfen"])
        self.assertEqual(engine.fachliche_hinweise(antworten, self.logik), hinweise)
        # Flachdach: Hinweis, keine Position KL007, keine Ampel
        a = kl_auslegung.auslegen(self.logik, antworten)
        self.assertEqual(a.ampeln, [])
        self.assertNotIn("KL007", mengen(self.positionen(antworten)))
        h2 = kl_auslegung.fachliche_hinweise(self.logik, kl_basis(
            [raum()], KA01="noch unklar", KA02="über 4 m", KZ01="vom Boden/Leiter"))
        self.assertEqual(h2, ["Montageort Außengerät klären",
                              "Montagehöhe über 4 m ohne Gerüst/Bühne prüfen"])
        self.assertEqual(kl_auslegung.fachliche_hinweise(self.logik, kl_basis([raum()], KO05=13)),
                         ["Mehr als 12 Räume – Rest händisch"])
        # Parameter Gewerbe-Verhalten „AMPEL“: Ampel statt Hinweis
        logik2 = copy.copy(self.logik)
        logik2.kl_parameter = dict(self.logik.kl_parameter)
        logik2.kl_parameter["Gewerbe-Verhalten"] = ("AMPEL", "")
        gewerbe = kl_basis([raum()], KO01="Gewerbe")
        self.assertNotIn("Gewerbeobjekt – Auslegung prüfen",
                         kl_auslegung.fachliche_hinweise(logik2, gewerbe))
        self.assertEqual(kl_auslegung.ampel_gruende(logik2, gewerbe),
                         ["Gewerbeobjekt – individuelle Auslegung"])
        # Hinweise aus dem gespeicherten Protokoll (Anzeige-Strings)
        prot = engine.protokoll(self.logik, antworten)
        self.assertEqual(engine.hinweise_aus_protokoll(prot)[:2],
                         ["Zustimmung des Eigentümers einholen", "Gewerbeobjekt – Auslegung prüfen"])

    # -- Protokoll / Einbindung konfigurator --------------------------------------------
    def test_protokoll_zeilen_fall_a(self):
        prot = kl_auslegung.protokoll_zeilen(self.logik, fall_a())
        self.assertEqual([z["frage"] for z in prot],
                         ["Auslegung Klimaanlage", "Raum 1 Wohnzimmer", "Raum 2 Schlafzimmer",
                          "Raum 3 Büro", "Außengerät 1", "Montage"])
        self.assertEqual(prot[1]["antwort"],
                         "Fläche 25,0 m² · Deckenhöhe bis 2,5 m (Faktor 1,0) · Wärmelast normal "
                         "(60 W/m²) · Kühllast 1,50 kW · Klasse 9 (Innengerät 2,6 kW) · Außengerät 1")
        self.assertEqual(prot[4]["antwort"],
                         "CL5000M 79/3 E · Multi-Split, Kombination 9+9+9 · 3 Innengeräte (Raum 1, 2, 3)")
        self.assertEqual(prot[5]["antwort"], "3 Innengeräte / 1 Außengerät – Pauschale KL023")
        self.assertTrue(all(z["ampel_grund"] == "" and z["frage_id"] == "" for z in prot))
        voll = engine.protokoll(self.logik, fall_a())
        self.assertEqual([z for z in voll if z["seite"] == "Auslegung"], prot)
        self.assertIn("KR08#1", {z["frage_id"] for z in voll})
        # unvollständiger Bogen: keine Auslegungszeilen, keine Gründe
        teil = {k: v for k, v in fall_a().items() if k != "KZ01"}
        self.assertEqual(kl_auslegung.protokoll_zeilen(self.logik, teil), [])
        self.assertEqual(engine.ampel_gruende(self.logik, teil), [])

    def test_vorbelegungen(self):
        f = self.logik.fragen
        self.assertEqual(engine.vorbelegung(f["KO06"], {}, self.logik), "Climate 3200i (Standard)")
        self.assertEqual(engine.vorbelegung(f["KO06"], {}), "Climate 3200i (Standard)")
        self.assertEqual(engine.vorbelegung(f["KO08"], {}), "als Eventualposition")
        klone = {k.id: k for k in engine._wiederhol_klone(f["KR10"], {"KO05": 2}, f)}
        self.assertEqual(engine.vorbelegung(klone["KR10#2"], {"KR02#2": "Dachgeschoss"}), STARK)
        self.assertIsNone(engine.vorbelegung(klone["KR10#2"], {"KR02#2": "Erdgeschoss"}))
        self.assertIsNone(engine.vorbelegung(klone["KR10#1"], {"KR02#2": "Dachgeschoss"}))
        # KR07-Klone zeigen nur so viele Optionen wie Außengeräte (KO04)
        kr07 = engine._wiederhol_klone(f["KR07"], {"KO05": 1, "KO04": "2"}, f)
        self.assertEqual(kr07[0].antworten, ["1", "2"])

    def test_hinweis_aktionen_ohne_position(self):
        hinweise = [a for a in self.logik.kl_aktionen if a.typ == "hinweis"]
        self.assertTrue(hinweise)
        for aktion in hinweise:
            self.assertEqual(aktion.artikel, [])
        self.assertFalse([a for a in self.logik.kl_aktionen
                          if a.frage in ("KA02", "KR04", "KR03") and a.typ == "ampel"])

    def test_ampel_clone_matching_engine(self):
        """Raum-Klone treffen die Aktionszeile ihrer Basisfrage (konfigurator.aktion_finden)."""
        f = self.logik.fragen
        klon = engine._wiederhol_klone(f["KR05"], {"KO05": 1}, f)[0]
        treffer = engine.aktion_finden(self.logik, klon, "über 15 m", {})
        self.assertIsNotNone(treffer)
        self.assertEqual(treffer[0].typ, "ampel")
        treffer = engine.aktion_finden(self.logik, klon, "6–10 m", {})
        self.assertEqual([r.ref for r in treffer[0].artikel], ["KL017"])


# --- dieselben Tests gegen die echte Logik (sobald Agent A/B fertig sind) ----------

class RechenwegEcht(RechenwegAttrappe):
    echt = True

    @classmethod
    def setUpClass(cls):
        logik = echte_logik()
        if logik is None:
            raise unittest.SkipTest("Echte KL-Logik noch nicht verfügbar (Blätter „Aktionen KL“/"
                                    "„Fragen KL“ v24 oder Lader kl_* fehlen) – Attrappe gilt.")
        cls.logik = logik
        cls.artikel = artikel_attrappe()


# --- Datenbank: Angebot anlegen, Router ---------------------------------------------

def _kl_artikel_vorhanden(session) -> bool:
    from app.models import Artikel
    quelle = getattr(__import__("app.models", fromlist=["QUELLE_KL"]), "QUELLE_KL", "kl")
    return (session.query(Artikel).filter(Artikel.quelle == quelle,
                                          Artikel.aktiv.is_(True)).count() >= 49)


def aufraeumen(s):
    from app.models import Angebot, Erfassung, Kunde, Vorgang
    for k in s.query(Kunde).filter(Kunde.email == TEST_EMAIL):
        for a in s.query(Angebot).filter_by(kunde_id=k.id):
            s.delete(a)
        for e in s.query(Erfassung).filter_by(kunde_id=k.id):
            s.delete(e)
        for v in s.query(Vorgang).filter_by(kunde_id=k.id):
            s.delete(v)
        s.delete(k)
    s.commit()


class Einbindung(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from app.db import SessionLocal, init_db
        from app.models import Kunde
        init_db()
        cls.s = SessionLocal()
        aufraeumen(cls.s)
        cls.kunde = Kunde(anrede="Frau", vorname="Klara", nachname="KL-Test v24 Auslegung",
                          strasse="Teststr. 2", plz="47139", ort="Duisburg", email=TEST_EMAIL,
                          objektart="RH")
        cls.s.add(cls.kunde)
        cls.s.commit()
        cls.echt = echte_logik()
        cls.logik = cls.echt or attrappe()

    @classmethod
    def tearDownClass(cls):
        aufraeumen(cls.s)
        cls.s.close()

    def test_models_und_db(self):
        from app import models
        from app.models import Angebot, ust_standard
        self.assertEqual(getattr(models, "QUELLE_KL", None), "kl")
        self.assertTrue(hasattr(Angebot, "kl_json"))
        self.assertEqual(ust_standard("KL"), 19.0)
        from app.db import _NACHTRAEGLICHE_SPALTEN
        self.assertIn("kl_json", _NACHTRAEGLICHE_SPALTEN["angebote"])

    def test_angebot_anlegen_fall_a(self):
        if not _kl_artikel_vorhanden(self.s):
            self.skipTest("KL-Artikel (KL001–KL050) noch nicht importiert – Import Agent A "
                          "(Artikel → Klima-Positionslisten importieren) abwarten.")
        from app import angebot_aufbau
        angebot = angebot_aufbau.angebot_anlegen(self.s, self.kunde.id, antworten=fall_a(),
                                                 logik=self.logik, sparte="KL")
        self.assertEqual((angebot.konfigurator_typ, angebot.ust_satz, angebot.kfw_json),
                         ("KL", 19.0, "{}"))
        su = angebot.summen()
        self.assertEqual(su["netto"], 651866)
        # Toleranz ± 0,01 € (PLAN_V16 Phase 116): Angebot.summen() schneidet die USt
        # ab (int), das Musterangebot rundet kaufmännisch (1.238,55 / 7.757,21)
        self.assertAlmostEqual(su["ust"], 123855, delta=1)
        self.assertAlmostEqual(su["brutto"], 775721, delta=1)
        self.assertEqual([p.pos_nr for p in angebot.positionen],
                         ["KL023", "KL008", "KL038", "KL026", "KL013", "KL049", ""])
        kl = json.loads(angebot.kl_json)
        self.assertEqual(kl["aussengeraete"][0]["kombination"], "9+9+9")
        self.assertEqual(kl["montage_ref"], "KL023")
        self.assertEqual({p.gruppe for p in angebot.positionen if p.block_nr == 3},
                         {"Klimaanlage Bosch"})
        self.assertEqual(angebot.positionen[-1].bezeichnung, "Auslegung der Klimaanlage")
        # keine WP-Artikel 014–017 (positionsregeln_anwenden, Agent D)
        self.assertFalse({"014", "015", "016", "017"} & {p.pos_nr for p in angebot.positionen})
        # Fall B
        b = angebot_aufbau.angebot_anlegen(self.s, self.kunde.id, antworten=fall_b(),
                                           logik=self.logik, sparte="KL")
        self.assertEqual(b.summen()["netto"], 549500)

    def test_sparten_start_objektart_ko01(self):
        """KO01 aus kunden.objektart (RH → RMH) – Mapping im Router, lead_kartei unverändert."""
        from fastapi.testclient import TestClient

        from app.main import app
        from app.models import Erfassung
        client = TestClient(app)
        client.post("/login", data={"benutzer_id": "1", "pin": "1234"})
        r = client.post("/erfassung/sparten-start",
                        data={"kunde_id": str(self.kunde.id), "sparte_KL": "on"},
                        follow_redirects=False)
        self.assertEqual(r.status_code, 303)
        e = (self.s.query(Erfassung).filter(Erfassung.kunde_id == self.kunde.id,
                                            Erfassung.sparte == "KL")
             .order_by(Erfassung.id.desc()).first())
        self.assertIsNotNone(e)
        antworten = json.loads(e.antworten_json or "{}")
        self.assertEqual(antworten.get("KO01"), "RMH")
        self.assertNotIn("O01", antworten)
        self.assertIn("KO01", json.loads(e.vorbelegt_json or "{}"))

    def test_seite_raeume_mit_raumueberschrift(self):
        """Bogenseite „Räume“: Zwischenüberschrift „Raum <i>“ je Raumblock, KR08 als
        Dezimaleingabe (inputmode=decimal, step=0.1), Dezimalkomma beim Speichern."""
        from fastapi.testclient import TestClient

        from app.main import app
        from app.models import Erfassung
        logik_voll, _ = logik_modul.hole_logik(self.s)
        kl = logik_modul.logik_fuer_sparte(logik_voll, "KL")
        if kl is None or "KR08" not in kl.fragen or "Räume" not in kl.seiten:
            self.skipTest("Blatt „Fragen KL“ (v24) noch nicht in der Live-Excel.")
        client = TestClient(app)
        client.post("/login", data={"benutzer_id": "1", "pin": "1234"})
        e = Erfassung(kunde_id=self.kunde.id, benutzer_id=1, sparte="KL",
                      konfigurator_typ="KL",
                      antworten_json=json.dumps({"KO01": "EFH", "KO02": 1995, "KO03": "Ja",
                                                 "KO06": "Climate 3200i (Standard)",
                                                 "KO04": "2", "KO05": 2, "KO07": "Nein",
                                                 "KO08": "als Eventualposition", "KO09": ""}))
        self.s.add(e)
        self.s.commit()
        nr = kl.seiten.index("Räume")
        r = client.get(f"/erfassung/{e.id}/seite/{nr}")
        self.assertEqual(r.status_code, 200)
        self.assertIn('<h2 class="raum-titel">Raum 1</h2>', r.text)
        self.assertIn('<h2 class="raum-titel">Raum 2</h2>', r.text)
        self.assertEqual(r.text.count('class="raum-titel"'), 2)
        self.assertIn('name="f_KR08#1" inputmode="decimal"', r.text)
        self.assertIn('step="0.1"', r.text)
        # Seite speichern mit Dezimalkomma; KR07 nur Optionen 1 und 2 (KO04 = 2)
        self.assertIn('name="f_KR07#1" value="2"', r.text)
        self.assertNotIn('name="f_KR07#1" value="3"', r.text)
        daten = {"richtung": "weiter"}
        for i, (name, flaeche, zu) in enumerate((("Wohnzimmer", "22,5", "1"),
                                                 ("Küche", "0", "2")), 1):
            daten.update({f"f_KR01#{i}": name, f"f_KR08#{i}": flaeche,
                          f"f_KR09#{i}": "bis 2,5 m", f"f_KR10#{i}": "normal",
                          f"f_KR02#{i}": "Erdgeschoss", f"f_KR03#{i}": "hauptsächlich kühlen",
                          f"f_KR04#{i}": "Außenwand", f"f_KR05#{i}": "bis 5 m",
                          f"f_KR06#{i}": "Nein",
                          f"f_KR11#{i}": "Außenwand bis 32 cm (in Pauschale)",
                          f"f_KR07#{i}": zu})
        r = client.post(f"/erfassung/{e.id}/seite/{nr}", data=daten, follow_redirects=False)
        self.assertEqual(r.status_code, 200)                 # Raum 2: Fläche 0 → Fehler
        self.assertIn("Raum 2: Raumfläche fehlt", r.text)
        daten["f_KR08#2"] = "18"
        r = client.post(f"/erfassung/{e.id}/seite/{nr}", data=daten, follow_redirects=False)
        self.assertEqual(r.status_code, 303)
        self.s.refresh(e)
        antworten = json.loads(e.antworten_json)
        self.assertEqual((antworten["KR08#1"], antworten["KR08#2"]), (22.5, 18.0))
        # Vorbelegung KO06/KO08 auf der Seite „Objekt & Anlage“ (leere Erfassung)
        leer = Erfassung(kunde_id=self.kunde.id, benutzer_id=1, sparte="KL",
                         konfigurator_typ="KL")
        self.s.add(leer)
        self.s.commit()
        r = client.get(f"/erfassung/{leer.id}/seite/{kl.seiten.index('Objekt & Anlage')}")
        self.assertRegex(r.text, r'name="f_KO06" value="Climate 3200i \(Standard\)"\s+checked')
        self.assertRegex(r.text, r'name="f_KO08" value="als Eventualposition"\s+checked')

    def test_absenden_validierung_fall_g(self):
        """Absenden blockiert mit der Meldung aus Phase 114; nach Korrektur grün/„Neu“."""
        if self.echt is None:
            self.skipTest("Echte KL-Logik noch nicht verfügbar – Router-Test folgt mit Agent A/B.")
        from fastapi.testclient import TestClient

        from app.main import app
        from app.models import Erfassung
        client = TestClient(app)
        client.post("/login", data={"benutzer_id": "1", "pin": "1234"})
        g = kl_basis([raum("A", 20), raum("B", 20), raum("C", 20)], KO04="2")
        e = Erfassung(kunde_id=self.kunde.id, benutzer_id=1, sparte="KL",
                      konfigurator_typ="KL", antworten_json=json.dumps(g, ensure_ascii=False))
        self.s.add(e)
        self.s.commit()
        r = client.post(f"/erfassung/{e.id}/absenden", follow_redirects=False)
        self.assertEqual(r.status_code, 200)
        self.assertIn("Außengerät 2 hat keinen zugeordneten Raum", r.text)
        self.s.refresh(e)
        self.assertEqual(e.status, "Entwurf")
        g["KR07#3"] = "2"
        e.antworten_json = json.dumps(g, ensure_ascii=False)
        self.s.commit()
        r = client.post(f"/erfassung/{e.id}/absenden", follow_redirects=False)
        self.assertEqual(r.status_code, 200)
        self.s.refresh(e)
        self.assertEqual((e.ampel, e.status), ("gruen", "Neu"))
        # Prüfseite zeigt die Auslegung
        r = client.get(f"/erfassung/{e.id}/pruefen")
        self.assertIn("Auslegung", r.text)


if __name__ == "__main__":
    unittest.main()
