# Tests PLAN_PROJ_V3 (Phasen 84–86): Auftragsdaten für TAIFUN-Aufträge,
# Bestandsimport, Go-live-Hilfen. Laufen gegen die Entwicklungs-DB und räumen
# ihre Testdaten (Kunden mit E-Mail projekt-v3@test.local) wieder auf.
import io
import unittest
import warnings
from datetime import datetime

warnings.filterwarnings("ignore")

from fastapi.testclient import TestClient

from app import projektierung as kern
from app import projektierung_logik
from app.db import SessionLocal, init_db
from app.main import app
from app.models import (Angebot, Aufgabe, AufgabenpaketInstanz, Gewerk, Kunde,
                        Projekt, ProjektDokument, ProjektTermin, ProjektVerlauf,
                        SteckbriefWert, Vorgang, angebot_status_setzen)
from tests.test_projektierung_v4 import aufraeumen as _aufraeumen_v4

TEST_EMAIL = "projekt-v3@test.local"
MINI_PDF = (b"%PDF-1.4\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF\n")


def aufraeumen(s):
    import tests.test_projektierung_v4 as v4
    alt = v4.TEST_EMAIL
    v4.TEST_EMAIL = TEST_EMAIL
    try:
        for k in s.query(Kunde).filter(Kunde.email == TEST_EMAIL):
            for p in s.query(Projekt).filter_by(kunde_id=k.id):
                s.query(ProjektDokument).filter_by(projekt_id=p.id).delete()
        s.commit()
        _aufraeumen_v4(s)
    finally:
        v4.TEST_EMAIL = alt


class Basis(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        init_db()
        cls.s = SessionLocal()
        aufraeumen(cls.s)
        cls.client = TestClient(app)
        cls.client.post("/login", data={"benutzer_id": "1", "pin": "1234"})

    @classmethod
    def tearDownClass(cls):
        aufraeumen(cls.s)
        cls.s.close()

    def taifun_angebot(self, sparte="WP", pdf=False):
        kunde = Kunde(anrede="Frau", vorname="Tina", nachname="Taifun-V3-Test",
                      strasse="Hafenstr. 1", plz="47119", ort="Duisburg",
                      email=TEST_EMAIL, telefon="0203 2")
        self.s.add(kunde)
        self.s.flush()
        angebot = Angebot(nummer="EXT-tmp", kunde_id=kunde.id, extern=True,
                          taifun_nummer="AN261234", extern_endbetrag_cent=2500000,
                          datum=datetime.now(), konfigurator_typ=sparte)
        self.s.add(angebot)
        self.s.flush()
        angebot.nummer = f"EXT-{angebot.id}"
        angebot_status_setzen(angebot, "Angenommen")
        self.s.commit()
        return angebot


class Phase84Auftragsdaten(Basis):
    def test_felder_aus_blatt(self):
        logik = projektierung_logik.hole_logik(self.s, erzwingen=True)
        self.assertEqual(logik.fehler, [])
        felder = {f.feld: f for f in logik.auftragsdaten_felder("WP")}
        self.assertEqual(felder["hersteller"].typ, "auswahl")
        self.assertIn("Vaillant", felder["hersteller"].optionen)
        self.assertEqual(felder["oeltank"].typ, "ja_nein")
        self.assertEqual(felder["leistungsklasse"].typ, "zahl")
        self.assertIn("wallbox", {f.feld for f in logik.auftragsdaten_felder("PV")})
        # zusätzliche Felder erscheinen im Steckbrief der Akte
        self.assertIn("serie_modell", [f for f, _ in kern.steckbrief_felder("WP")])

    def test_taifun_projekt_mit_auftragsdaten(self):
        angebot = self.taifun_angebot()
        # Anlage über den Dialog → Weiterleitung auf die Auftragsdaten
        antwort = self.client.post(f"/projektierung/angebot/{angebot.id}/projekt",
                                   data={"projekt_wahl": "neu", "sparte_WP": "on",
                                         "projektleiter_id": "1", "strasse": "Hafenstr. 1",
                                         "plz": "47119", "ort": "Duisburg"},
                                   follow_redirects=False)
        self.assertEqual(antwort.status_code, 303)
        self.assertIn("/auftragsdaten", antwort.headers["location"])
        self.s.expire_all()
        gewerk = self.s.query(Gewerk).filter(Gewerk.angebot_id == angebot.id).one()
        seite = self.client.get(f"/projektierung/gewerk/{gewerk.id}/auftragsdaten")
        self.assertEqual(seite.status_code, 200)
        self.assertIn("Angebots-PDF (Pflicht)", seite.text)
        self.assertIn("KfW-gefördert", seite.text)
        # ohne PDF → abgewiesen
        ohne = self.client.post(f"/projektierung/gewerk/{gewerk.id}/auftragsdaten",
                                data={f"g{gewerk.id}__hersteller": "Bosch"},
                                follow_redirects=False)
        self.assertIn("Pflicht", ohne.headers["location"].replace("%20", " ")
                      .replace("+", " "))
        p = f"g{gewerk.id}__"
        daten = {p + "hersteller": "Vaillant", p + "serie_modell": "aroTHERM plus",
                 p + "leistungsklasse": "10", p + "innengeraet": "uniTOWER",
                 p + "oeltank": "ja", p + "oeltank_groesse": "3.000",
                 p + "oeltank_material": "Stahl, Kellerzugang", p + "folierung": "ja",
                 p + "zaehlerschrank": "3-Feld", p + "stemmarbeiten": "ja",
                 p + "erdleitung_m": "12", "kfw_gefoerdert": "ja"}
        mit = self.client.post(f"/projektierung/gewerk/{gewerk.id}/auftragsdaten",
                               data=daten,
                               files={"pdf_datei": ("AN261234.pdf", io.BytesIO(MINI_PDF),
                                                    "application/pdf")},
                               follow_redirects=False)
        self.assertEqual(mit.status_code, 303)
        self.s.expire_all()
        gewerk = self.s.get(Gewerk, gewerk.id)
        angebot = self.s.get(Angebot, angebot.id)
        self.assertIsNotNone(gewerk.auftragsdaten_am)
        self.assertTrue(angebot.extern_pdf_pfad)
        self.assertEqual(angebot.kfw_gefoerdert, "ja")
        werte = kern.steckbrief_daten(self.s, [gewerk.id])[gewerk.id]
        self.assertEqual(werte["hersteller"].wert, "Vaillant")
        self.assertEqual(werte["hersteller"].quelle, "auftragsdaten")
        self.assertFalse(werte["hersteller"].manuell)
        self.assertEqual(werte["oeltank_groesse"].wert, "3000")
        # Sub-Mail-Platzhalter
        platz = kern.sub_mail_platzhalter(self.s, gewerk)
        self.assertEqual(platz["geraet"], "Vaillant · aroTHERM plus · 10 · uniTOWER")
        self.assertEqual(platz["oeltank"], "ja · 3000 · Stahl, Kellerzugang")
        self.assertIn("12 m", platz["erdarbeiten"])
        self.assertTrue(platz["stemmarbeiten"].startswith("Stemmarbeiten sind"))
        # BzA: gefördert = ja → Buttons (kein „unbekannt“ mehr)
        from app import bza
        self.assertIs(bza.ist_gefoerdert(self.s, gewerk), True)
        # PDF auch in der Galerie „Allgemein“
        from app.models import GalerieDatei
        projekt = self.s.get(Projekt, gewerk.projekt_id)
        self.assertTrue(self.s.query(GalerieDatei).filter_by(
            vorgang_id=projekt.vorgang_id, ordner="Allgemein").count())
        # Material-Aufgabe: Hinweis statt UGL-Button
        akte = self.client.get(f"/projektierung/projekt/{projekt.id}").text
        self.assertIn("Auftragsdaten bearbeiten", akte)
        material = (self.s.query(Aufgabe).filter(Aufgabe.gewerk_id == gewerk.id,
                                                Aufgabe.aktion_wert == "ugl_collin").first())
        if material is not None:
            self.assertIn("Bestellung aus TAIFUN-Positionen", akte)
        # „Neu ableiten“ überschreibt Auftragsdaten nicht …
        kern.steckbrief_ableiten(self.s, gewerk)
        self.s.commit()
        self.assertEqual(kern.steckbrief_daten(self.s, [gewerk.id])[gewerk.id]
                         ["hersteller"].wert, "Vaillant")
        # … die FP-Erfassung schon (fp_frage-Regel FP-O01 → oeltank)
        kern.steckbrief_ableiten(self.s, gewerk, fp_antworten={"FP-O01": "Nein"})
        self.s.commit()
        oel = kern.steckbrief_daten(self.s, [gewerk.id])[gewerk.id]["oeltank"]
        self.assertEqual((oel.wert, oel.quelle), ("nein", ""))

    def test_steckbrief_schritte_nachziehen(self):
        """Schritte mit steckbrief:-Bedingung (z. B. Folierung) entstehen nach
        dem Speichern der Auftragsdaten nachträglich."""
        logik = projektierung_logik.hole_logik(self.s)
        bedingte = [(p.key, sch.nr, sch.sichtbar_wenn) for p in logik.pakete.values()
                    for sch in p.schritte
                    if (sch.sichtbar_wenn or "").lower().startswith("steckbrief:")
                    and p.sparte in ("ALLE", "WP")]
        if not bedingte:
            self.skipTest("keine steckbrief-abhängigen Schritte im Blatt")
        angebot = self.taifun_angebot()
        projekt = kern.projekt_anlegen(self.s, angebot, projektleiter_id=1)
        gewerk = kern.gewerk_anlegen(self.s, projekt, angebot, "WP")
        self.s.commit()
        key, nr, bedingung = bedingte[0]
        feld, _, soll = bedingung[len("steckbrief:"):].partition("=")
        instanz = (self.s.query(AufgabenpaketInstanz)
                   .filter_by(gewerk_id=gewerk.id, paket_key=key).first())
        if instanz is None:
            self.skipTest("Paket mit bedingtem Schritt ist kein IMMER-Paket")
        vorher = {a.reihenfolge for a in self.s.query(Aufgabe)
                  .filter_by(paket_instanz_id=instanz.id)}
        self.assertNotIn(nr, vorher)
        feld_def = next((f for f in kern.auftragsdaten_felder(self.s, "WP")
                         if f.feld == feld.strip()), None)
        if feld_def is None:
            self.skipTest("Bedingungsfeld nicht im Auftragsdaten-Formular")
        kern.auftragsdaten_speichern(self.s, gewerk, {feld.strip(): soll.strip()})
        self.s.commit()
        nachher = {a.reihenfolge for a in self.s.query(Aufgabe)
                   .filter_by(paket_instanz_id=instanz.id)}
        self.assertIn(nr, nachher)

    def test_tool_angebot_ohne_auftragsdaten(self):
        """Tool-Angebote kennen keine Auftragsdaten-Seite."""
        from tests.test_projektierung_v4 import Basis as V4
        gewerk = V4.gewerk_neu(self)
        projekt = self.s.get(Projekt, gewerk.projekt_id)
        self.s.get(Kunde, projekt.kunde_id).email = TEST_EMAIL   # Aufräumen
        self.s.commit()
        antwort = self.client.get(f"/projektierung/gewerk/{gewerk.id}/auftragsdaten",
                                  follow_redirects=False)
        self.assertEqual(antwort.status_code, 303)
        self.assertNotIn("/auftragsdaten", antwort.headers["location"])

    def test_pdf_stub(self):
        self.assertEqual(kern.auftragsdaten_aus_pdf(b"%PDF"), {})


# --- Phase 85 ---------------------------------------------------------------

class Phase85Bestandsimport(Basis):
    def _datei(self, zeilen):
        import openpyxl
        from app import bestandsimport
        wb = openpyxl.load_workbook(io.BytesIO(bestandsimport.vorlage_bytes()))
        ws = wb["Projekte"]
        ws.delete_rows(2, 1)                       # Beispielzeile raus
        titel = [t for _, t, _, _ in bestandsimport.SPALTEN]
        for werte in zeilen:
            ws.append([werte.get(t, "") for t in titel])
        puffer = io.BytesIO()
        wb.save(puffer)
        return puffer.getvalue()

    def test_import_vorschau_zweiter_lauf_rueckgaengig(self):
        import re

        from app import bestandsimport
        from app.models import Benutzer, Bestandsimport, Team
        kern.teams_vorbelegen(self.s)
        self.s.commit()
        team = self.s.query(Team).filter(Team.typ == "montage").order_by(Team.id).first()
        pl = self.s.get(Benutzer, 1).name
        # vorhandener Kunde (Dublette über Name + PLZ)
        alt = Kunde(anrede="Herr", vorname="Otto", nachname="Bestand-V3-Alt",
                    plz="47051", ort="Duisburg", email=TEST_EMAIL)
        self.s.add(alt)
        self.s.commit()
        basis = {"Kunde Anrede": "Frau", "Vorname": "Berta", "Nachname": "Bestand-V3-Test",
                 "Straße": "Ruhrort 3", "PLZ": "47119", "Ort": "Duisburg",
                 "E-Mail": TEST_EMAIL, "Auftragswert brutto": "32.450,00",
                 "Auftragsdatum": "01.08.2026", "Projektleiter (Name)": pl}
        zeilen = [
            {**basis, "Sparte": "WP", "TAIFUN-Angebotsnummer": "AN26B001",
             "Phase": "Montagevorbereitung", "Montage von": "20.10.2026",
             "Montageteam": team.name, "Termin bestätigt (J/N)": "J",
             "Hersteller": "Bosch", "Leistungsklasse": "10", "Öltankentsorgung (J/N)": "J"},
            {**basis, "Sparte": "PV", "TAIFUN-Angebotsnummer": "AN26B002",
             "Phase": "Planung", "Auftragswert brutto": "18.000"},
            {**basis, "Nachname": "Bestand-V3-Fehler", "Sparte": "XX",
             "TAIFUN-Angebotsnummer": "AN26B003", "Phase": "Abnahme & Freigabe"},
            {**basis, "Vorname": "Otto", "Nachname": "Bestand-V3-Alt", "PLZ": "47051",
             "Sparte": "WP", "TAIFUN-Angebotsnummer": "AN26B004", "Phase": "Abnahme",
             "Montage von": "01.09.2026", "Montage bis": "05.09.2026"},
            {**basis, "Nachname": "Bestand-V3-Zwei", "PLZ": "47057", "Sparte": "KL",
             "TAIFUN-Angebotsnummer": "AN26B005", "Phase": "Auftragseingang"},
        ]
        inhalt = self._datei(zeilen)
        # Vorschau über die Route
        vorschau = self.client.post("/parametrierung/bestandsimport/vorschau",
                                    files={"datei": ("bestand.xlsx", io.BytesIO(inhalt))})
        self.assertEqual(vorschau.status_code, 200)
        self.assertIn("4</strong> von 5 Zeilen fehlerfrei", vorschau.text)
        self.assertIn("seit V4 getrennt", vorschau.text)
        self.assertIn("vorhanden #", vorschau.text)
        kennung = re.search(r'name="kennung" value="(\w+)"', vorschau.text).group(1)
        antwort = self.client.post("/parametrierung/bestandsimport/ausfuehren",
                                   data={"kennung": kennung, "dateiname": "bestand.xlsx"},
                                   follow_redirects=False)
        self.assertEqual(antwort.status_code, 303)
        self.s.expire_all()
        imp = self.s.query(Bestandsimport).order_by(Bestandsimport.id.desc()).first()
        self.assertEqual((imp.angelegt, imp.aktualisiert, imp.uebersprungen), (4, 0, 1))
        wp = self.s.query(Angebot).filter_by(taifun_nummer="AN26B001").one()
        pv = self.s.query(Angebot).filter_by(taifun_nummer="AN26B002").one()
        self.assertTrue(wp.bestand and wp.extern and wp.status == "Angenommen")
        self.assertEqual(wp.kunde_id, pv.kunde_id)          # zwei Sparten, ein Kunde
        g_wp = self.s.get(Gewerk, wp.projekt_gewerk_id)
        g_pv = self.s.get(Gewerk, pv.projekt_gewerk_id)
        self.assertEqual(g_wp.projekt_id, g_pv.projekt_id)  # ein Projekt
        self.assertEqual((g_wp.phase, g_pv.phase), ("montagevorbereitung", "planung"))
        self.assertEqual(g_wp.bestand_import_id, imp.id)
        # Pflichtaufgaben vor der Phase erledigt, Aufgaben der Phase offen
        instanzen = {i.id: i for i in self.s.query(AufgabenpaketInstanz)
                     .filter_by(gewerk_id=g_wp.id)}
        vorher_offen = []
        for a in self.s.query(Aufgabe).filter_by(gewerk_id=g_wp.id, pflicht=True):
            if a.paket_instanz_id not in instanzen or a.status == "entfaellt":
                continue
            rang = kern.paket_rang(instanzen[a.paket_instanz_id].paket_name)
            if rang < 3 and a.status != "erledigt":
                vorher_offen.append(a.titel)
        self.assertEqual(vorher_offen, [])
        self.assertTrue(any(a.status != "erledigt" for a in self.s.query(Aufgabe)
                            .filter_by(gewerk_id=g_wp.id, pflicht=True)
                            if a.paket_instanz_id in instanzen and kern.paket_rang(
                                instanzen[a.paket_instanz_id].paket_name) >= 3))
        # Termin mit Team + Bestätigung, Steckbrief
        termin = self.s.query(ProjektTermin).filter_by(gewerk_id=g_wp.id, typ="montage").one()
        self.assertTrue(termin.kunde_bestaetigt)
        self.assertEqual(termin.team_id, team.id)
        werte = kern.steckbrief_daten(self.s, [g_wp.id])[g_wp.id]
        self.assertEqual(werte["hersteller"].wert, "Bosch")
        self.assertEqual(werte["oeltank"].wert, "ja")
        # Dubletten-Kunde wiederverwendet
        alt_ang = self.s.query(Angebot).filter_by(taifun_nummer="AN26B004").one()
        self.assertEqual(alt_ang.kunde_id, alt.id)
        # Badge in Akte + Board, Statistik ausgenommen
        self.assertIn(">Bestand<", self.client.get(
            f"/projektierung/projekt/{g_wp.projekt_id}").text)
        self.assertIn(">Bestand<", self.client.get("/projektierung").text)
        from app.routers import statistik
        self.assertNotIn(wp.id, statistik._vertriebler_map(self.s))
        # zweiter Lauf: aktualisiert statt dupliziert
        anzahl_vorher = self.s.query(Gewerk).count()
        zeilen[0]["Phase"] = "Montage"
        zweit = bestandsimport.einlesen(self._datei(zeilen))[0]
        bestandsimport.pruefen(self.s, zweit)
        imp2 = bestandsimport.importieren(self.s, zweit, "bestand2.xlsx")
        self.s.commit()
        self.assertEqual((imp2.angelegt, imp2.aktualisiert), (0, 4))
        self.assertEqual(self.s.query(Gewerk).count(), anzahl_vorher)
        self.assertEqual(self.s.query(Angebot).filter_by(taifun_nummer="AN26B001").count(), 1)
        self.assertEqual(self.s.get(Gewerk, g_wp.id).phase, "montage")
        # Rückgängig Import 2 entfernt nichts (nur aktualisiert)
        entfernt, _ = bestandsimport.rueckgaengig(self.s, imp2)
        self.assertEqual(entfernt, 0)
        # Rückgängig Import 1: alle Gewerke tragen jetzt Import 2 → bleiben;
        # Import 1 wirkt nur auf das, was seither unverändert ist
        imp = self.s.get(Bestandsimport, imp.id)
        entfernt, behalten = bestandsimport.rueckgaengig(self.s, imp)
        self.s.commit()
        self.assertEqual(entfernt, 0)
        self.assertEqual(len(behalten), 4)
        self.assertIsNotNone(self.s.get(Gewerk, g_wp.id))

    def test_rueckgaengig_unveraendert(self):
        """Direkt nach dem Import: Rückgängig entfernt alles Angelegte, der
        vorher vorhandene Kunde bleibt."""
        from app import bestandsimport
        from app.models import Benutzer
        alt = Kunde(anrede="Herr", vorname="Uwe", nachname="Bestand-V3-Undo",
                    plz="47058", ort="Duisburg", email=TEST_EMAIL)
        self.s.add(alt)
        self.s.commit()
        zeilen = [{"Vorname": "Uwe", "Nachname": "Bestand-V3-Undo", "PLZ": "47058",
                   "E-Mail": TEST_EMAIL, "Sparte": "WP", "TAIFUN-Angebotsnummer": "AN26U001",
                   "Auftragswert brutto": "20000", "Auftragsdatum": "02.08.2026",
                   "Projektleiter (Name)": self.s.get(Benutzer, 1).name,
                   "Phase": "Planung"},
                  {"Vorname": "Ina", "Nachname": "Bestand-V3-Undo-Neu", "PLZ": "47059",
                   "E-Mail": TEST_EMAIL, "Sparte": "PV", "TAIFUN-Angebotsnummer": "AN26U002",
                   "Auftragswert brutto": "15000", "Auftragsdatum": "02.08.2026",
                   "Projektleiter (Name)": self.s.get(Benutzer, 1).name,
                   "Phase": "Auftragseingang"}]
        z = bestandsimport.einlesen(self._datei(zeilen))[0]
        bestandsimport.pruefen(self.s, z)
        imp = bestandsimport.importieren(self.s, z, "undo.xlsx")
        self.s.commit()
        self.assertEqual(imp.angelegt, 2)
        entfernt, behalten = bestandsimport.rueckgaengig(self.s, imp)
        self.s.commit()
        self.assertEqual((entfernt, behalten), (2, []))
        self.assertIsNone(self.s.query(Angebot).filter_by(taifun_nummer="AN26U001").first())
        self.assertIsNotNone(self.s.get(Kunde, alt.id))
        self.assertIsNone(self.s.query(Kunde).filter_by(
            nachname="Bestand-V3-Undo-Neu").first())

    def test_vorlage(self):
        import openpyxl
        from app import bestandsimport
        wb = openpyxl.load_workbook(io.BytesIO(bestandsimport.vorlage_bytes()))
        self.assertEqual(wb.sheetnames, ["Projekte", "Anleitung"])
        kopf = [c.value for c in wb["Projekte"][1]]
        self.assertIn("TAIFUN-Angebotsnummer", kopf)
        self.assertIn("Termin bestätigt (J/N)", kopf)


if __name__ == "__main__":
    unittest.main()
