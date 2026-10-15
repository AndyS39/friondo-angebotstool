# Tests PLAN_V21 Phase 178 (v34, Agent W – Stand Nachtrag 15.10.2026): Import der Wallbox-
# Positionsliste (app/import_wallbox.py) gegen die synthetische Fixture
# tests/fixtures/wallbox/WB Angebot Excel 1.xlsx (Aufbau wie der echte TAIFUN-Export: Blatt
# „Tabelle 1“, Textzeile ohne Position, Positionen 001–008, Werte als Text, `_x000D_`,
# „?“-Platzhalter; Langtexte als Platzhalter, EK frei erfunden). Stichproben gegen die echte
# Zulieferung in Artikel-Preislisten/Wallbox/ laufen nur, wenn sie vorliegt – EK-Werte werden
# dabei nie geprüft oder ausgegeben. Läuft gegen die Datenbank des Testaufrufs
# (DB_PFAD_OVERRIDE); der Import legt die WB-Artikel an (idempotent, Anker WB-Nr.).
import re
import shutil
import tempfile
import unittest
import warnings
from pathlib import Path
from unittest import mock

warnings.filterwarnings("ignore")

import openpyxl

from app import config, import_klima, import_preisliste, import_pv, import_wallbox
from app import logik as logik_modul
from app.db import SessionLocal, init_db
from app.models import Artikel, einstellung_holen, einstellung_setzen

FIXTURES = Path(__file__).resolve().parent / "fixtures" / "wallbox"
PFLICHT = "WB Angebot Excel 1.xlsx"
ECHT = import_wallbox.wb_ordner()          # Zulieferung (nur im Projektordner, Stichproben)
QUELLE_WB = import_wallbox.QUELLE_WB
WB_NUMMERN = [f"WB{n:03d}" for n in range(1, 9)]
VK = {"WB001": 93850, "WB002": 139055, "WB003": 129720, "WB004": 0, "WB005": 4934,
      "WB006": 12350, "WB007": 102000, "WB008": 16900}
EINHEIT = {"WB001": "Stück", "WB002": "Stück", "WB003": "psl.", "WB004": "psl.", "WB005": "m",
           "WB006": "m", "WB007": "psl.", "WB008": "Stück"}
POSITION = {"WB001": "001", "WB002": "002", "WB003": "003", "WB004": "005", "WB005": "006",
            "WB006": "007", "WB007": "008", "WB008": "004"}
# Fixture-EK (frei erfunden – nur zur Kontrolle, dass sie in artikel.ek_cent landen und nie im Protokoll)
EK_FIXTURE = {"WB001": 61120, "WB002": 90380, "WB003": 41200, "WB004": 0, "WB005": 2110,
              "WB006": 6000, "WB007": 70000, "WB008": 8000}
EK_FIXTURE_TEXTE = ("611,20", "903,80", "412,00", "21,10", "60,00", "700,00", "80,00")
MUSTER_SUMME = 362625      # Musterangebot: Pos. 001 + 002 + 003 + 005 = 3.626,25 €


def _zweite_liste(ordner: Path, preis_003: str = "1597.2") -> Path:
    """Kopie der Fixture als optionale zweite Liste mit anderem Preis für Pos. 003."""
    wb = openpyxl.load_workbook(FIXTURES / PFLICHT)
    ws = wb["Tabelle 1"]
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 3).value == "003":
            ws.cell(r, 7).value = preis_003
    ziel = ordner / "Zweite Liste.xlsx"
    wb.save(ziel)
    return ziel


class Bereinigung(unittest.TestCase):
    def test_x000d_fragezeichen_und_komma(self):
        roh = ("Mennekes AMTRON 4You 310_x000D_\n? Produktvorteile_x000D_\n  ?  Ladeleistung 11 kW_x000D_\n"
               "Reicht die Zuleitung aus?_x000D_\n   ")
        text = import_wallbox.text_bereinigen(roh)
        self.assertEqual(text, "Mennekes AMTRON 4You 310\nProduktvorteile\nLadeleistung 11 kW\n"
                               "Reicht die Zuleitung aus?")
        self.assertNotIn("_x000D_", text)
        self.assertNotIn("  ", text)
        self.assertEqual(import_wallbox.text_bereinigen(None), "")
        self.assertEqual(import_wallbox.text_bereinigen("?? Nur Platzhalter ?"), "Nur Platzhalter")
        # Leerzeichen vor Satzzeichen (TAIFUN „E-Auto , sondern“) und „?“-Symbole mitten im Satz
        self.assertEqual(import_wallbox.text_bereinigen("Ihr E-Auto , sondern ? auch  mehr ."),
                         "Ihr E-Auto, sondern auch mehr.")

    def test_anfang_passt(self):
        self.assertTrue(import_wallbox.anfang_passt(
            "Anschlussvorbereitung Ladestation mit Ladeanschlussleistung bis zu 11 KW",
            "anschlussvorbereitung   ladestation mit LADEANSCHLUSSLEISTUNG bis zu 11 kW\nLeistungsumfang"))
        self.assertTrue(import_wallbox.anfang_passt("Netzanmeldung beim EVU nach §14a EnWG",
                                                    "Netzanmeldung beim EVU nach §14a EnWG"))
        self.assertFalse(import_wallbox.anfang_passt("Netzanmeldung beim EVU nach §14a EnWG",
                                                     "Netzanmeldung beim Netzbetreiber"))
        self.assertFalse(import_wallbox.anfang_passt("", "irgendwas"))
        self.assertEqual(len(import_wallbox.anfang_norm("x" * 80)), 40)

    def test_einheiten_und_dateinamen(self):
        self.assertEqual(import_wallbox.einheit_norm("Stk"), "Stück")
        self.assertEqual(import_wallbox.einheit_norm("psl"), "psl.")
        self.assertEqual(import_wallbox.einheit_norm("lfm"), "m")
        self.assertEqual(import_wallbox.einheit_norm("", "m"), "m")
        self.assertEqual(import_wallbox.datei_finden(FIXTURES, "wb angebot excel 1.xlsx").name, PFLICHT)
        self.assertEqual(import_wallbox.datei_finden(FIXTURES, "WB Angebot Excel 1").name, PFLICHT)
        self.assertIsNone(import_wallbox.datei_finden(FIXTURES, "Extraleitung"))
        self.assertEqual(import_wallbox.anker_dateien(), ["WB Angebot Excel 1"])   # eine Pflichtdatei
        self.assertEqual(import_wallbox.dateien_vorhanden(FIXTURES), (True, []))
        with tempfile.TemporaryDirectory() as leer:
            self.assertEqual(import_wallbox.dateien_vorhanden(Path(leer)), (False, [PFLICHT]))
        self.assertEqual(import_wallbox.dateien_vorhanden(Path("C:/gibt/es/nicht"))[0], False)
        self.assertEqual(import_wallbox.weitere_listen(FIXTURES, ["WB Angebot Excel 1"]), [])
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            shutil.copyfile(FIXTURES / PFLICHT, tmp / PFLICHT)
            _zweite_liste(tmp)
            (tmp / "~$WB Angebot Excel 1.xlsx").write_bytes(b"")      # Excel-Sperrdatei
            self.assertEqual([p.name for p in import_wallbox.weitere_listen(tmp, ["WB Angebot Excel 1"])],
                             ["Zweite Liste.xlsx"])


class Lesen(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.erg = import_wallbox.lese_wb_dateien(FIXTURES)
        cls.nach = {a["pos_nr"]: a for a in cls.erg.artikel}

    def test_anker_aus_blatt(self):
        anker = import_wallbox.anker_lesen()
        self.assertEqual([a["nr"] for a in anker], WB_NUMMERN)
        self.assertEqual({a["nr"]: a["position"] for a in anker}, POSITION)
        self.assertTrue(all(a["datei"] == "WB Angebot Excel 1" for a in anker))
        self.assertEqual({a["nr"]: a["einheit"] for a in anker}, EINHEIT)

    def test_acht_artikel_und_summe(self):
        nach = self.nach
        self.assertEqual(sorted(nach), WB_NUMMERN)
        self.assertEqual({nr: a["e_preis_cent"] for nr, a in nach.items()}, VK)
        self.assertEqual({nr: a["einheit"] for nr, a in nach.items()}, EINHEIT)
        self.assertEqual({nr: a["_position"] for nr, a in nach.items()}, POSITION)
        self.assertEqual(sum(nach[nr]["e_preis_cent"] for nr in ("WB001", "WB002", "WB003", "WB004")),
                         MUSTER_SUMME)
        self.assertEqual(self.erg.konflikte, [])
        self.assertFalse([w for w in self.erg.warnungen if "FEHLER" in w or "fehlt" in w], self.erg.warnungen)
        for a in self.erg.artikel:
            self.assertEqual(a["quelle"], QUELLE_WB)
            self.assertEqual(a["kategorie"], "WB · Wallbox")
            self.assertFalse(a["ep_flag"])
            self.assertEqual(a["_datei"], "WB Angebot Excel 1")
            self.assertNotIn("_x000D_", a["bezeichnung"] + a["beschreibung"])
            self.assertNotIn("?", a["bezeichnung"])
            self.assertNotIn("\n", a["bezeichnung"])
            self.assertFalse(re.search(r"(^|\n) ", a["beschreibung"]))   # keine führenden Leerzeichen
            self.assertNotIn(" ,", a["bezeichnung"] + a["beschreibung"])
        self.assertEqual(nach["WB001"]["bezeichnung"], "Mennekes AMTRON 4You 310 11 kW mit 7,5 m Ladekabel")
        self.assertIn("Produktvorteile\n", nach["WB001"]["beschreibung"])
        # echtes Fragezeichen am Wortende bleibt, Platzhalter-„?“ sind weg
        self.assertIn("reicht die Zuleitung aus?", nach["WB003"]["beschreibung"])
        self.assertNotIn("? ", nach["WB003"]["beschreibung"])
        self.assertEqual(nach["WB008"]["bezeichnung"],
                         "Zusätzliche erforderte Bohrung/Durchbrüche für den Leitungsdurchgang o.ä., "
                         "bis 320mm Wandstärke. Zuschlag für Beton.")
        self.assertEqual(nach["WB006"]["bezeichnung"],
                         "Erdarbeiten zwischen Wallbox bis zur Hauseinführung Boden auf eine Tiefe "
                         "von ca. 40 cm ausheben.")
        self.assertTrue(nach["WB007"]["bezeichnung"].startswith("Fundament und Standsäule"))
        # Angebots-GUID des Exports mitgenommen (Fixture-Muster)
        self.assertTrue(nach["WB001"]["guid"].startswith("{F"))
        self.assertIn("-A4B1-", nach["WB005"]["guid"])

    def test_einheit_blatt_ueberstimmt_export(self):
        """Pos. 008 trägt im Export „m“, das Blatt „WB-Artikel“ sagt psl. – das Blatt gilt."""
        self.assertEqual(self.nach["WB007"]["einheit"], "psl.")
        zeilen = [p for p in self.erg.protokoll if "Einheit laut Export" in p]
        self.assertEqual(zeilen, ["WB007: Einheit laut Export „m“, laut Blatt „WB-Artikel“ „psl.“ – das Blatt gilt"])

    def test_protokoll_ohne_ek(self):
        prot = "\n".join(self.erg.protokoll + self.erg.warnungen + self.erg.konflikte)
        self.assertIn("WB Angebot Excel 1 Pos. 001 → WB001", prot)
        self.assertIn("VK 938,50 €", prot)
        self.assertIn("WB Angebot Excel 1 Pos. 004 → WB008", prot)
        self.assertIn("VK 169,00 €", prot)
        self.assertIn("Textzeile ohne Position", prot)
        for ek in EK_FIXTURE_TEXTE + ("Einkaufspreis", "EK "):
            self.assertNotIn(ek, prot, ek)

    def test_komplettpaket_text(self):
        text = import_wallbox.komplettpaket_text(FIXTURES)
        self.assertTrue(text.startswith("Komplettpaket Wallbox inkl. Montage\n"), text[:60])
        self.assertIn("Die Zukunft der Energie beginnt bei Ihnen zuhause und wir begleiten Sie dabei.", text)
        self.assertIn("E-Auto, ", text)              # „E-Auto ,“ bereinigt
        for verboten in ("?", "  ", "_x000D_", " ,"):
            self.assertNotIn(verboten, text, verboten)
        self.assertEqual(self.erg.textzeilen, [text])
        with tempfile.TemporaryDirectory() as leer:
            self.assertEqual(import_wallbox.komplettpaket_text(Path(leer)), "")

    def test_zweite_liste_preiskonflikt(self):
        """Optionale weitere Liste im Ordner: nur Preisvergleich – anderer Preis = Konflikt
        (Ankerdatei gewinnt), gleicher Preis = „identisch“."""
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            shutil.copyfile(FIXTURES / PFLICHT, tmp / PFLICHT)
            _zweite_liste(tmp)
            erg = import_wallbox.lese_wb_dateien(tmp)
            nach = {a["pos_nr"]: a for a in erg.artikel}
            self.assertEqual(sorted(nach), WB_NUMMERN)
            self.assertEqual(len(erg.konflikte), 1)
            k = erg.konflikte[0]
            for teil in ("WB003", "1.597,20 €", "1.297,20 €", "WB Angebot Excel 1", "Zweite Liste"):
                self.assertIn(teil, k)
            self.assertIn(k, erg.warnungen)
            self.assertEqual(nach["WB003"]["e_preis_cent"], 129720)         # Ankerdatei gewinnt
            prot = "\n".join(erg.protokoll)
            self.assertIn("Zweite Liste: weitere Liste (optional) – nur Preisvergleich", prot)
            self.assertIn("identisch mit WB001", prot)
            self.assertEqual(sum("identisch mit" in p for p in erg.protokoll), 7)
            # identische zweite Liste: kein Konflikt
            _zweite_liste(tmp, "1297.2")
            erg = import_wallbox.lese_wb_dateien(tmp)
            self.assertEqual(erg.konflikte, [])
            self.assertEqual(sum("identisch mit" in p for p in erg.protokoll), 8)

    def test_beschreibungsanfang_abweichung(self):
        anker = import_wallbox.anker_lesen()
        anker[0]["anfang"] = "Tesla Wall Connector Gen 3"
        erg = import_wallbox.lese_wb_dateien(FIXTURES, anker)
        self.assertNotIn("WB001", {a["pos_nr"] for a in erg.artikel})
        self.assertTrue(any("FEHLER" in w and "Beschreibungsanfang" in w and "WB001" in w
                            for w in erg.warnungen), erg.warnungen)

    def test_ohne_dateien_und_ohne_anker(self):
        with tempfile.TemporaryDirectory() as leer:
            erg = import_wallbox.lese_wb_dateien(Path(leer))
            self.assertEqual(erg.artikel, [])
            self.assertEqual([w for w in erg.warnungen if "fehlt" in w],
                             [f"Wallbox-Positionsliste fehlt: {Path(leer) / PFLICHT}"])
        erg = import_wallbox.lese_wb_dateien(FIXTURES, [])
        self.assertEqual(erg.artikel, [])
        self.assertTrue(any("WB-Artikel" in w for w in erg.warnungen))


@unittest.skipUnless((ECHT / PFLICHT).exists(),
                     "Zulieferung Artikel-Preislisten/Wallbox/WB Angebot Excel 1.xlsx liegt nicht vor")
class EchteDatei(unittest.TestCase):
    """Stichproben gegen die echte Zulieferung (nur lesen, kein Import; EK nie geprüft/ausgegeben)."""

    def test_stichproben_ohne_ek(self):
        erg = import_wallbox.lese_wb_dateien(ECHT)
        nach = {a["pos_nr"]: a for a in erg.artikel}
        self.assertEqual(sorted(nach), WB_NUMMERN)
        self.assertEqual({nr: a["e_preis_cent"] for nr, a in nach.items()}, VK)
        self.assertEqual({nr: a["einheit"] for nr, a in nach.items()}, EINHEIT)
        self.assertEqual(sum(nach[nr]["e_preis_cent"] for nr in ("WB001", "WB002", "WB003", "WB004")),
                         MUSTER_SUMME)
        self.assertFalse([w for w in erg.warnungen if "FEHLER" in w], "Beschreibungsanfang passt nicht")
        self.assertEqual(erg.konflikte, [])
        for a in erg.artikel:
            self.assertNotIn("_x000D_", a["bezeichnung"] + a["beschreibung"])
            self.assertNotIn("?", a["bezeichnung"])
            self.assertIsNotNone(a["ek_cent"])        # EK vorhanden – der Wert bleibt ungeprüft
        text = import_wallbox.komplettpaket_text(ECHT)
        self.assertTrue(text.startswith("Komplettpaket Wallbox inkl. Montage\n"))
        self.assertIn("Die Zukunft der Energie beginnt bei Ihnen zuhause", text)
        for verboten in ("?", "  ", "_x000D_"):
            self.assertNotIn(verboten, text, verboten)
        prot = "\n".join(erg.protokoll + erg.warnungen)
        self.assertNotIn("Einkaufspreis", prot)
        for a in erg.artikel:   # kein EK-Betrag im Protokoll (Meldung ohne den Wert)
            if a["ek_cent"]:
                self.assertFalse(import_wallbox._euro(a["ek_cent"]) in prot, "EK-Betrag im Protokoll")


class Import(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        init_db()
        cls.s = SessionLocal()
        cls.diff, cls.meldung = import_wallbox.import_ausfuehren(cls.s, FIXTURES)

    @classmethod
    def tearDownClass(cls):
        cls.s.close()

    def aktive(self):
        return {a.pos_nr: a for a in self.s.query(Artikel)
                .filter(Artikel.quelle == QUELLE_WB, Artikel.aktiv.is_(True))}

    def test_acht_artikel_aktiv(self):
        wb = self.aktive()
        self.assertEqual(sorted(wb), WB_NUMMERN)
        self.assertTrue(all(a.quelle == "wb" for a in wb.values()))
        self.assertTrue(all(not a.ep_flag for a in wb.values()))
        self.assertTrue(self.meldung.startswith("WB: "))
        self.assertNotIn("Preiskonflikt", self.meldung)
        self.assertEqual(import_wallbox.fehlende_wb_artikel(self.s), [])

    def test_stichproben(self):
        wb = self.aktive()
        self.assertEqual({nr: a.e_preis_cent for nr, a in wb.items()}, VK)
        self.assertEqual({nr: a.einheit for nr, a in wb.items()}, EINHEIT)
        self.assertEqual(sum(wb[nr].e_preis_cent for nr in ("WB001", "WB002", "WB003", "WB004")), MUSTER_SUMME)
        # EK nur in artikel.ek_cent (Fixture-Werte frei erfunden)
        for nr, ek in EK_FIXTURE.items():
            self.assertEqual(wb[nr].ek_cent, ek, nr)
        for a in wb.values():
            self.assertNotIn("_x000D_", a.bezeichnung + a.beschreibung)
            self.assertNotIn("\n", a.bezeichnung)
        self.assertTrue(wb["WB001"].guid.startswith("{F"))

    def test_reimport_idempotent(self):
        vorher = {nr: a.id for nr, a in self.aktive().items()}
        diff, meldung = import_wallbox.import_ausfuehren(self.s, FIXTURES)
        self.assertEqual((len(diff.neu), len(diff.geaendert), diff.unveraendert, len(diff.entfallen)),
                         (0, 0, 8, 0))
        nachher = self.aktive()
        self.assertEqual({nr: a.id for nr, a in nachher.items()}, vorher)
        self.assertEqual(self.s.query(Artikel).filter(Artikel.pos_nr.like("WB%")).count(), 8)
        diff2 = import_wallbox.berechne_diff(self.s, ordner=FIXTURES)
        self.assertEqual((len(diff2.neu), len(diff2.geaendert), len(diff2.entfallen)), (0, 0, 0))
        self.assertEqual(diff2.konflikte, [])
        self.assertTrue(diff2.protokoll)
        self.assertEqual(len(diff2.textzeilen), 1)

    def test_guid_wechsel_bei_neuem_export(self):
        """Ein neuer Export trägt andere Angebots-GUIDs – Anker bleibt die WB-Nr.,
        die GUID wird nachgeführt, kein Doppel."""
        try:
            with tempfile.TemporaryDirectory() as ordner:
                ordner = Path(ordner)
                wb = openpyxl.load_workbook(FIXTURES / PFLICHT)
                ws = wb["Tabelle 1"]
                for r in range(2, ws.max_row + 1):
                    wert = ws.cell(r, 1).value
                    if wert:
                        ws.cell(r, 1).value = str(wert).replace("-A4B1-", "-FFFF-")
                wb.save(ordner / PFLICHT)
                diff, _m = import_wallbox.import_ausfuehren(self.s, ordner)
                self.assertEqual((len(diff.neu), len(diff.entfallen)), (0, 0))
                self.assertEqual(len(diff.geaendert), 8)
                self.assertTrue(all("GUID" in felder for _v, _n, felder in diff.geaendert))
                self.assertTrue(any("GUID-Wechsel" in p and "-FFFF-" in p for p in diff.protokoll),
                                diff.protokoll)
                self.assertTrue(all("-FFFF-" in a.guid for a in self.aktive().values()))
                self.assertEqual(self.s.query(Artikel).filter(Artikel.pos_nr.like("WB%")).count(), 8)
        finally:
            # zurück auf die Fixture-GUIDs (auch wenn eine Prüfung oben scheitert)
            import_wallbox.import_ausfuehren(self.s, FIXTURES)
        self.assertTrue(all("-A4B1-" in a.guid for a in self.aktive().values()))

    def test_guid_nachfuehren_in_kopie_der_excel(self):
        """„GUID (zuletzt gesehen)“ wird in das Blatt geschrieben – geprüft an einer
        Kopie der Logik-Excel (die Live-Excel bleibt unangetastet)."""
        with tempfile.TemporaryDirectory() as ordner:
            kopie = Path(ordner) / "logik.xlsx"
            kopie.write_bytes(config.LOGIK_EXCEL_PFAD.read_bytes())
            text = import_wallbox.guid_nachfuehren_blatt({"WB001": "{TEST-GUID-1}", "WB009": "{X}"}, kopie)
            self.assertIn("1 Zeilen", text)
            wb = openpyxl.load_workbook(kopie, read_only=True)
            ws = wb["WB-Artikel"]
            zeilen = list(ws.iter_rows(values_only=True))
            kopf = [str(v or "") for v in zeilen[0]]
            i_guid = kopf.index("GUID (zuletzt gesehen)")
            self.assertEqual(zeilen[1][i_guid], "{TEST-GUID-1}")
            wb.close()
            # zweiter Lauf: nichts mehr zu ändern
            self.assertIn("0 Zeilen", import_wallbox.guid_nachfuehren_blatt({"WB001": "{TEST-GUID-1}"}, kopie))
        # Live-Excel: Spalte weiterhin ohne Test-GUID
        anker = import_wallbox.anker_lesen()
        self.assertNotEqual(anker[0]["guid"], "{TEST-GUID-1}")

    def test_wp_pv_kl_importe_fassen_wb_nicht_an(self):
        ergebnis = import_preisliste.lese_dateien()
        diff = import_preisliste.berechne_diff(self.s, ergebnis)
        self.assertFalse([a for a in diff.entfallen if a.quelle == QUELLE_WB or a.pos_nr.startswith("WB")])
        pv_diff = import_pv.berechne_diff(self.s)
        self.assertFalse([a for a in pv_diff.entfallen if a.quelle == QUELLE_WB])
        kl_diff = import_klima.berechne_diff(self.s)
        self.assertFalse([a for a in kl_diff.entfallen if a.quelle == QUELLE_WB])
        # und umgekehrt: der WB-Import sieht nur seine Quelle
        wb_diff = import_wallbox.berechne_diff(self.s, ordner=FIXTURES)
        self.assertTrue(all(a.quelle == QUELLE_WB for a in wb_diff.entfallen))

    def test_logik_pruefen_mit_artikeln_gruen(self):
        logik, bericht = logik_modul.neu_einlesen(self.s)
        self.assertTrue(bericht.ok, bericht.fehler)
        self.assertFalse(logik.wb_artikel_fehlen)
        self.assertFalse([w for w in bericht.warnungen if "Wallbox-Artikel" in w])
        sicht = logik_modul.logik_fuer_sparte(logik, "WB")
        self.assertEqual(len(sicht.aktionen), 13)
        self.assertEqual(import_wallbox.wb_referenzen(), WB_NUMMERN)


class Migration(unittest.TestCase):
    """migration_v34_wallbox: ohne Dateien einmal melden, zweiter Lauf still; mit
    Dateien automatischer Import, wenn WB-Artikel fehlen."""

    def setUp(self):
        init_db()
        self.s = SessionLocal()
        self.marker_alt = einstellung_holen(self.s, import_wallbox.MARKER_HINWEIS, "")

    def tearDown(self):
        einstellung_setzen(self.s, import_wallbox.MARKER_HINWEIS, self.marker_alt)
        self.s.commit()
        self.s.close()

    def test_ohne_dateien_meldung_nur_beim_ersten_lauf(self):
        einstellung_setzen(self.s, import_wallbox.MARKER_HINWEIS, "")
        self.s.commit()
        with tempfile.TemporaryDirectory() as leer:
            with mock.patch.object(import_wallbox, "wb_ordner", return_value=Path(leer)), \
                    mock.patch.object(import_wallbox, "fehlende_wb_artikel",
                                      return_value=["WB001", "WB002"]):
                m1 = import_wallbox.migration_v34_wallbox(self.s)
                self.s.commit()
                m2 = import_wallbox.migration_v34_wallbox(self.s)
        self.assertIn(import_wallbox.HINWEIS_DATEIEN_FEHLEN, m1)
        self.assertNotIn(import_wallbox.HINWEIS_DATEIEN_FEHLEN, m2)
        self.assertFalse([m for m in m2 if "Wallbox" in m and "fehlen" in m])
        self.assertTrue(einstellung_holen(self.s, import_wallbox.MARKER_HINWEIS, ""))

    def test_mit_dateien_import_wenn_artikel_fehlen(self):
        einstellung_setzen(self.s, import_wallbox.MARKER_HINWEIS, "alt")
        self.s.commit()
        aufrufe = []

        def fake_import(session, ordner=None, guid_nachfuehren=False):
            aufrufe.append(guid_nachfuehren)
            return import_wallbox.WbDiff(), "WB: 8 neu, 0 geändert, 0 unverändert, 0 deaktiviert"

        with mock.patch.object(import_wallbox, "wb_ordner", return_value=FIXTURES), \
                mock.patch.object(import_wallbox, "fehlende_wb_artikel", return_value=["WB005"]), \
                mock.patch.object(import_wallbox, "import_ausfuehren", side_effect=fake_import):
            meldungen = import_wallbox.migration_v34_wallbox(self.s)
        self.assertTrue(any("WB-Artikel fehlten (1, z. B. WB005)" in m for m in meldungen), meldungen)
        self.assertEqual(aufrufe, [False])                     # Live-Excel wird nie beschrieben
        self.assertEqual(einstellung_holen(self.s, import_wallbox.MARKER_HINWEIS, ""), "")

    def test_nichts_zu_tun_wenn_artikel_da(self):
        with mock.patch.object(import_wallbox, "fehlende_wb_artikel", return_value=[]):
            meldungen = import_wallbox.migration_v34_wallbox(self.s)
        self.assertFalse([m for m in meldungen if "fehlen" in m or "Import" in m], meldungen)


if __name__ == "__main__":
    unittest.main()
