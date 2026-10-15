# Tests PLAN_V21 Phase 178/179 (v34, Agent W – Stand Nachtrag + Ergänzung 15.10.2026):
# Wallbox-Blätter der Live-Excel, Lader (app/logik.py: wb_*-Felder, Mengenwörter „× Mehrlänge“,
# „× WB04“, „× (WB07 − 2)“, logik_fuer_sparte(…, "WB"), „Logik prüfen“) und der reine
# Rechenkern app/wb_auslegung.py (ohne Datenbank).
import copy
import re
import unittest
import warnings
from datetime import datetime

warnings.filterwarnings("ignore")

import openpyxl

from app import config, konfigurator as engine, sparten, wb_auslegung
from app import logik as logik_modul
from app.logik import Logik, Pruefbericht

WB_BLAETTER = ("Fragen WB", "WB-Artikel", "Aktionen WB", "Angebotsaufbau WB", "WB-Parameter")
WB310 = "Mennekes AMTRON 4You 310 – 11 kW, 7,5 m Ladekabel"
WB560 = "Mennekes AMTRON 4You 560 – 11 kW, 7,5 m Ladekabel, MID-Zähler, App, OCPP"
FRAGEN_IDS = ["WB01", "WB02", "WB03", "WB04", "WB05", "WB07", "WB08", "S01", "S02"]   # WB06 unbelegt
WB_NUMMERN = [f"WB{n:03d}" for n in range(1, 9)]
PARAMETER = ["leitung_inklusive_m", "meterposition_leitung", "erdarbeiten_position",
             "fundament_position", "durchbrueche_inklusive", "durchbruch_position",
             "wallbox_standard"]
AKTIONEN_ANZAHL = 13
BLOCK2 = ["WB003", "WB008", "WB005", "WB006", "WB007"]   # Reihenfolge Ergänzung 15.10.2026


def _zeilen(wb, blatt):
    rows = [tuple(r) for r in wb[blatt].iter_rows(values_only=True)
            if any(v is not None and str(v).strip() for v in r)]
    return [str(v or "").strip() for v in rows[0]], rows[1:]


def _text(v) -> str:
    return str(v or "").strip()


def _wb():
    return openpyxl.load_workbook(config.LOGIK_EXCEL_PFAD, data_only=True)


def _einlesen(wb) -> tuple[Logik, Pruefbericht]:
    """Nur die WB-Blätter eines (ggf. veränderten) Workbooks einlesen."""
    bericht = Pruefbericht()
    logik = Logik({}, [], [], [], {}, datetime.now())
    logik.sparten_fragen["WB"] = logik_modul._fragen_einlesen(wb, bericht, "Fragen WB")
    logik_modul._wb_einlesen(wb, logik, bericht)
    return logik, bericht


def _zeile_finden(ws, spalte: int, wert: str) -> int:
    for r in range(2, ws.max_row + 1):
        if _text(ws.cell(r, spalte).value) == wert:
            return r
    raise AssertionError(f"{wert} nicht in Spalte {spalte}")


def wb_sicht():
    logik, bericht = logik_modul.logik_einlesen()
    return logik_modul.logik_fuer_sparte(logik, "WB"), bericht


# --- Blätter der Live-Excel ----------------------------------------------------------

class Blaetter(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.wb = openpyxl.load_workbook(config.LOGIK_EXCEL_PFAD, read_only=True, data_only=True)

    @classmethod
    def tearDownClass(cls):
        cls.wb.close()

    def test_blaetter_vorhanden_und_alte_unveraendert(self):
        for name in WB_BLAETTER + ("Fragen KL", "KL-Artikel", "Aktionen KL", "Anhänge", "Lesehilfe"):
            self.assertIn(name, self.wb.sheetnames)
        # bestehende Blätter nicht umbenannt (Reihenfolge der v24-Blätter bleibt)
        self.assertLess(self.wb.sheetnames.index("KL-Parameter"), self.wb.sheetnames.index("Fragen WB"))

    def test_fragen_wb(self):
        kopf, zeilen = _zeilen(self.wb, "Fragen WB")
        self.assertEqual(kopf[:7], ["Seite", "ID", "Fragetext", "Typ", "Antwortmöglichkeiten",
                                    "Anzeigen wenn", "Hinweis"])
        self.assertEqual([_text(z[1]) for z in zeilen], FRAGEN_IDS)
        nach = {_text(z[1]): z for z in zeilen}
        self.assertEqual(_text(nach["WB01"][3]), "Auswahl")
        self.assertEqual(_text(nach["WB01"][4]), f"{WB310} | {WB560}")
        self.assertEqual(_text(nach["WB02"][2]), "Wie viel Meter Leitungsverlegung erforderlich?")
        self.assertEqual(_text(nach["WB02"][3]), "Zahleneingabe")
        self.assertIn("Vorbelegung 10 m", _text(nach["WB02"][6]))
        self.assertEqual((_text(nach["WB03"][2]), _text(nach["WB03"][3]), _text(nach["WB03"][4])),
                         ("Erdarbeiten erforderlich?", "Auswahl", "Ja | Nein"))
        self.assertIn("Vorbelegung Nein", _text(nach["WB03"][6]))
        self.assertEqual((_text(nach["WB04"][2]), _text(nach["WB04"][3]), _text(nach["WB04"][5])),
                         ("Erdarbeiten: wie viele Meter?", "Zahleneingabe", "nur wenn WB03 = Ja"))
        self.assertEqual((_text(nach["WB05"][2]), _text(nach["WB05"][3]), _text(nach["WB05"][4])),
                         ("Wandmontage oder Säule?", "Auswahl", "Wandmontage | Säule"))
        self.assertIn("Vorbelegung Wandmontage", _text(nach["WB05"][6]))
        self.assertEqual((_text(nach["WB07"][2]), _text(nach["WB07"][3])),
                         ("Wie viele Durchbrüche erforderlich?", "Zahleneingabe"))
        self.assertIn("2 Durchbrüche sind in der Anschlussvorbereitung enthalten", _text(nach["WB07"][6]))
        self.assertEqual(_text(nach["WB08"][3]), "Freitext groß")
        self.assertIn("optional", _text(nach["WB08"][2]))
        # Einschätzung (S01/S02) liegt im Sparten-Blatt – der Bogen-Rahmen liefert sie nicht
        # spartenübergreifend (Fragen KL/PV führen sie ebenfalls selbst)
        self.assertEqual((_text(nach["S01"][0]), _text(nach["S02"][3])), ("Einschätzung", "Datum"))
        for z in zeilen:
            if _text(z[1]) != "WB04":
                self.assertEqual(_text(z[5]), "immer", z)
        # Pflichtfragen: kein „optional“/„leer lassen“ in Text + Hinweis (naechste_pflichtfrage)
        for fid in ("WB01", "WB02", "WB03", "WB04", "WB05", "WB07"):
            text = (_text(nach[fid][2]) + " " + _text(nach[fid][6])).lower()
            self.assertNotIn("optional", text, fid)
            self.assertNotIn("leer lassen", text, fid)

    def test_wb_artikel(self):
        kopf, zeilen = _zeilen(self.wb, "WB-Artikel")
        self.assertEqual(kopf[:7], ["WB-Nr.", "Datei", "Position", "Beschreibungsanfang",
                                    "GUID (zuletzt gesehen)", "Einheit", "Bemerkung"])
        self.assertEqual([_text(z[0]) for z in zeilen], WB_NUMMERN)
        nach = {_text(z[0]): z for z in zeilen}
        for nr, pos, einheit in (("WB001", "001", "Stück"), ("WB002", "002", "Stück"),
                                 ("WB003", "003", "psl."), ("WB004", "005", "psl."),
                                 ("WB005", "006", "m"), ("WB006", "007", "m"),
                                 ("WB007", "008", "psl."), ("WB008", "004", "Stück")):
            self.assertEqual((_text(nach[nr][1]), _text(nach[nr][2]), _text(nach[nr][5])),
                             ("WB Angebot Excel 1", pos, einheit), nr)
        self.assertTrue(_text(nach["WB001"][3]).startswith("Mennekes AMTRON 4You 310"))
        self.assertTrue(_text(nach["WB003"][3]).startswith("Anschlussvorbereitung Ladestation"))
        self.assertTrue(_text(nach["WB004"][3]).startswith("Netzanmeldung beim EVU"))
        self.assertTrue(_text(nach["WB005"][3]).lower().startswith("lfd. meter nyy-j 5x10mm"))
        self.assertTrue(_text(nach["WB006"][3]).startswith("Erdarbeiten zwischen Wallbox"))
        self.assertTrue(_text(nach["WB007"][3]).startswith("Fundament und Standsäule"))
        self.assertTrue(_text(nach["WB008"][3]).startswith("Zusätzliche erforderte Bohrung/Durchbrüche"))
        self.assertIn("1.297,20", _text(nach["WB003"][6]))     # Preis bestätigt (R1)
        self.assertIn("psl.", _text(nach["WB007"][6]))         # Blatt überstimmt Export-Einheit „m“
        alle = " ".join(_text(v) for z in zeilen for v in z)
        self.assertNotIn("Extraleitung", alle)                 # entfällt seit dem Nachtrag
        # GUIDs sind nur „zuletzt gesehen“ (leer, bis der Import lief)
        for z in zeilen:
            self.assertFalse(_text(z[4]).startswith("{") and len(_text(z[4])) != 38, z)

    def test_aktionen_wb(self):
        kopf, zeilen = _zeilen(self.wb, "Aktionen WB")
        self.assertEqual(kopf[:5], ["Frage", "Antwort / Bedingung", "Aktion", "Bemerkung",
                                    "Zusatzbedingung"])
        aktionen = [(_text(z[0]), _text(z[1]), _text(z[2])) for z in zeilen]
        self.assertEqual(len(aktionen), AKTIONEN_ANZAHL)
        for erwartet in (("WB01", WB310, "Artikel: WB001 ×1"), ("WB01", WB560, "Artikel: WB002 ×1"),
                         ("Grundpaket", "immer", "Artikel: WB003 ×1"),
                         ("Grundpaket", "immer", "Artikel: WB004 ×1"),
                         ("WB02", "bis 10", "–"), ("WB02", "über 10", "Artikel: WB005 × Mehrlänge"),
                         ("WB03", "Ja", "Artikel: WB006 × WB04"), ("WB03", "Nein", "–"),
                         ("WB05", "Säule", "Artikel: WB007 ×1"), ("WB05", "Wandmontage", "–"),
                         ("WB07", "bis 2", "–"), ("WB07", "über 2", "Artikel: WB008 × (WB07 − 2)"),
                         ("WB08", "(Freitext)", "–")):
            self.assertIn(erwartet, aktionen)
        self.assertFalse([a for a in aktionen if a[2].startswith("AMPEL")])   # keine AMPEL-Gründe
        self.assertFalse([a for a in aktionen if "WB006 × Mehrlänge" in a[2]])  # kein Montagekanal mehr

    def test_angebotsaufbau_und_parameter(self):
        kopf, zeilen = _zeilen(self.wb, "Angebotsaufbau WB")
        self.assertEqual(kopf[:4], ["Block", "Überschrift im Angebot", "Inhalt", "Wann"])
        bloecke = {int(float(_text(z[0]))): z for z in zeilen
                   if _text(z[0]).replace(".0", "").isdigit()}
        self.assertEqual(sorted(bloecke), [1, 2, 3])
        self.assertEqual([_text(bloecke[n][1]) for n in (1, 2, 3)],
                         ["Wallbox", "Installation und Anschluss", "Netzanmeldung"])
        self.assertEqual(re.findall(r"WB\d{3}", _text(bloecke[1][2])), ["WB001", "WB002"])
        self.assertEqual(re.findall(r"WB\d{3}", _text(bloecke[2][2])), BLOCK2)
        self.assertEqual(re.findall(r"WB\d{3}", _text(bloecke[3][2])), ["WB004"])
        sonstige = [_text(z[0]) for z in zeilen if not _text(z[0]).replace(".0", "").isdigit()]
        for name in ("Summenblock", "Nachtexte", "Anhänge", "Protokoll"):
            self.assertIn(name, sonstige)
        self.assertNotIn("Auslegungszeile", sonstige)   # v34 Phase 180: keine Auslegungs-Position
        kopf, zeilen = _zeilen(self.wb, "WB-Parameter")
        self.assertEqual(kopf[:4], ["Parameter", "Wert", "Einheit", "Bemerkung"])
        werte = {_text(z[0]): _text(z[1]) for z in zeilen}
        self.assertEqual(list(werte), PARAMETER)
        self.assertEqual((float(werte["leitung_inklusive_m"]), werte["meterposition_leitung"],
                          werte["erdarbeiten_position"], werte["fundament_position"],
                          float(werte["durchbrueche_inklusive"]), werte["durchbruch_position"],
                          werte["wallbox_standard"]),
                         (10.0, "WB005", "WB006", "WB007", 2.0, "WB008", ""))
        einheiten = {_text(z[0]): _text(z[2]) for z in zeilen}
        self.assertEqual((einheiten["leitung_inklusive_m"], einheiten["durchbrueche_inklusive"]),
                         ("m", "Stück"))

    def test_anhaenge_und_lesehilfe(self):
        kopf, zeilen = _zeilen(self.wb, "Anhänge")
        zeile = next((z for z in zeilen if _text(z[0]) == "Mennekes AMTRON 4You Datenblatt.pdf"), None)
        self.assertIsNotNone(zeile)
        self.assertEqual(_text(zeile[1]), "wenn Sparte = WB")
        self.assertIn("Zulieferung", _text(zeile[2]))
        # KL-Zeile unverändert, Schlusszeile bleibt die letzte
        self.assertTrue(any(_text(z[0]) == "Bosch Climate 3200i.pdf" for z in zeilen))
        self.assertTrue(_text(zeilen[-1][0]).startswith("(weitere Zeilen"))
        texte = " ".join(_text(v) for row in self.wb["Lesehilfe"].iter_rows(values_only=True)
                         for v in row)
        for stichwort in ("Wallbox (v34)", "Mehrlänge", "Erdarbeiten", "WB07", "Klima (v24)"):
            self.assertIn(stichwort, texte)


# --- Lader ---------------------------------------------------------------------------

class Lader(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.logik, cls.bericht = logik_modul.logik_einlesen()
        cls.wb_sicht = logik_modul.logik_fuer_sparte(cls.logik, "WB")

    def test_logik_pruefen_gruen(self):
        self.assertTrue(self.bericht.ok, self.bericht.fehler)
        self.assertFalse([w for w in self.bericht.warnungen if "WB" in w or "Wallbox" in w],
                         self.bericht.warnungen)

    def test_felder(self):
        l = self.logik
        self.assertEqual([(a.frage, a.antwort[:20]) for a in l.wb_aktionen][:2],
                         [("WB01", WB310[:20]), ("WB01", WB560[:20])])
        self.assertEqual(len(l.wb_aktionen), AKTIONEN_ANZAHL)
        mengen = {(a.frage, a.antwort): [(r.ref, r.menge) for r in a.artikel] for a in l.wb_aktionen}
        self.assertEqual(mengen[("WB02", "über 10")], [("WB005", "Mehrlänge")])
        self.assertEqual(mengen[("WB03", "Ja")], [("WB006", "WB04")])
        self.assertEqual(mengen[("WB05", "Säule")], [("WB007", "1")])
        self.assertEqual(mengen[("WB07", "über 2")], [("WB008", "(WB07 − 2)")])
        self.assertEqual(mengen[("WB02", "bis 10")], [])
        self.assertEqual([(b.nr, b.ueberschrift) for b in l.wb_bloecke],
                         [(1, "Wallbox"), (2, "Installation und Anschluss"), (3, "Netzanmeldung")])
        self.assertEqual(l.wb_parameter["leitung_inklusive_m"][0], "10")
        self.assertEqual(l.wb_parameter_einheit["leitung_inklusive_m"], "m")
        self.assertEqual(l.wb_parameter["durchbrueche_inklusive"][0], "2")
        self.assertEqual(l.wb_parameter_einheit["durchbrueche_inklusive"], "Stück")
        self.assertEqual(l.wb_artikel_nummern, WB_NUMMERN)
        self.assertEqual((l.wb_artikel["WB005"]["datei"], l.wb_artikel["WB005"]["position"]),
                         ("WB Angebot Excel 1", "006"))
        self.assertEqual(l.wb_artikel["WB008"]["position"], "004")
        self.assertEqual(l.wb_artikel["WB007"]["einheit"], "psl.")
        self.assertEqual(len(l.wb_artikel_zeilen), 8)
        self.assertEqual(logik_modul.wb_parameter_fehlende(l), [])
        self.assertEqual(logik_modul.wb_standardwert_text("wallbox_standard"), "(leer)")
        self.assertEqual(logik_modul.wb_standardwert_text("leitung_inklusive_m"), "10")
        self.assertEqual(logik_modul.wb_standardwert_text("durchbruch_position"), "WB008")

    def test_refs_extrahieren_mengenwoerter(self):
        refs = logik_modul.refs_extrahieren("Artikel: WB005 × Mehrlänge")
        self.assertEqual([(r.ref, r.menge, r.ep) for r in refs], [("WB005", "Mehrlänge", False)])
        refs = logik_modul.refs_extrahieren("Artikel: WB006 × WB04")
        self.assertEqual([(r.ref, r.menge) for r in refs], [("WB006", "WB04")])
        refs = logik_modul.refs_extrahieren("Artikel: WB008 × (WB07 − 2)")
        self.assertEqual([(r.ref, r.menge) for r in refs], [("WB008", "(WB07 − 2)")])
        refs = logik_modul.refs_extrahieren("Artikel: WB001 ×1")
        self.assertEqual([(r.ref, r.menge) for r in refs], [("WB001", "1")])
        refs = logik_modul.refs_extrahieren("WB003 (EP) · WB004 ×2 als EP")
        self.assertEqual([(r.ref, r.menge, r.ep) for r in refs],
                         [("WB003", "1", True), ("WB004", "2", True)])
        # Frage-IDs (WB01/WB02) sind keine Artikelreferenzen
        self.assertEqual(logik_modul.refs_extrahieren("WB001 · WB002 Wallbox lt. WB01"),
                         logik_modul.refs_extrahieren("WB001 · WB002"))

    def test_wb_referenzen(self):
        refs = logik_modul.wb_referenzen(self.logik)
        self.assertEqual(sorted(refs), WB_NUMMERN)
        for nr in ("WB005", "WB006", "WB007", "WB008"):
            self.assertTrue(any("WB-Parameter" in q for q in refs[nr]), nr)
        alle = logik_modul.artikel_referenzen(self.logik)
        self.assertIn("WB008", alle)

    def test_sicht_wb(self):
        s = self.wb_sicht
        self.assertEqual(s.sparte, "WB")
        self.assertEqual(list(s.fragen), FRAGEN_IDS)
        self.assertEqual(len(s.aktionen), AKTIONEN_ANZAHL)
        self.assertEqual([b.nr for b in s.bloecke], [1, 2, 3])
        self.assertEqual(s.seiten, ["Wallbox", "Einschätzung"])
        self.assertTrue(engine.ist_wb_konfigurator(s))
        self.assertFalse(engine.ist_wb_konfigurator(self.logik))
        self.assertTrue(any(a.art == "sparte" and a.antwort == "WB" for a in s.anhaenge))
        # WB04 nur bei Erdarbeiten = Ja (Spalte „Anzeigen wenn“)
        self.assertFalse(engine.ist_sichtbar(s.fragen["WB04"], {"WB03": "Nein"}, s.fragen, s))
        self.assertTrue(engine.ist_sichtbar(s.fragen["WB04"], {"WB03": "Ja"}, s.fragen, s))
        self.assertTrue(engine.ist_sichtbar(s.fragen["WB07"], {"WB03": "Nein"}, s.fragen, s))

    def test_sicht_ohne_artikel_ist_reiner_bogen(self):
        l = copy.copy(self.logik)
        l.wb_artikel_fehlen = True
        s = logik_modul.logik_fuer_sparte(l, "WB")
        self.assertEqual((s.sparte, s.aktionen, s.wb_aktionen), ("WB", [], []))
        self.assertFalse(engine.ist_wb_konfigurator(s))

    def test_ist_freitext_sparte(self):
        self.assertFalse(sparten.ist_freitext_sparte("WB", self.logik))
        for code in ("GW", "MS", "WA", "DMS"):
            self.assertTrue(sparten.ist_freitext_sparte(code, self.logik), code)
        ohne = Logik({}, [], [], [], {}, datetime.now())
        self.assertTrue(sparten.ist_freitext_sparte("WB", ohne))
        self.assertNotIn("WB", sparten.freitext_sparten(self.logik))
        self.assertIn("GW", sparten.freitext_sparten(self.logik))

    def test_vorbelegung(self):
        s = self.wb_sicht
        self.assertEqual(engine.vorbelegung(s.fragen["WB02"], {}, s), "10")
        self.assertEqual(engine.vorbelegung(s.fragen["WB07"], {}, s), "2")
        self.assertEqual(engine.vorbelegung(s.fragen["WB03"], {}, s), "Nein")
        self.assertEqual(engine.vorbelegung(s.fragen["WB05"], {}, s), "Wandmontage")
        self.assertIsNone(engine.vorbelegung(s.fragen["WB01"], {}, s))
        self.assertIsNone(engine.vorbelegung(s.fragen["WB04"], {}, s))
        self.assertEqual(engine.vorbelegung(s.fragen["WB02"], {}, None), "10")
        self.assertEqual(engine.vorbelegung(s.fragen["WB07"], {}, None), "2")

    def test_anhang_regel_artikel_mit_praefix(self):
        wb = _wb()
        ws = wb["Anhänge"]
        ws.append(["Mennekes AMTRON 4You 310.pdf", "wenn Artikel WB001 im Angebot", "Test", None])
        ws.append(["Mennekes AMTRON 4You 560.pdf", "wenn Pos. WB002 im Angebot", "Test", None])
        bericht = Pruefbericht()
        anhaenge = {a.datei: a for a in logik_modul._anhaenge_einlesen(wb, bericht)}
        self.assertEqual((anhaenge["Mennekes AMTRON 4You 310.pdf"].art,
                          anhaenge["Mennekes AMTRON 4You 310.pdf"].positionen), ("position", ["WB001"]))
        self.assertEqual(anhaenge["Mennekes AMTRON 4You 560.pdf"].positionen, ["WB002"])
        self.assertFalse([w for w in bericht.warnungen if "Mennekes" in w])


# --- „Logik prüfen“ an veränderten Blättern ---------------------------------------

class Fehlerfaelle(unittest.TestCase):
    def test_unveraendert_gruen(self):
        logik, bericht = _einlesen(_wb())
        self.assertTrue(bericht.ok, bericht.fehler)
        self.assertEqual(bericht.warnungen, [])
        self.assertEqual(len(logik.wb_aktionen), AKTIONEN_ANZAHL)

    def test_fehlende_referenz(self):
        wb = _wb()
        ws = wb["Aktionen WB"]
        ws.cell(_zeile_finden(ws, 3, "Artikel: WB004 ×1"), 3).value = "Artikel: WB099 ×1"
        _logik, bericht = _einlesen(wb)
        self.assertTrue(any("WB-Referenz WB099 fehlt im Blatt „WB-Artikel“" in f for f in bericht.fehler),
                        bericht.fehler)

    def test_mehrlaenge_nur_bei_wb02(self):
        wb = _wb()
        ws = wb["Aktionen WB"]
        ws.cell(_zeile_finden(ws, 3, "Artikel: WB001 ×1"), 3).value = "Artikel: WB001 × Mehrlänge"
        _logik, bericht = _einlesen(wb)
        self.assertTrue(any("Mehrlänge" in f and "WB01" in f for f in bericht.fehler), bericht.fehler)

    def test_mengenwort_muss_zahlenfrage_sein(self):
        """„× WB09“ (unbekannte Frage) und „× WB01“ (Auswahl statt Zahl) sind Fehler."""
        wb = _wb()
        ws = wb["Aktionen WB"]
        ws.cell(_zeile_finden(ws, 3, "Artikel: WB006 × WB04"), 3).value = "Artikel: WB006 × WB09"
        _logik, bericht = _einlesen(wb)
        self.assertTrue(any("Mengenwort" in f and "WB09" in f for f in bericht.fehler), bericht.fehler)
        wb = _wb()
        ws = wb["Aktionen WB"]
        ws.cell(_zeile_finden(ws, 3, "Artikel: WB008 × (WB07 − 2)"), 3).value = "Artikel: WB008 × (WB01 − 2)"
        _logik, bericht = _einlesen(wb)
        self.assertTrue(any("Mengenwort" in f and "WB01" in f for f in bericht.fehler), bericht.fehler)

    def test_schwelle_weicht_vom_parameter_ab(self):
        wb = _wb()
        ws = wb["WB-Parameter"]
        ws.cell(_zeile_finden(ws, 1, "leitung_inklusive_m"), 2).value = "15"
        _logik, bericht = _einlesen(wb)
        self.assertTrue(bericht.ok, bericht.fehler)
        self.assertTrue(any("Schwelle" in w and "leitung_inklusive_m" in w for w in bericht.warnungen),
                        bericht.warnungen)
        wb = _wb()
        ws = wb["WB-Parameter"]
        ws.cell(_zeile_finden(ws, 1, "durchbrueche_inklusive"), 2).value = "3"
        _logik, bericht = _einlesen(wb)
        self.assertTrue(bericht.ok, bericht.fehler)
        self.assertTrue(any("Schwelle" in w and "durchbrueche_inklusive" in w for w in bericht.warnungen),
                        bericht.warnungen)

    def test_parameter_fehlt_und_unlesbar(self):
        wb = _wb()
        ws = wb["WB-Parameter"]
        ws.delete_rows(_zeile_finden(ws, 1, "fundament_position"))
        logik, bericht = _einlesen(wb)
        self.assertTrue(bericht.ok, bericht.fehler)
        self.assertTrue(any("fundament_position" in w and "Standardwert WB007" in w
                            for w in bericht.warnungen), bericht.warnungen)
        self.assertEqual(logik_modul.wb_parameter_fehlende(logik), ["fundament_position"])
        self.assertEqual(wb_auslegung.parameter(logik)["fundament_position"], "WB007")
        wb = _wb()
        ws = wb["WB-Parameter"]
        ws.cell(_zeile_finden(ws, 1, "leitung_inklusive_m"), 2).value = "zehn"
        ws.cell(_zeile_finden(ws, 1, "durchbrueche_inklusive"), 2).value = "zwei"
        _logik, bericht = _einlesen(wb)
        self.assertTrue(any("leitung_inklusive_m" in f and "keine Zahl" in f for f in bericht.fehler),
                        bericht.fehler)
        self.assertTrue(any("durchbrueche_inklusive" in f and "keine Zahl" in f for f in bericht.fehler),
                        bericht.fehler)

    def test_wallbox_standard_muss_option_sein(self):
        wb = _wb()
        ws = wb["WB-Parameter"]
        ws.cell(_zeile_finden(ws, 1, "wallbox_standard"), 2).value = "Mennekes AMTRON 4You 560"
        logik, bericht = _einlesen(wb)
        self.assertTrue(bericht.ok, bericht.fehler)                       # eindeutiger Präfix
        sicht = logik_modul.logik_fuer_sparte(logik, "WB")
        self.assertEqual(engine.vorbelegung(sicht.fragen["WB01"], {}, sicht), WB560)
        ws.cell(_zeile_finden(ws, 1, "wallbox_standard"), 2).value = "Tesla Wall Connector"
        _logik, bericht = _einlesen(wb)
        self.assertTrue(any("wallbox_standard" in f for f in bericht.fehler), bericht.fehler)

    def test_unbekannte_frage_und_fehlende_pflichtfrage(self):
        wb = _wb()
        ws = wb["Aktionen WB"]
        ws.append(["WB09", "Ja", "Artikel: WB001 ×1", None, None])
        _logik, bericht = _einlesen(wb)
        self.assertTrue(any("unbekannte Frage „WB09“" in f for f in bericht.fehler), bericht.fehler)
        for fid in ("WB02", "WB04", "WB07"):
            wb = _wb()
            ws = wb["Fragen WB"]
            ws.delete_rows(_zeile_finden(ws, 2, fid))
            _logik, bericht = _einlesen(wb)
            self.assertTrue(any(f"Frage {fid} fehlt" in f for f in bericht.fehler), (fid, bericht.fehler))

    def test_abdeckung_wb01(self):
        wb = _wb()
        ws = wb["Aktionen WB"]
        ws.delete_rows(_zeile_finden(ws, 3, "Artikel: WB002 ×1"))
        _logik, bericht = _einlesen(wb)
        self.assertTrue(any("keine Aktionszeile für Option" in w and "560" in w
                            for w in bericht.warnungen), bericht.warnungen)

    def test_ohne_aktionen_wb_bleibt_freitext(self):
        wb = _wb()
        wb.remove(wb["Aktionen WB"])
        logik, bericht = _einlesen(wb)
        self.assertTrue(bericht.ok, bericht.fehler)
        self.assertEqual(logik.wb_aktionen, [])
        self.assertTrue(sparten.ist_freitext_sparte("WB", logik))
        sicht = logik_modul.logik_fuer_sparte(logik, "WB")
        self.assertEqual(sicht.aktionen, [])          # reiner Bogen (Fragen WB vorhanden)


# --- Rechenkern (ohne Datenbank) ----------------------------------------------------

class Stamm:
    def __init__(self, nr, preis, einheit="Stück", ek=None):
        self.pos_nr = nr
        self.bezeichnung = f"Artikel {nr}"
        self.beschreibung = "Langtext"
        self.einheit = einheit
        self.e_preis_cent = preis
        self.ep_flag = False
        self.ek_cent = ek
        self.guid = None


# VK laut Zulieferung „WB Angebot Excel 1“ (Nachtrag 15.10.2026)
ARTIKEL_MAP = {"WB001": Stamm("WB001", 93850), "WB002": Stamm("WB002", 139055),
               "WB003": Stamm("WB003", 129720, "psl."), "WB004": Stamm("WB004", 0, "psl."),
               "WB005": Stamm("WB005", 4934, "m"), "WB006": Stamm("WB006", 12350, "m"),
               "WB007": Stamm("WB007", 102000, "psl."), "WB008": Stamm("WB008", 16900)}


def antworten(wallbox=WB310, leitung=10, erdarbeiten="Nein", erd_m="", montage="Wandmontage",
              durchbrueche=2, bemerkung=""):
    a = {"WB01": wallbox, "WB02": leitung, "WB03": erdarbeiten, "WB05": montage,
         "WB07": durchbrueche, "WB08": bemerkung, "S01": "warm", "S02": ""}
    if erdarbeiten == "Ja":
        a["WB04"] = erd_m
    return a


FALL_A = antworten()                                       # 310, 10 m, keine Erdarbeiten, Wand, 2 Durchbrüche
FALL_B = antworten(WB560, 25, "Ja", 8, "Säule")            # 560, 25 m, Erdarbeiten 8 m, Säule
FALL_B2 = antworten(WB560, 25, "Ja", 8, "Säule", 4)        # Fall B + 4 Durchbrüche
FALL_C = antworten(WB560, 8)
FALL_D = antworten(WB310, "12,4")


class Rechenkern(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.logik, _b = wb_sicht()

    def test_mehrlaenge_und_ganze_zahl_kaufmaennisch(self):
        self.assertEqual(wb_auslegung.mehrlaenge(10, 10.0), 0)
        self.assertEqual(wb_auslegung.mehrlaenge(8, 10.0), 0)
        self.assertEqual(wb_auslegung.mehrlaenge(25, 10.0), 15)
        self.assertEqual(wb_auslegung.mehrlaenge(12.4, 10.0), 2)
        self.assertEqual(wb_auslegung.mehrlaenge(12.5, 10.0), 3)        # kaufmännisch, nicht Banker
        self.assertEqual(wb_auslegung.mehrlaenge("12,4", 10.0), 2)
        self.assertEqual(wb_auslegung.mehrlaenge("", 10.0), 0)
        self.assertEqual(wb_auslegung.mehrlaenge(None, 10.0), 0)
        self.assertEqual(wb_auslegung.mehrlaenge(10.4, 10.0), 0)
        for wert, erwartet in ((8, 8), (8.5, 9), ("7,4", 7), ("7,5", 8), (0, 0), ("", 0), (None, 0),
                               (-3, 0), ("abc", 0), (4, 4)):
            self.assertEqual(wb_auslegung.ganze_zahl(wert), erwartet, wert)
        self.assertIs(wb_auslegung.ganze_meter, wb_auslegung.ganze_zahl)

    def test_parameter_und_kurzname(self):
        p = wb_auslegung.parameter(self.logik)
        self.assertEqual((p["leitung_inklusive_m"], p["meterposition_leitung"], p["erdarbeiten_position"],
                          p["fundament_position"], p["durchbrueche_inklusive"], p["durchbruch_position"],
                          p["wallbox_standard"], p["fehlende"]),
                         (10.0, "WB005", "WB006", "WB007", 2, "WB008", "", []))
        self.assertEqual(wb_auslegung.wallbox_kurz(WB560), "AMTRON 4You 560")
        self.assertEqual(wb_auslegung.wallbox_kurz(WB310), "AMTRON 4You 310")
        self.assertEqual(wb_auslegung.wallbox_kurz(""), "")
        self.assertEqual(wb_auslegung.meter_text(25), "25 m")
        self.assertEqual(wb_auslegung.meter_text(12.4), "12,4 m")

    def test_auslegen_fall_b(self):
        a = wb_auslegung.auslegen(self.logik, FALL_B)
        self.assertEqual((a.wallbox, a.wallbox_ref, a.leitung_m, a.mehrlaenge_m),
                         ("AMTRON 4You 560", "WB002", 25.0, 15))
        self.assertEqual((a.erdarbeiten, a.erdarbeiten_m, a.montage, a.saeule), (True, 8, "Säule", True))
        self.assertEqual((a.durchbrueche, a.durchbrueche_inklusive, a.durchbrueche_zusatz), (2, 2, 0))
        self.assertEqual(a.positionen, ["WB002", "WB003", "WB005", "WB006", "WB007", "WB004"])
        self.assertEqual(a.warnungen, [])
        self.assertEqual(wb_auslegung.protokoll_zeilen(a),
                         ["Wallbox: AMTRON 4You 560",
                          "Leitungsverlegung gesamt: 25 m, davon 10 m enthalten, Mehrlänge 15 m",
                          "Erdarbeiten: 8 m",
                          "Montage: Säule (Fundament und Standsäule)",
                          "Durchbrüche: 2 (in der Anschlussvorbereitung enthalten)"])
        self.assertEqual(wb_auslegung.auslegungs_text(a), "\n".join(wb_auslegung.protokoll_zeilen(a)))
        d = a.als_dict()
        self.assertEqual((d["wallbox"], d["mehrlaenge_m"], d["leitung_inklusive_m"], d["erdarbeiten"],
                          d["erdarbeiten_m"], d["montage"], d["durchbrueche"], d["durchbrueche_zusatz"]),
                         ("AMTRON 4You 560", 15, 10.0, True, 8, "Säule", 2, 0))

    def test_auslegen_fall_b2_durchbrueche(self):
        a = wb_auslegung.auslegen(self.logik, FALL_B2)
        self.assertEqual((a.durchbrueche, a.durchbrueche_zusatz), (4, 2))
        self.assertEqual(a.positionen, ["WB002", "WB003", "WB008", "WB005", "WB006", "WB007", "WB004"])
        self.assertEqual(wb_auslegung.protokoll_zeilen(a)[-1],
                         "Durchbrüche: 4 (2 in der Anschlussvorbereitung enthalten, 2 zusätzlich)")
        # „3,5“ Durchbrüche → kaufmännisch 4; 1 Durchbruch → kein Zusatz, nie negativ
        a = wb_auslegung.auslegen(self.logik, antworten(WB310, 10, durchbrueche="3,5"))
        self.assertEqual((a.durchbrueche, a.durchbrueche_zusatz), (4, 2))
        a = wb_auslegung.auslegen(self.logik, antworten(WB310, 10, durchbrueche=1))
        self.assertEqual((a.durchbrueche, a.durchbrueche_zusatz), (1, 0))
        self.assertNotIn("WB008", a.positionen)

    def test_auslegen_fall_a_c_und_erdarbeiten(self):
        a = wb_auslegung.auslegen(self.logik, FALL_A)
        self.assertEqual(a.positionen, ["WB001", "WB003", "WB004"])
        self.assertEqual(wb_auslegung.protokoll_zeilen(a)[1:],
                         ["Leitungsverlegung gesamt: 10 m – bis 10 m enthalten, keine Mehrlänge",
                          "Erdarbeiten: keine", "Montage: Wandmontage",
                          "Durchbrüche: 2 (in der Anschlussvorbereitung enthalten)"])
        c = wb_auslegung.auslegen(self.logik, FALL_C)
        self.assertEqual(c.positionen, ["WB002", "WB003", "WB004"])
        self.assertEqual((c.mehrlaenge_m, c.erdarbeiten, c.saeule), (0, False, False))
        # Erdarbeiten unabhängig von Wandmontage/Säule; „7,5“ m → 8 m
        e = wb_auslegung.auslegen(self.logik, antworten(WB310, 10, "Ja", "7,5", "Wandmontage"))
        self.assertEqual(e.positionen, ["WB001", "WB003", "WB006", "WB004"])
        self.assertEqual(e.erdarbeiten_m, 8)
        self.assertEqual(wb_auslegung.erdarbeiten_text(e), "Erdarbeiten: 8 m")
        # Erdarbeiten = Ja ohne Meter: keine Position, Protokoll „Meter fehlen“
        g = wb_auslegung.auslegen(self.logik, antworten(WB310, 10, "Ja", ""))
        self.assertNotIn("WB006", g.positionen)
        self.assertEqual(wb_auslegung.erdarbeiten_text(g), "Erdarbeiten: Meter fehlen")

    def test_positionen_bauen_summen(self):
        def netto(pos):
            return sum(round(p["menge"] * p["e_preis_cent"]) for p in pos)
        a = wb_auslegung.positionen_bauen(self.logik, FALL_A, ARTIKEL_MAP)
        self.assertEqual([(p["pos_nr"], p["menge"], p["gruppe"], p["block_nr"]) for p in a],
                         [("WB001", 1.0, "Wallbox", 1), ("WB003", 1.0, "Installation und Anschluss", 2),
                          ("WB004", 1.0, "Netzanmeldung", 3)])
        self.assertEqual(netto(a), 223570)                                 # Fall A 2.235,70 €
        b = wb_auslegung.positionen_bauen(self.logik, FALL_B, ARTIKEL_MAP)
        self.assertEqual([(p["pos_nr"], p["menge"]) for p in b],
                         [("WB002", 1.0), ("WB003", 1.0), ("WB005", 15.0), ("WB006", 8.0),
                          ("WB007", 1.0), ("WB004", 1.0)])
        self.assertEqual(netto(b), 543585)                                 # Fall B 5.435,85 €
        self.assertEqual([p["einheit"] for p in b], ["Stück", "psl.", "m", "m", "psl.", "psl."])
        b2 = wb_auslegung.positionen_bauen(self.logik, FALL_B2, ARTIKEL_MAP)
        self.assertEqual([(p["pos_nr"], p["menge"], p["block_nr"]) for p in b2],
                         [("WB002", 1.0, 1), ("WB003", 1.0, 2), ("WB008", 2.0, 2), ("WB005", 15.0, 2),
                          ("WB006", 8.0, 2), ("WB007", 1.0, 2), ("WB004", 1.0, 3)])
        self.assertEqual(netto(b2), 577385)                                # Fall B2 5.773,85 €
        self.assertEqual(b2[2]["gruppe"], "Installation und Anschluss")
        c = wb_auslegung.positionen_bauen(self.logik, FALL_C, ARTIKEL_MAP)
        self.assertEqual(netto(c), 268775)                                 # Fall C 2.687,75 €
        d = wb_auslegung.positionen_bauen(self.logik, FALL_D, ARTIKEL_MAP)
        self.assertEqual([(p["pos_nr"], p["menge"]) for p in d if p["pos_nr"] == "WB005"], [("WB005", 2.0)])
        self.assertEqual(netto(d), 233438)                                 # Fall D 2.334,38 €
        self.assertFalse(any(p["bezeichnung"].startswith("Auslegung") for p in b2))   # keine Auslegungszeile
        self.assertTrue(all(p["sort"] == i for i, p in enumerate(b2, 1)))

    def test_validierung_und_warnungen(self):
        self.assertEqual(wb_auslegung.validierung(FALL_A), [])
        self.assertEqual(wb_auslegung.validierung(FALL_B2), [])
        self.assertEqual(wb_auslegung.validierung({"WB01": WB310, "WB02": ""}),
                         ["Leitungslänge fehlt oder ist keine Zahl"])
        self.assertEqual(wb_auslegung.validierung({"WB01": WB310, "WB02": "abc"}),
                         ["Leitungslänge fehlt oder ist keine Zahl"])
        self.assertEqual(wb_auslegung.validierung({"WB01": "", "WB02": 12}), ["Wallbox auswählen"])
        self.assertEqual(wb_auslegung.validierung({}),
                         ["Wallbox auswählen", "Leitungslänge fehlt oder ist keine Zahl"])
        # Fall G: Erdarbeiten = Ja ohne Meter (leer, 0, keine Zahl) → Meldung
        for erd_m in ("", 0, "0", "abc", None):
            self.assertEqual(wb_auslegung.validierung(antworten(WB310, 10, "Ja", erd_m)),
                             ["Erdarbeiten: Meter fehlen"], erd_m)
        self.assertEqual(wb_auslegung.validierung(antworten(WB310, 10, "Ja", "0,4")),
                         ["Erdarbeiten: Meter fehlen"])          # rundet auf 0 m
        self.assertEqual(wb_auslegung.validierung(antworten(WB310, 10, "Nein", "")), [])
        self.assertEqual(wb_auslegung.validierung(antworten(WB310, 10, durchbrueche="abc")),
                         ["Durchbrüche: Anzahl fehlt oder ist keine Zahl"])
        self.assertEqual(wb_auslegung.validierung(antworten(WB310, 10, durchbrueche=0)), [])
        self.assertEqual(wb_auslegung.warnungen(antworten(WB310, 100)), [])
        self.assertEqual(wb_auslegung.warnungen(antworten(WB310, 100.5)),
                         ["Leitungslänge über 100 m – bitte prüfen"])
        a = wb_auslegung.auslegen(self.logik, antworten(WB310, 120))
        self.assertEqual(a.warnungen, ["Leitungslänge über 100 m – bitte prüfen"])
        self.assertEqual(a.mehrlaenge_m, 110)                              # keine Sperre
        self.assertIn("Leitungslänge über 100 m – bitte prüfen",
                      engine.fachliche_hinweise(antworten(WB310, 120), self.logik))
        self.assertEqual(engine.ampel_gruende(self.logik, antworten(WB310, 120)), [])   # keine AMPEL
        self.assertEqual(engine.ampel_gruende(self.logik, FALL_B2), [])

    def test_protokoll_eintraege(self):
        self.assertEqual(wb_auslegung.protokoll_eintraege(self.logik, {"WB01": WB310}), [])   # unvollständig
        eintraege = wb_auslegung.protokoll_eintraege(self.logik, FALL_B2)
        self.assertEqual([(e["seite"], e["frage"], e["antwort"]) for e in eintraege],
                         [("Auslegung", "Wallbox", "AMTRON 4You 560"),
                          ("Auslegung", "Leitungsverlegung gesamt", "25 m, davon 10 m enthalten, Mehrlänge 15 m"),
                          ("Auslegung", "Erdarbeiten", "8 m"),
                          ("Auslegung", "Montage", "Säule (Fundament und Standsäule)"),
                          ("Auslegung", "Durchbrüche", "4 (2 in der Anschlussvorbereitung enthalten, 2 zusätzlich)")])
        prot = engine.protokoll(self.logik, FALL_B2)
        self.assertEqual([e["frage_id"] for e in prot][:7],
                         ["WB01", "WB02", "WB03", "WB04", "WB05", "WB07", "WB08"])
        self.assertEqual([e["seite"] for e in prot][-5:], ["Auslegung"] * 5)
        prot_a = engine.protokoll(self.logik, FALL_A)
        self.assertNotIn("WB04", [e["frage_id"] for e in prot_a])     # unsichtbar bei Nein
        prot_warn = engine.protokoll(self.logik, antworten(WB310, 150))
        self.assertEqual(prot_warn[-1]["antwort"], "Leitungslänge über 100 m – bitte prüfen")


if __name__ == "__main__":
    unittest.main()
